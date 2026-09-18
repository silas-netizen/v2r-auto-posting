from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

import pytest

from app.alerts.telegram import TelegramChannel
from app.browser.publish import RecordingBrowser, publish_planned_slots
from app.collect.public import PublicPageReader
from app.importers import load_adapted_csv, load_account_workbook
from app.models import JobStatus, TaskSpec
from app.providers.anthropic_provider import AnthropicProvider
from app.providers.router import ModelRouter
from app.sources import SourceRef, parse_daily_rows, sync_source
from app.store import JobStore
from app.worker import CommandRuntime


class FakeResponse:
    def __init__(self, body: bytes, headers: dict[str, str] | None = None, url: str = ""):
        self._body = body
        self.headers = headers or {"Content-Type": "text/html"}
        self._url = url

    def read(self) -> bytes:
        return self._body

    def geturl(self) -> str:
        return self._url

    def __enter__(self) -> "FakeResponse":
        return self

    def __exit__(self, *_args) -> None:
        return None


def test_store_is_idempotent_and_leases_one_job(tmp_path: Path) -> None:
    store = JobStore(tmp_path / "v2r.sqlite")
    first = store.enqueue(TaskSpec(task="status"), idempotency_key="same")
    second = store.enqueue(TaskSpec(task="status"), idempotency_key="same")
    assert first.id == second.id
    taken = store.acquire("pc-a")
    assert taken is not None
    blocked = store.acquire("pc-b")
    assert blocked is None
    store.finish(taken.id, JobStatus.PUBLISHED, result={"ok": True})
    later = store.acquire("pc-b")
    assert later is None


def test_runtime_handles_korean_daily_command(tmp_path: Path) -> None:
    root = tmp_path
    (root / "config").mkdir()
    (root / "config" / "accounts.json").write_text(
        json.dumps(
            [
                {
                    "login_id": f"own{index}",
                    "work_type": "자사 카페",
                    "linked": "V2R",
                    "excluded": False,
                    "grade": "",
                    "nickname": f"닉{index}",
                    "shade": "",
                }
                for index in range(1, 6)
            ]
        ),
        encoding="utf-8",
    )
    store = JobStore(root / "data" / "v2r.sqlite")
    store.save_source(
        "일상글목록",
        {
            "manuscripts": [
                {
                    "title": f"아침 {index}",
                    "body": f"본문 {index}",
                    "cafe": "고요한 아침",
                    "board": "자유게시판",
                    "source": "sheet",
                    "keyword": "",
                    "account": "",
                    "comments": [],
                }
                for index in range(1, 6)
            ]
        },
    )
    (root / "config" / "sources.json").write_text(
        json.dumps({"sources": [{"name": "일상글목록", "url": "https://example.invalid", "kind": "sheet"}]}),
        encoding="utf-8",
    )
    runtime = CommandRuntime(root, store=store, owner="pc-test")
    result = runtime.handle_text(
        "내일 오전 9시부터 오후 6시까지만 일상 글 3개 올려줘. 아이디 3개 자동으로 쓰고 10분 간격으로 해줘."
    )
    assert result["ok"] is True
    processed = result["processed"]
    assert processed["status"] == "published"
    assert processed["result"]["planned"] == 3
    assert processed["result"]["dry_run"] is True
    assert processed["result"]["registered"] == 0


def test_telegram_rejects_unknown_chat_and_unknown_task() -> None:
    channel = TelegramChannel(token="x", allowed_chat_ids=["111"])
    spec = channel.parse_update(
        {"message": {"chat": {"id": 111}, "text": "상태 알려줘"}}
    )
    assert spec is not None
    assert spec.task == "status"
    with pytest.raises(PermissionError):
        channel.parse_update({"message": {"chat": {"id": 999}, "text": "상태 알려줘"}})
    with pytest.raises(ValueError):
        channel.parse_update({"message": {"chat": {"id": 111}, "text": "서버 셸 열어줘"}})


def test_public_reader_stops_at_login_wall() -> None:
    def opener(_request, timeout=8):
        return FakeResponse(
            "<html><title>Login</title><body>로그인 해주세요</body></html>".encode("utf-8"),
            url="https://example.invalid/login",
        )

    result = PublicPageReader(opener=opener).fetch("https://example.invalid/login")
    assert result.blocked is True
    assert result.reason == "authentication required"


def test_public_reader_extracts_public_article() -> None:
    html = """
    <html><head>
    <meta property="og:title" content="공개 일상">
    <meta property="og:description" content="공원 산책">
    <script type="application/ld+json">{"@type":"Article","articleBody":"오늘 날씨가 좋아 걷다 왔어요"}</script>
    </head><body>본문</body></html>
    """

    def opener(_request, timeout=8):
        return FakeResponse(html.encode("utf-8"), url="https://example.invalid/post")

    result = PublicPageReader(opener=opener).fetch("https://example.invalid/post")
    assert result.blocked is False
    assert result.title == "공개 일상"
    assert "걷다" in result.text


def test_claude_collect_uses_provider_json(tmp_path: Path) -> None:
    def opener(_request, timeout=30):
        payload = {
            "content": [
                {
                    "type": "text",
                    "text": '[{"title":"수집제목","body":"수집본문","cafe":"고요한 아침","board":"자유게시판"}]',
                }
            ]
        }
        return FakeResponse(json.dumps(payload).encode("utf-8"))

    router = ModelRouter(anthropic=AnthropicProvider(api_key="test-key", opener=opener))
    store = JobStore(tmp_path / "v2r.sqlite")
    runtime = CommandRuntime(tmp_path, store=store, router=router, owner="pc-claude")
    (tmp_path / "config").mkdir()
    (tmp_path / "config" / "sources.json").write_text(
        json.dumps(
            {
                "sources": [
                    {
                        "name": "공개피드",
                        "kind": "public",
                        "url": "https://example.invalid/feed",
                    }
                ]
            }
        ),
        encoding="utf-8",
    )

    def page_opener(_request, timeout=8):
        return FakeResponse(
            "<rss><item><title>원본</title><description>공개 글</description></item></rss>".encode("utf-8"),
            headers={"Content-Type": "application/rss+xml"},
            url="https://example.invalid/feed",
        )

    from app.collect.daily import collect_daily_manuscripts
    from app.collect.public import PublicPageReader

    manuscripts = collect_daily_manuscripts(
        ["https://example.invalid/feed"],
        router,
        reader=PublicPageReader(opener=page_opener),
        cafe="고요한 아침",
        board="자유게시판",
    )
    assert manuscripts[0].title == "수집제목"
    assert manuscripts[0].source == "claude-collect"


def test_sync_keeps_cache_when_source_times_out(tmp_path: Path) -> None:
    store = JobStore(tmp_path / "v2r.sqlite")
    store.save_source("일상글목록", {"manuscripts": [{"title": "캐시", "body": "유지"}]})
    ref = SourceRef(name="일상글목록", url="https://docs.google.com/spreadsheets/d/abc/edit", gid="1")

    def opener(_request, timeout=5):
        raise TimeoutError("slow")

    payload = sync_source(ref, store, opener=opener)
    assert payload["status"] == "cache"


def test_parse_daily_rows_accepts_combined_title_body() -> None:
    csv_text = '제목,본문,카페\n,"제목 : 첫 글\n본문 : 내용입니다",씨씨앙\n'
    rows = parse_daily_rows(csv_text)
    assert rows[0].title == "첫 글"
    assert "내용" in rows[0].body


def test_self_owned_publish_registers_and_verifies_result() -> None:
    browser = RecordingBrowser()
    results = publish_planned_slots(
        browser,
        [
            {
                "title": "자사 테스트",
                "body": "자사 본문",
                "cafe": "고요한 아침",
                "board": "자유게시판",
                "account": "own1",
                "scheduled_at": datetime(2026, 9, 19, 9, 0),
                "comments": [
                    {
                        "text": "첫 댓글",
                        "parent_text": None,
                        "scheduled_at": datetime(2026, 9, 19, 9, 3),
                    }
                ],
            }
        ],
        dry_run=False,
    )
    assert results[0]["status"] == "registered"
    assert results[0]["url"].endswith("/1")
    assert any(step.startswith("verify:") for step in browser.steps)
    assert any(step.startswith("comment:첫 댓글") for step in browser.steps)


def test_affiliate_publish_links_daily_revision_and_comments() -> None:
    browser = RecordingBrowser()
    results = publish_planned_slots(
        browser,
        [
            {
                "workflow": "affiliate",
                "title": "브랜드 수정 글",
                "body": "브랜드 본문",
                "daily": {"title": "오늘 일상", "body": "오늘의 일상 본문"},
                "cafe": "씨씨앙",
                "board": "자유 수다방",
                "account": "aff1",
                "scheduled_at": datetime(2026, 9, 19, 9, 0),
                "revision_at": datetime(2026, 9, 19, 13, 0),
                "comments": [
                    {
                        "text": "제휴 댓글",
                        "parent_text": None,
                        "scheduled_at": datetime(2026, 9, 19, 13, 3),
                    }
                ],
            }
        ],
        dry_run=False,
    )
    result = results[0]
    assert result["daily_url"].endswith("/1")
    assert result["revision_url"].endswith("/2")
    assert result["url"] == result["revision_url"]
    assert any(step.startswith("reserve_revision:") for step in browser.steps)
    assert any(step.startswith("comment:제휴 댓글") for step in browser.steps)


def test_adapted_csv_rows_have_stable_source_identity(tmp_path: Path) -> None:
    path = tmp_path / "adapted.csv"
    path.write_text(
        "카페명,게시판명,게시판링크,작성계정,각색제목,각색본문,등록시각\n"
        "고요한아침,자유게시판,,,첫 제목,첫 본문,#REF!\n",
        encoding="utf-8-sig",
    )
    rows = load_adapted_csv(path, source_key="20260903_각색_전체_1.xlsx")
    assert len(rows) == 1
    assert rows[0].source_row == 2
    assert rows[0].source == "20260903_각색_전체_1.xlsx"
    assert len(rows[0].content_hash) == 64
    assert rows[0].cafe == "고요한 아침"


def test_account_workbook_excludes_gray_rows(tmp_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill

    path = tmp_path / "accounts.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "아이디 리스트"
    sheet.append(["번호", "ID", "PW", "이름", "", "", "", "작업 구분", "연동"])
    sheet.append([1, "own-ok", "secret", "", "", "", "", "자사 카페", "V2R"])
    sheet.append([2, "own-gray", "secret", "", "", "", "", "자사 카페", "V2R"])
    sheet.cell(3, 2).fill = PatternFill("solid", fgColor="D9D9D9")
    sheet.append([3, "affiliate-ok", "secret", "", "", "", "", "제휴 작업", "V2R"])
    workbook.save(path)

    accounts, counts = load_account_workbook(path)
    assert counts["self_v2r"] == 1
    assert counts["affiliate_v2r"] == 1
    assert counts["gray_excluded"] == 1
    assert accounts[1].excluded is True
    assert all(not hasattr(account, "password") for account in accounts)


def test_publication_identity_prevents_repeat(tmp_path: Path) -> None:
    store = JobStore(tmp_path / "v2r.sqlite")
    assert not store.publication_exists(
        source_key="source",
        row_number=2,
        content_hash="abc",
    )
    store.mark_publication(
        source_key="source",
        row_number=2,
        content_hash="abc",
        status="registered",
        url="https://v2r.example/posts/1",
        account="own1",
        cafe="고요한 아침",
    )
    assert store.publication_exists(
        source_key="source",
        row_number=2,
        content_hash="abc",
    )
    assert store.publication_count("source") == 1


def test_uncertain_checkpoint_blocks_automatic_repeat(tmp_path: Path) -> None:
    store = JobStore(tmp_path / "v2r.sqlite")
    store.mark_publication(
        source_key="source",
        row_number=3,
        content_hash="def",
        status="uncertain",
        account="own1",
        cafe="고요한 아침",
    )
    assert store.publication_exists(
        source_key="source",
        row_number=3,
        content_hash="def",
    )
    assert store.publication_count("source") == 0


def test_checkpoint_is_written_before_registration_attempt() -> None:
    class DisconnectingBrowser(RecordingBrowser):
        def register(self) -> str:
            raise ConnectionError("network disconnected")

    events: list[str] = []
    with pytest.raises(ConnectionError):
        publish_planned_slots(
            DisconnectingBrowser(),
            [
                {
                    "title": "끊김 테스트",
                    "body": "본문",
                    "cafe": "고요한 아침",
                    "board": "자유게시판",
                    "account": "own1",
                    "source_key": "source",
                    "source_row": 2,
                    "content_hash": "abc",
                    "scheduled_at": datetime(2026, 9, 19, 9, 0),
                }
            ],
            dry_run=False,
            checkpoint=lambda stage, _item: events.append(stage),
        )
    assert events == ["submitting"]
