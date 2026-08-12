import logging
import threading
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook

from v2r_auto.immediate_api import (
    BOARD_ALIASES,
    MANAGER_ACCOUNTS,
    ImmediateApiPublisher,
    SELF_COMMENT_ACCOUNTS,
)
from v2r_auto.cafe_catalog import CafeMenu, match_catalog_name, normalized_name
from v2r_auto.immediate_inputs import (
    format_daily_body,
    is_informational_sheet,
    load_account_test_jobs,
    load_brand_immediate_jobs,
    load_daily_excel_jobs,
)
from v2r_auto.models import JobStatus
from v2r_auto.history import HistoryStore
from v2r_auto.runner import ImmediateRunner, assign_immediate_schedules


def test_load_brand_sheet_with_board_and_optional_comments(tmp_path: Path) -> None:
    path = tmp_path / "brand.csv"
    path.write_text(
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        '"키 워드","제목 : 제목\n본문 : 본문\n댓글1: 질문",고요한아침,,질문형,,,비실명,,극복후기\n'
        '"일상","제목 : 일상 제목\n본문 : 일상 본문",고요한아침,writer,,,,비실명,Y,오늘의 한 끼\n',
        encoding="utf-8-sig",
    )

    jobs = load_brand_immediate_jobs(path, brand="팥순이")

    assert len(jobs) == 2
    assert jobs[0].board == "극복후기"
    assert jobs[0].brand == "팥순이"
    assert jobs[0].comments
    assert jobs[0].use_comment_ai is False
    assert jobs[1].comments == []
    assert jobs[1].image_disabled is True


def test_brand_sheet_preserves_body_unless_special_format_is_enabled(
    tmp_path: Path,
) -> None:
    path = tmp_path / "brand.csv"
    path.write_text(
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        '"키워드","제목 : 제목\n본문 : 첫 문장이에요. 두 번째 문장이에요, 세 번째 문장이에요😊",'
        "고요한아침,writer,,,,,,가입인사\n",
        encoding="utf-8-sig",
    )

    regular = load_brand_immediate_jobs(path, brand="팥순이")[0]
    informational = load_brand_immediate_jobs(
        path,
        brand="",
        format_body=True,
        use_comment_ai=True,
    )[0]

    assert "." in regular.body and "," in regular.body and "😊" in regular.body
    assert "." not in informational.body
    assert "," not in informational.body
    assert "😊" not in informational.body
    assert regular.use_comment_ai is False
    assert informational.use_comment_ai is True
    publisher = ImmediateApiPublisher(None, logging.getLogger("test"))
    for job in (regular, informational):
        job.cafe_id = 14567700
        job.menu_id = 34
        job.account = "writer"
        job.canonical_cafe_name = "고요한 아침"
        job.canonical_board_name = "가입인사"
        job.scheduled_at = datetime(2026, 8, 12, 12, 0, tzinfo=timezone.utc)
    assert publisher._destination(regular)["use_comment_ai"] is False
    assert publisher._destination(informational)["use_comment_ai"] is True


def test_one_malformed_sheet_row_does_not_abort_other_rows(tmp_path: Path) -> None:
    path = tmp_path / "brand.csv"
    path.write_text(
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        '"오류","제목 :\n본문 : 본문",고요한아침,writer,,,,,,가입인사\n'
        '"정상","제목 : 정상 제목\n본문 : 정상 본문",고요한아침,writer,,,,,,가입인사\n',
        encoding="utf-8-sig",
    )

    jobs = load_brand_immediate_jobs(path, brand="", format_body=True)

    assert len(jobs) == 2
    assert jobs[0].status == JobStatus.FAILED
    assert "제목" in jobs[0].message
    assert jobs[1].status == JobStatus.PENDING
    assert jobs[1].title == "정상 제목"


def test_load_daily_excel_a_to_d(tmp_path: Path) -> None:
    path = tmp_path / "daily.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "일상"
    sheet.append(["카페명", "게시판명", "각색제목", "각색본문"])
    sheet.append(["고요한아침", "가입인사", "안녕하세요", "반갑습니다"])
    workbook.save(path)

    jobs = load_daily_excel_jobs(path)

    assert len(jobs) == 1
    assert jobs[0].source_kind == "daily"
    assert jobs[0].title == "안녕하세요"
    assert jobs[0].body == "반갑습니다"
    assert jobs[0].comments == []
    assert jobs[0].use_comment_ai is True


def test_loads_only_checked_one_line_account_tests(tmp_path: Path) -> None:
    path = tmp_path / "account-tests.csv"
    path.write_text(
        "번호,ID,작업 구분,연동,가아사 조건,실/비실,테스트 선택,테스트 결과,테스트 링크,테스트 일시\n"
        "12,checked-id,,,,,TRUE,,,\n"
        "13,unchecked-id,,,,,FALSE,,,\n",
        encoding="utf-8-sig",
    )

    jobs = load_account_test_jobs(path)

    assert len(jobs) == 1
    assert jobs[0].account == "checked-id"
    assert jobs[0].title == "김천kb보험 그라래12"
    assert jobs[0].body == "김천kb보험 그라래12\n김천kb보험 그라래12"
    assert jobs[0].source_kind == "account_test"
    assert jobs[0].use_comment_ai is False
    assert jobs[0].validate() == []


def test_formats_punctuation_free_daily_body_into_two_sentence_paragraphs() -> None:
    body = (
        "어제 남편과 드라이브를 다녀왔어요 "
        "오랜만에 날씨가 좋아서 기분이 좋았어요 "
        "근처 카페 분위기가 정말 좋더라고요 "
        "다음에는 친구들과 함께 가보고 싶어요"
    )

    assert format_daily_body(body) == (
        "어제 남편과 드라이브를 다녀왔어요\n"
        "오랜만에 날씨가 좋아서 기분이 좋았어요\n\n"
        "근처 카페 분위기가 정말 좋더라고요\n"
        "다음에는 친구들과 함께 가보고 싶어요"
    )


def test_daily_body_keeps_one_two_sentences_and_existing_lines() -> None:
    short = "오늘 산책을 다녀왔어요 날씨가 정말 좋았어요"
    formatted = "오늘 산책을 다녀왔어요\n\n날씨가 정말 좋았어요"

    assert format_daily_body(short) == short
    assert format_daily_body(formatted) == formatted


def test_daily_body_cleans_punctuation_and_picture_emoji_but_keeps_numbers() -> None:
    body = (
        "체온은 37.5도였어요😊 "
        "가격은 1,500만원이었어요... "
        "정말 놀랐어요ㅠㅠ"
    )

    formatted = format_daily_body(body)

    assert "37.5" in formatted
    assert "1,500" in formatted
    assert "😊" not in formatted
    assert "..." not in formatted
    assert "ㅠㅠ" in formatted
    assert formatted.count("\n") == 3


def test_long_two_sentence_daily_body_adds_paragraph_break() -> None:
    body = (
        "오늘은 가족들과 오랜만에 멀리 있는 공원까지 산책을 다녀와서 "
        "이야기도 많이 나누고 여유롭게 시간을 보냈어요 "
        "집으로 돌아오는 길에는 근처 시장에도 들러서 저녁거리와 과일을 "
        "사고 다음 주말 계획도 같이 정했어요"
    )

    assert "\n\n" in format_daily_body(body)


def test_only_known_information_sheet_url_enables_special_formatting() -> None:
    assert is_informational_sheet(
        "https://docs.google.com/spreadsheets/d/"
        "1vSON0Rej9anDQXcAOXyBrCr50B4MMqZ79FahF4cDPJw/edit"
        "?gid=1193993260#gid=1193993260"
    )
    assert not is_informational_sheet(
        "https://docs.google.com/spreadsheets/d/other/edit?gid=1193993260"
    )


class FakeImmediatePublisher(ImmediateApiPublisher):
    def _capture_authorization(self) -> None:
        return None

    def _request(self, method, path, payload=None, query=None):
        if path == "/naver_cafes/naver_join_cafes":
            return {
                "naver_join_cafes": [
                    {"cafe_id": 14567700, "pc_cafe_name": "고요한 아침"}
                ]
            }
        if path == "/navers/accounts":
            accounts = [
                {
                    "naver_login_id": "writer-a",
                    "my_info_v2": {"is_real_name": True},
                },
                {
                    "naver_login_id": "writer-b",
                    "my_info_v2": {"is_real_name": True},
                },
                {
                    "naver_login_id": "writer-alias",
                    "my_info_v2": {"is_real_name": False},
                },
                *[
                    {
                        "naver_login_id": account,
                        "my_info_v2": {"is_real_name": True},
                    }
                    for account in MANAGER_ACCOUNTS
                ],
                *[
                    {
                        "naver_login_id": account,
                        "my_info_v2": {"is_real_name": False},
                    }
                    for account in SELF_COMMENT_ACCOUNTS
                ],
            ]
            return {"accounts": accounts}
        if path == "/naver_cafes/naver_join_cafe":
            return {
                "naver_join_cafe": [
                    {"login_id": "writer-a"},
                    {"login_id": "writer-b"},
                    {"login_id": "writer-alias"},
                    *[{"login_id": account} for account in MANAGER_ACCOUNTS],
                    *[{"login_id": account} for account in SELF_COMMENT_ACCOUNTS],
                ]
            }
        if path == "/naver_cafes/board_histories" or path == "/naver_cafe_articles/board_histories":
            return {"histories": []}
        if path == "/naver_cafes/menus":
            return {
                "cafe_menus": [
                    {"menuId": 34, "menuName": "가입인사🌱", "writable": True},
                    {"menuId": 10, "menuName": "몸 증상", "writable": True},
                ]
            }
        raise AssertionError((method, path, query))


def test_account_tests_resolve_registration_membership_and_alternate_cafes(
    tmp_path: Path,
) -> None:
    path = tmp_path / "account-tests.csv"
    path.write_text(
        "번호,ID,작업 구분,연동,가아사 조건,실/비실,테스트 선택,테스트 결과,테스트 링크,테스트 일시\n"
        "1,both-id,,,,,TRUE,,,\n"
        "2,both-id,,,,,TRUE,,,\n"
        "3,nojoin-id,,,,,TRUE,,,\n"
        "4,missing-id,,,,,TRUE,,,\n"
        "5,login-fail-id,,,,,TRUE,,,\n",
        encoding="utf-8-sig",
    )
    jobs = load_account_test_jobs(path)

    class AccountTestPublisher(ImmediateApiPublisher):
        def _capture_authorization(self):
            return None

        def _request(self, method, request_path, payload=None, query=None):
            if request_path == "/naver_cafes/naver_join_cafes":
                return {
                    "naver_join_cafes": [
                        {
                            "cafe_id": 31670254,
                            "pc_cafe_name": "태극마케팅센터",
                        },
                        {
                            "cafe_id": 31670256,
                            "pc_cafe_name": "소나무마케팅센터",
                        },
                    ]
                }
            if request_path == "/navers/accounts":
                return {
                    "accounts": [
                        {"naver_login_id": "both-id"},
                        {"naver_login_id": "nojoin-id"},
                        {
                            "naver_login_id": "login-fail-id",
                            "is_login_fail": True,
                        },
                    ]
                }
            if request_path == "/naver_cafes/naver_join_cafe":
                return {"naver_join_cafe": [{"login_id": "both-id"}]}
            if request_path == "/naver_cafes/menus":
                return {
                    "cafe_menus": [
                        {
                            "menuId": 1,
                            "menuName": "자유게시판",
                            "writable": True,
                        }
                    ]
                }
            raise AssertionError((method, request_path, query))

    publisher = AccountTestPublisher(None, logging.getLogger("test"))
    publisher.prepare_jobs(jobs)

    assert jobs[0].status == JobStatus.PENDING
    assert jobs[1].status == JobStatus.PENDING
    assert {jobs[0].cafe_id, jobs[1].cafe_id} == {31670254, 31670256}
    assert jobs[2].message == "테스트 카페 미가입"
    assert jobs[3].message == "V2R 미등록 계정"
    assert jobs[4].message == "네이버 로그인 실패"


def test_code_27000_result_includes_detection_and_release_dates() -> None:
    publisher = ImmediateApiPublisher(None, logging.getLogger("test"))
    detected_at = datetime(2026, 8, 12, 3, 0, tzinfo=timezone.utc)
    publisher.restrictions.observe_code_27000(
        source_id="source-27000",
        account="restricted-id",
        reason="error_code 27000",
        now=detected_at,
    )

    result = publisher._restriction_result("restricted-id")

    assert "발견 2026-08-12" in result
    assert "제외 종료 2026-09-11" in result


def test_prepare_jobs_matches_live_ids_and_rotates_all_writers(tmp_path: Path) -> None:
    path = tmp_path / "daily.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["카페명", "게시판명", "각색제목", "각색본문"])
    for index in range(3):
        sheet.append(["고요한아침", "가입인사", f"제목{index}", f"본문{index}"])
    workbook.save(path)
    jobs = load_daily_excel_jobs(path)
    publisher = FakeImmediatePublisher(None, logging.getLogger("test"))

    publisher.prepare_jobs(jobs)

    assert [job.account for job in jobs] == ["writer-a", "writer-b", "writer-a"]
    assert not ({job.account for job in jobs} & MANAGER_ACCOUNTS)
    assert "writer-alias" not in {job.account for job in jobs}
    assert all(job.cafe_id == 14567700 for job in jobs)
    assert all(job.menu_id == 34 for job in jobs)
    assert all(job.status == JobStatus.PENDING for job in jobs)


def test_writer_rotation_does_not_restart_for_each_board(tmp_path: Path) -> None:
    path = tmp_path / "daily.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["카페명", "게시판명", "각색제목", "각색본문"])
    sheet.append(["고요한아침", "가입인사", "첫 글", "첫 본문"])
    sheet.append(["고요한아침", "몸 증상", "둘째 글", "둘째 본문"])
    workbook.save(path)
    jobs = load_daily_excel_jobs(path)
    publisher = FakeImmediatePublisher(None, logging.getLogger("test"))

    publisher.prepare_jobs(jobs)

    assert [job.account for job in jobs] == ["writer-a", "writer-b"]


def test_prepare_jobs_unions_menus_from_all_healthy_accounts(tmp_path: Path) -> None:
    path = tmp_path / "daily.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["카페명", "게시판명", "각색제목", "각색본문"])
    sheet.append(["고요한아침", "신혼 가전 후기", "제목", "본문"])
    workbook.save(path)
    jobs = load_daily_excel_jobs(path)

    class PartialMenuPublisher(FakeImmediatePublisher):
        def _request(self, method, request_path, payload=None, query=None):
            if request_path == "/naver_cafes/menus":
                if query["naver_login_id"] == "writer-a":
                    return {
                        "cafe_menus": [
                            {"menuId": 1, "menuName": "🪻가입인사", "writable": True}
                        ]
                    }
                return {
                    "cafe_menus": [
                        {
                            "menuId": 29,
                            "menuName": "🌸신혼 가전 후기",
                            "writable": True,
                        }
                    ]
                }
            return super()._request(method, request_path, payload, query)

    publisher = PartialMenuPublisher(None, logging.getLogger("test"))
    publisher.prepare_jobs(jobs)

    assert jobs[0].menu_id == 29
    assert jobs[0].canonical_board_name == "🌸신혼 가전 후기"
    assert jobs[0].account == "writer-b"


def test_verified_wedding_board_alias_resolves_v2r_typo() -> None:
    wanted = BOARD_ALIASES[(26680163, normalized_name("웨딩홀 탐방기"))]
    menu = match_catalog_name(
        wanted,
        [CafeMenu(menu_id=5, name="웨딩홀 탑방기")],
        label="게시판",
    )

    assert menu.menu_id == 5


def test_one_unknown_board_does_not_abort_other_rows(tmp_path: Path) -> None:
    path = tmp_path / "daily.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["카페명", "게시판명", "각색제목", "각색본문"])
    sheet.append(["고요한아침", "없는 게시판", "실패 제목", "실패 본문"])
    sheet.append(["고요한아침", "가입인사", "정상 제목", "정상 본문"])
    workbook.save(path)
    jobs = load_daily_excel_jobs(path)
    publisher = FakeImmediatePublisher(None, logging.getLogger("test"))

    publisher.prepare_jobs(jobs)

    assert jobs[0].status == JobStatus.FAILED
    assert "게시판을 찾지 못했습니다" in jobs[0].message
    assert jobs[1].status == JobStatus.PENDING
    assert jobs[1].menu_id == 34


def test_brand_job_can_still_select_non_real_name_account(tmp_path: Path) -> None:
    path = tmp_path / "brand.csv"
    path.write_text(
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        '"키워드","제목 : 제목\n본문 : 본문",고요한아침,,,,,비실명,,가입인사\n',
        encoding="utf-8-sig",
    )
    job = load_brand_immediate_jobs(path, brand="팥순이")[0]
    publisher = FakeImmediatePublisher(None, logging.getLogger("test"))

    publisher.prepare_jobs([job])

    assert job.account == "writer-alias"


def test_schedules_accumulate_five_to_fifteen_minutes_per_cafe(
    tmp_path: Path,
) -> None:
    path = tmp_path / "daily.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["카페명", "게시판명", "각색제목", "각색본문"])
    for index in range(4):
        sheet.append(["고요한아침", "가입인사", f"제목{index}", f"본문{index}"])
    workbook.save(path)
    jobs = load_daily_excel_jobs(path)
    jobs[0].cafe_id = 1
    jobs[1].cafe_id = 1
    jobs[2].cafe_id = 2
    jobs[3].cafe_id = 31670256

    class FixedRandom:
        values = iter((5, 15, 7))

        def randint(self, _minimum, _maximum):
            return next(self.values)

    now = datetime(2026, 8, 11, 1, 0, tzinfo=timezone.utc)
    assign_immediate_schedules(jobs, now=now, rng=FixedRandom())

    assert jobs[0].scheduled_at == now + timedelta(minutes=5)
    assert jobs[1].scheduled_at == now + timedelta(minutes=20)
    assert jobs[2].scheduled_at == now + timedelta(minutes=7)
    assert jobs[3].scheduled_at is None


def test_duplicate_skip_reason_is_written_to_live_log(
    tmp_path: Path,
    caplog,
) -> None:
    path = tmp_path / "daily.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["카페명", "게시판명", "각색제목", "각색본문"])
    sheet.append(["고요한아침", "가입인사", "제목", "본문"])
    workbook.save(path)
    job = load_daily_excel_jobs(path)[0]
    history_path = tmp_path / "history.json"
    HistoryStore(history_path).record(job)

    class DuplicateBrowser:
        def ensure_v2r_login(self, _email, _password):
            return None

        def prepare_immediate_jobs(self, jobs):
            jobs[0].cafe_id = 14567700
            jobs[0].menu_id = 34
            jobs[0].account = "writer"
            jobs[0].canonical_cafe_name = "고요한 아침"
            jobs[0].canonical_board_name = "가입인사"

        def consume_failed_immediate_urls(self):
            return set()

    runner = ImmediateRunner(
        browser=DuplicateBrowser(),
        history_path=history_path,
        report_dir=tmp_path,
        logger=logging.getLogger("duplicate-log-test"),
    )
    with caplog.at_level(logging.INFO, logger="duplicate-log-test"):
        result, _report = runner.run(
            [job],
            dry_run=False,
            stop_event=threading.Event(),
            progress=lambda _current, _total: None,
        )

    assert result.jobs[0].status == JobStatus.SKIPPED
    assert "행 2 건너뜀: 이전에 발행한 동일 글" in caplog.text


def test_pause_waits_then_resumes_before_next_publish(tmp_path: Path) -> None:
    path = tmp_path / "daily.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["카페명", "게시판명", "각색제목", "각색본문"])
    sheet.append(["고요한아침", "가입인사", "제목", "본문"])
    workbook.save(path)
    job = load_daily_excel_jobs(path)[0]

    class PauseBrowser:
        published = 0

        def ensure_v2r_login(self, _email, _password):
            return None

        def prepare_immediate_jobs(self, jobs):
            jobs[0].cafe_id = 14567700
            jobs[0].menu_id = 34
            jobs[0].account = "writer"
            jobs[0].canonical_cafe_name = "고요한 아침"
            jobs[0].canonical_board_name = "가입인사"

        def consume_failed_immediate_urls(self):
            return set()

        def publish_immediate(self, _job, _dry_run):
            self.published += 1
            return ""

    browser = PauseBrowser()
    pause_event = threading.Event()
    pause_event.set()
    runner = ImmediateRunner(
        browser=browser,
        history_path=tmp_path / "history.json",
        report_dir=tmp_path,
        logger=logging.getLogger("pause-test"),
    )
    thread = threading.Thread(
        target=lambda: runner.run(
            [job],
            dry_run=True,
            stop_event=threading.Event(),
            pause_event=pause_event,
            progress=lambda _current, _total: None,
        )
    )
    thread.start()
    time.sleep(0.1)

    assert browser.published == 0
    scheduled_before_resume = job.scheduled_at
    pause_event.clear()
    thread.join(timeout=2)

    assert not thread.is_alive()
    assert browser.published == 1
    assert job.scheduled_at == scheduled_before_resume


def test_sheet_write_failure_does_not_block_v2r_publish(tmp_path: Path) -> None:
    path = tmp_path / "brand.csv"
    path.write_text(
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        '"키워드","제목 : 제목\n본문 : 본문",고요한아침,,질문형,,,실명,,가입인사\n',
        encoding="utf-8-sig",
    )
    job = load_brand_immediate_jobs(path, brand="팥순이")[0]

    class SheetFailureBrowser:
        published = 0

        def ensure_v2r_login(self, _email, _password):
            return None

        def prepare_immediate_jobs(self, jobs):
            jobs[0].cafe_id = 14567700
            jobs[0].menu_id = 34
            jobs[0].account = "writer"
            jobs[0].canonical_cafe_name = "고요한 아침"
            jobs[0].canonical_board_name = "가입인사"

        def consume_failed_immediate_urls(self):
            return set()

        def update_sheet_cell(self, *_args):
            raise RuntimeError("sheet unavailable")

        def publish_immediate(self, _job, _dry_run):
            self.published += 1
            return "https://v2r.example/source"

    browser = SheetFailureBrowser()
    runner = ImmediateRunner(
        browser=browser,
        history_path=tmp_path / "history.json",
        report_dir=tmp_path,
        logger=logging.getLogger("sheet-failure-test"),
    )
    result, _report = runner.run(
        [job],
        dry_run=False,
        stop_event=threading.Event(),
        progress=lambda _current, _total: None,
        source_sheet_url="https://docs.google.com/spreadsheets/d/example/edit?gid=0",
    )

    assert browser.published == 1
    assert result.jobs[0].status == JobStatus.RESERVED
    assert "D열 작성계정 저장 실패" in result.jobs[0].message
    assert "F열 완료 링크 저장 실패" in result.jobs[0].message


def test_previous_account_test_success_restores_sheet_without_republish(
    tmp_path: Path,
) -> None:
    path = tmp_path / "account-tests.csv"
    path.write_text(
        "번호,ID,작업 구분,연동,가아사 조건,실/비실,테스트 선택,테스트 결과,테스트 링크,테스트 일시\n"
        "12,test-id,,,,,TRUE,,,\n",
        encoding="utf-8-sig",
    )
    old_job = load_account_test_jobs(path)[0]
    old_job.cafe = "태극마케팅센터"
    old_job.post_url = "https://v2r.example/old-success"
    history_path = tmp_path / "history.json"
    HistoryStore(history_path).record(old_job)
    job = load_account_test_jobs(path)[0]

    class RecoveryBrowser:
        published = 0
        sheet_updates = []

        def ensure_v2r_login(self, _email, _password):
            return None

        def prepare_immediate_jobs(self, jobs):
            jobs[0].cafe = "태극마케팅센터"
            jobs[0].cafe_id = 31670254
            jobs[0].menu_id = 1
            jobs[0].canonical_cafe_name = "태극마케팅센터"
            jobs[0].canonical_board_name = "자유게시판"

        def consume_failed_immediate_urls(self):
            return set()

        def publish_immediate(self, _job, _dry_run):
            self.published += 1
            return "https://v2r.example/new"

        def update_sheet_cell(
            self,
            _url,
            column,
            row,
            value,
            **_kwargs,
        ):
            self.sheet_updates.append((column, row, value))

    browser = RecoveryBrowser()
    runner = ImmediateRunner(
        browser=browser,
        history_path=history_path,
        report_dir=tmp_path,
        logger=logging.getLogger("recovery-test"),
    )
    result, _report = runner.run(
        [job],
        dry_run=False,
        stop_event=threading.Event(),
        progress=lambda _current, _total: None,
        source_sheet_url="https://docs.google.com/spreadsheets/d/example/edit?gid=0",
    )

    assert browser.published == 0
    assert result.jobs[0].status == JobStatus.SUCCESS
    assert result.jobs[0].post_url == old_job.post_url
    assert ("I", 2, old_job.post_url) in browser.sheet_updates
    assert {column for column, _row, _value in browser.sheet_updates} == {
        "H",
        "I",
        "J",
    }


def test_failed_account_test_writes_only_result_column(tmp_path: Path) -> None:
    path = tmp_path / "account-tests.csv"
    path.write_text(
        "번호,ID,작업 구분,연동,가아사 조건,실/비실,테스트 선택,테스트 결과,테스트 링크,테스트 일시\n"
        "12,missing-id,,,,,TRUE,,,\n",
        encoding="utf-8-sig",
    )
    job = load_account_test_jobs(path)[0]

    class FailureResultBrowser:
        sheet_updates = []

        def ensure_v2r_login(self, _email, _password):
            return None

        def prepare_immediate_jobs(self, jobs):
            jobs[0].status = JobStatus.FAILED
            jobs[0].message = "V2R 미등록 계정"

        def consume_failed_immediate_urls(self):
            return set()

        def update_sheet_cell(
            self,
            _url,
            column,
            row,
            value,
            **_kwargs,
        ):
            self.sheet_updates.append((column, row, value))

    browser = FailureResultBrowser()
    runner = ImmediateRunner(
        browser=browser,
        history_path=tmp_path / "history.json",
        report_dir=tmp_path,
        logger=logging.getLogger("failure-result-test"),
    )
    runner.run(
        [job],
        dry_run=False,
        stop_event=threading.Event(),
        progress=lambda _current, _total: None,
        source_sheet_url="https://docs.google.com/spreadsheets/d/example/edit?gid=0",
    )

    assert browser.sheet_updates == [("H", 2, "V2R 미등록 계정")]
