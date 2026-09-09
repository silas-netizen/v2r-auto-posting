from pathlib import Path
import random

import pytest
from openpyxl import Workbook, load_workbook

from v2r_auto.content import CommentNode, ParsedArticle, parse_article
from v2r_auto.gatling_accounts import (
    comment_auto_labels,
    normalize_auto_id_count,
    recognize_proxy_workbook,
)
from v2r_auto.images import ResolvedImage
from v2r_auto.gatling_paste import (
    AFFILIATE_BOARD_LINKS,
    SELF_OWNED_BOARD_LINKS,
    SELF_OWNED_CAFE_IDS,
    AFFILIATE_EXACT_BOARDS,
    EXPORT_BOTH,
    EXPORT_EXCEL,
    EXPORT_TXT,
    MASTER_HEADER_ROW,
    MASTER_HEADERS,
    TXT_COMMENTS_123,
    TXT_COMMENTS_45,
    TXT_REPLIES,
    TXT_TITLE_BODY,
    TYPE_COMMENT,
    TYPE_EDIT_POST,
    TYPE_NEW_POST,
    TYPE_REPLY,
    GatlingBrandJob,
    GatlingPasteError,
    assign_daily_posts,
    build_and_write_master,
    build_gatling_master,
    build_master_rows,
    collect_resolved_images,
    completion_txt_stem,
    create_master_template,
    exact_board_name,
    export_gatling_output,
    gatling_image_folder,
    gatling_txt_folder,
    load_gatling_brand_jobs,
    normalize_gatling_source,
    parse_completion_cafe_article,
    parse_export_mode,
    parse_gatling_article,
    paste_manuscripts_into_gatling,
    recognize_gatling_workbook,
    replace_image_tokens,
    reply_target_value,
    require_writable_gatling,
    safe_txt_keyword,
    self_owned_board_link,
    split_manuscript_txt_parts,
    write_gatling_txt_files,
)
from v2r_auto.models import DailyPost


QUESTION_SOURCE = """제목 :
실제 원고 제목

본문 :
실제 원고 본문

댓글1:
첫 댓글
대댓글1:
첫 답글
댓글2:
둘째 댓글
대댓글2:
둘째 답글
대대댓글2:
깊은 답글
대대대댓글2:
더 깊은 답글
"""


def write_brand_csv(tmp_path: Path, content: str) -> Path:
    path = tmp_path / "brand.csv"
    path.write_text(content, encoding="utf-8-sig")
    return path


def write_gatling_xlsx(tmp_path: Path, name: str = "gatling.xlsx") -> Path:
    path = tmp_path / name
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "마스터"
    for column, header in enumerate(MASTER_HEADERS, start=1):
        sheet.cell(MASTER_HEADER_ROW, column, header)
    sheet.cell(7, 1, 800)
    sheet.cell(7, 2, "딜레이")
    sheet.cell(8, 2, TYPE_NEW_POST)
    sheet.cell(8, 3, "이미 있는 글")
    sheet.cell(8, 4, "이미 있는 본문")
    workbook.save(path)
    return path


def write_daily_csv(tmp_path: Path) -> Path:
    path = tmp_path / "daily.csv"
    path.write_text(
        (
            "번호,제목,내용,카페\n"
            '1,분류,"제목 : 씨씨앙 일상\n본문 : 씨씨앙 일상 본문",씨씨앙\n'
            '2,분류,"제목 : 양평맘 일상\n본문 : 양평맘 일상 본문",양평맘\n'
        ),
        encoding="utf-8-sig",
    )
    return path


def make_job(
    *,
    cafe: str,
    board: str,
    cafe_article_url: str = "",
    daily_title: str = "일상 제목",
    daily_body: str = "일상 본문",
    with_daily: bool = True,
) -> GatlingBrandJob:
    job = GatlingBrandJob(
        row_number=2,
        keyword="단식원 가격",
        article=parse_article("단식원 가격", QUESTION_SOURCE),
        cafe=cafe,
        board=board,
        article_type="질문형",
        prefix="자유",
        cafe_article_url=cafe_article_url,
    )
    if with_daily:
        job.daily_post = DailyPost(
            row_number=2,
            cafe=cafe,
            title=daily_title,
            body=daily_body,
        )
    return job


def test_exact_board_name_ignores_spaces_and_keeps_cafe_spelling() -> None:
    assert exact_board_name("자유수다방", "씨씨앙") == "자유 수다방"
    assert exact_board_name("이모저모이야기", "양평맘") == "이모저모 이야기💕"
    assert exact_board_name("", "씨씨앙") == AFFILIATE_EXACT_BOARDS["씨씨앙"]


def test_exact_board_name_rejects_wrong_affiliate_board() -> None:
    with pytest.raises(GatlingPasteError, match="정확한 이름"):
        exact_board_name("다른게시판", "씨씨앙")


def test_self_owned_board_uses_known_exact_name() -> None:
    assert (
        exact_board_name("웨딩홀탐방기", "마이 웨딩 드림") == "웨딩홀탑방기"
    )
    assert (
        exact_board_name("가입인사", "고요한아침", extra_exact_names=("가입 인사",))
        == "가입 인사"
    )


def test_reply_targets_match_manuscript_labels() -> None:
    article = parse_article("단식원 가격", QUESTION_SOURCE)
    first_reply = article.comments[0].children[0]
    second_reply = article.comments[1].children[0]
    deep = second_reply.children[0]
    deeper = deep.children[0]

    assert first_reply.label == "대댓글1"
    assert reply_target_value(first_reply) == 1
    assert reply_target_value(second_reply) == 2
    assert deep.label == "대대댓글2"
    assert reply_target_value(deep) == 2.1
    assert deeper.label == "대대대댓글2"
    assert reply_target_value(deeper) == 2.2


def test_affiliate_daily_goes_to_new_post_and_manuscript_goes_to_edit() -> None:
    rows = build_master_rows(make_job(cafe="씨씨앙", board="자유수다방"))

    assert rows[0].type == TYPE_NEW_POST
    assert rows[0].title == "일상 제목"
    assert rows[0].body == "일상 본문"
    assert rows[0].board_name == "자유 수다방"
    assert rows[0].hashtag == ""
    assert rows[0].prefix == "자유"
    assert rows[0].link == AFFILIATE_BOARD_LINKS["씨씨앙"]
    assert rows[0].image_location == ""

    assert rows[1].type == TYPE_EDIT_POST
    assert rows[1].title == "실제 원고 제목"
    assert rows[1].body == "실제 원고 본문"
    assert rows[1].hashtag == "단식원 가격"
    assert rows[1].board_name == "자유 수다방"
    assert rows[1].link is None


def test_edit_and_comments_use_cafe_url_in_the_right_columns() -> None:
    url = "https://cafe.naver.com/cantsb/3398059"
    rows = build_master_rows(
        make_job(cafe="양평맘", board="이모저모이야기", cafe_article_url=url)
    )

    assert rows[0].type == TYPE_NEW_POST
    assert rows[0].link == AFFILIATE_BOARD_LINKS["양평맘"]
    assert rows[1].type == TYPE_EDIT_POST
    assert rows[1].link == url
    assert rows[1].board_name == "이모저모 이야기💕"

    comments = [row for row in rows if row.type == TYPE_COMMENT]
    replies = [row for row in rows if row.type == TYPE_REPLY]
    assert [row.link for row in comments] == [url, url]
    assert [row.link for row in replies] == [1, 2, 2.1, 2.2]
    assert [row.result_link for row in replies] == [url, url, url, url]
    assert replies[2].body == "깊은 답글"
    assert replies[3].body == "더 깊은 답글"


def test_self_owned_manuscript_is_new_post_only() -> None:
    rows = build_master_rows(
        make_job(cafe="고요한아침", board="가입인사", with_daily=False)
    )

    types = [row.type for row in rows]
    assert types[0] == TYPE_NEW_POST
    assert TYPE_EDIT_POST not in types
    assert rows[0].title == "실제 원고 제목"
    assert rows[0].board_name == "가입인사"
    assert rows[0].hashtag == "단식원 가격"
    assert all(not row.hashtag for row in rows[1:])


def test_affiliate_new_post_gets_board_link_only() -> None:
    ssi = build_master_rows(make_job(cafe="씨씨앙", board="자유수다방"))
    yang = build_master_rows(make_job(cafe="양평맘", board="이모저모이야기"))
    self_owned = build_master_rows(
        make_job(cafe="고요한아침", board="가입인사", with_daily=False)
    )

    assert ssi[0].type == TYPE_NEW_POST
    assert ssi[0].link == AFFILIATE_BOARD_LINKS["씨씨앙"]
    assert all(row.link != AFFILIATE_BOARD_LINKS["씨씨앙"] for row in ssi[1:])
    assert yang[0].link == AFFILIATE_BOARD_LINKS["양평맘"]
    assert all(row.link != AFFILIATE_BOARD_LINKS["양평맘"] for row in yang[1:])
    assert self_owned[0].type == TYPE_NEW_POST
    assert self_owned[0].link == SELF_OWNED_BOARD_LINKS["고요한아침"]


def test_self_owned_new_post_gets_board_link() -> None:
    love = build_master_rows(
        make_job(cafe="러브인썸", board="뷰티미용", with_daily=False)
    )
    wedding = build_master_rows(
        make_job(cafe="마이 웨딩 드림", board="뷰티다이어트", with_daily=False)
    )
    love_home = build_master_rows(
        make_job(cafe="러브 인썸 (Love in Some)", board="뷰티&미용", with_daily=False)
    )

    assert love[0].type == TYPE_NEW_POST
    assert love[0].link == (
        f"https://cafe.naver.com/f-e/cafes/{SELF_OWNED_CAFE_IDS['러브인썸']}"
        "/menus/18?viewType=L"
    )
    assert wedding[0].link == SELF_OWNED_BOARD_LINKS["마이웨딩드림"]
    assert love_home[0].link == love[0].link
    assert all(
        row.link != love[0].link for row in love[1:] if row.type != TYPE_NEW_POST
    )


def test_self_owned_board_link_is_one_url_per_cafe() -> None:
    assert self_owned_board_link("러브인썸", "뷰티미용") == SELF_OWNED_BOARD_LINKS["러브인썸"]
    assert self_owned_board_link("마이웨딩드림", "뷰티다이어트") == (
        SELF_OWNED_BOARD_LINKS["마이웨딩드림"]
    )
    assert self_owned_board_link("고요한 아침") == SELF_OWNED_BOARD_LINKS["고요한아침"]
    assert self_owned_board_link("헬씨 트리") == SELF_OWNED_BOARD_LINKS["헬씨트리"]
    assert self_owned_board_link("송도포털") == SELF_OWNED_BOARD_LINKS["송도포털"]
    assert self_owned_board_link("글로시 마이") == SELF_OWNED_BOARD_LINKS["글로시마이"]
    assert self_owned_board_link("웨딩 노트") == SELF_OWNED_BOARD_LINKS["웨딩노트"]
    assert self_owned_board_link("쌍둥이맘 모여라") == SELF_OWNED_BOARD_LINKS["쌍둥이맘모여라"]
    assert SELF_OWNED_BOARD_LINKS["고요한아침"].endswith("/menus/29?viewType=L")
    assert SELF_OWNED_BOARD_LINKS["헬씨트리"].endswith("/menus/27?viewType=L")
    assert SELF_OWNED_BOARD_LINKS["송도포털"].endswith("/menus/19?viewType=L")
    assert SELF_OWNED_BOARD_LINKS["글로시마이"].endswith("/menus/50?viewType=L")
    assert SELF_OWNED_BOARD_LINKS["웨딩노트"].endswith("/menus/32?viewType=L")
    assert SELF_OWNED_BOARD_LINKS["쌍둥이맘모여라"].endswith("/menus/664?viewType=L")
    assert self_owned_board_link("없는카페") == ""


def test_new_self_owned_cafes_write_stored_board_link() -> None:
    rows = build_master_rows(
        make_job(cafe="헬씨 트리", board="자유로운 건강 수다방", with_daily=False)
    )
    assert rows[0].type == TYPE_NEW_POST
    assert rows[0].link == SELF_OWNED_BOARD_LINKS["헬씨트리"]
    songdo = build_master_rows(
        make_job(cafe="송도포털", board="친해지는 수다", with_daily=False)
    )
    assert songdo[0].link == SELF_OWNED_BOARD_LINKS["송도포털"]
    twin = build_master_rows(
        make_job(cafe="쌍둥이맘 모여라", board="ㄴ가족업체 자유게시판", with_daily=False)
    )
    assert twin[0].link == SELF_OWNED_BOARD_LINKS["쌍둥이맘모여라"]
    assert twin[0].board_name == "ㄴ가족업체 자유게시판"


def test_affiliate_without_daily_post_raises() -> None:
    job = make_job(cafe="씨씨앙", board="자유 수다방", with_daily=False)
    with pytest.raises(GatlingPasteError, match="일상 글"):
        build_master_rows(job)


def test_load_brand_sheet_reads_board_and_skips_completed(tmp_path: Path) -> None:
    path = write_brand_csv(
        tmp_path,
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        f'"단식원 가격","{QUESTION_SOURCE}",씨씨앙,writer,질문형,,,비실명,,자유수다방\n'
        f'"완료키워드","제목 : 완료\n본문 : 완료본문",씨씨앙,writer,질문형,https://v2r.example/x,,,Y,자유수다방\n',
    )

    jobs, skipped = load_gatling_brand_jobs(path, skip_completed=True)

    assert len(jobs) == 1
    assert jobs[0].board == "자유수다방"
    assert jobs[0].article.comments[1].children[0].children[0].label == "대대댓글2"
    assert any("완료 링크" in item for item in skipped)

    kept, kept_skipped = load_gatling_brand_jobs(path)
    assert len(kept) == 2
    assert kept_skipped == []


def test_build_assigns_daily_posts_and_writes_master_columns(tmp_path: Path) -> None:
    brand = write_brand_csv(
        tmp_path,
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        f'"단식원 가격","{QUESTION_SOURCE}",씨씨앙,writer,질문형,,,비실명,,자유수다방\n'
        f'"뱃살 보조제","제목 : 자사 제목\n본문 : 자사 본문",고요한아침,writer,후기형,,,,,가입인사\n',
    )
    daily = write_daily_csv(tmp_path)
    output = tmp_path / "master.xlsx"

    result = build_and_write_master(brand, output, daily, rng=__import__("random").Random(1))

    assert result.type_counts()[TYPE_NEW_POST] == 2
    assert result.type_counts()[TYPE_EDIT_POST] == 1
    affiliate_new = next(
        row
        for row in result.rows
        if row.type == TYPE_NEW_POST and row.board_name == "자유 수다방"
    )
    assert affiliate_new.title == "씨씨앙 일상"
    self_owned_new = next(
        row
        for row in result.rows
        if row.type == TYPE_NEW_POST and row.board_name == "가입인사"
    )
    assert self_owned_new.title == "자사 제목"

    workbook = load_workbook(output)
    sheet = workbook["마스터"]
    headers = [sheet.cell(MASTER_HEADER_ROW, column).value for column in range(1, 20)]
    assert headers == list(MASTER_HEADERS)

    reply_row = None
    for row_number in range(MASTER_HEADER_ROW + 1, sheet.max_row + 1):
        if sheet.cell(row_number, 2).value == TYPE_REPLY and sheet.cell(row_number, 1).value == 2.1:
            reply_row = row_number
            break
    assert reply_row is not None
    assert sheet.cell(reply_row, 4).value == "깊은 답글"
    assert sheet.cell(reply_row, 16).value is None
    assert sheet.cell(reply_row, 7).value is None

    edit_row = next(
        row_number
        for row_number in range(MASTER_HEADER_ROW + 1, sheet.max_row + 1)
        if sheet.cell(row_number, 2).value == TYPE_EDIT_POST
    )
    assert sheet.cell(edit_row, 3).value == "실제 원고 제목"
    assert sheet.cell(edit_row, 10).value == "자유 수다방"


def test_assign_daily_posts_does_not_reuse_same_cafe_template(tmp_path: Path) -> None:
    jobs = [
        make_job(cafe="양평맘", board="이모저모이야기", with_daily=False),
        make_job(cafe="양평맘", board="이모저모이야기", with_daily=False),
    ]
    jobs[1].row_number = 3
    posts = [
        DailyPost(row_number=2, cafe="양평맘", title="일상1", body="본문1"),
        DailyPost(row_number=3, cafe="양평맘", title="일상2", body="본문2"),
    ]

    assign_daily_posts(jobs, posts, rng=__import__("random").Random(3))

    assert {job.daily_post.title for job in jobs if job.daily_post} == {"일상1", "일상2"}


def test_affiliate_sheet_without_daily_file_raises(tmp_path: Path) -> None:
    brand = write_brand_csv(
        tmp_path,
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        f'"단식원 가격","{QUESTION_SOURCE}",씨씨앙,writer,질문형,,,,,자유수다방\n',
    )

    with pytest.raises(GatlingPasteError, match="일상 글 시트"):
        build_gatling_master(brand)


def test_manuscript_only_uses_sheet_title_body_and_all_comments() -> None:
    rows = build_master_rows(
        make_job(cafe="씨씨앙", board="자유수다방", with_daily=False),
        include_daily_new_post=False,
    )

    assert [row.type for row in rows] == [
        TYPE_EDIT_POST,
        TYPE_COMMENT,
        TYPE_COMMENT,
        TYPE_REPLY,
        TYPE_REPLY,
        TYPE_REPLY,
        TYPE_REPLY,
    ]
    assert rows[0].title == "실제 원고 제목"
    assert rows[0].body == "실제 원고 본문"
    assert [row.body for row in rows[1:]] == [
        "첫 댓글",
        "둘째 댓글",
        "첫 답글",
        "둘째 답글",
        "깊은 답글",
        "더 깊은 답글",
    ]


def test_xlsm_template_is_recognized_and_writable(tmp_path: Path) -> None:
    path = create_master_template(tmp_path / "V2R-Gatling-Master.xlsm")
    info = recognize_gatling_workbook(path)

    assert path.suffix == ".xlsm"
    assert info.recognized is True
    assert info.writable is True
    assert info.kind == "xlsm"
    assert info.next_row == 7


def test_recognize_gatling_by_master_headers(tmp_path: Path) -> None:
    path = write_gatling_xlsx(tmp_path)
    info = recognize_gatling_workbook(path)

    assert info.recognized is True
    assert info.writable is True
    assert info.kind == "xlsx"
    assert info.last_data_row == 8
    assert info.next_row == 9
    assert "9행" in info.message


def test_recognize_headers_even_if_shifted_or_not_on_row_six(tmp_path: Path) -> None:
    path = tmp_path / "shifted.xlsm"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "마스터"
    sheet.cell(5, 2, "링크")
    sheet.cell(5, 3, "타입")
    sheet.cell(5, 4, "제목")
    sheet.cell(5, 5, "내용")
    sheet.cell(6, 3, "새글")
    sheet.cell(6, 4, "기존 제목")
    workbook.save(path)

    info = recognize_gatling_workbook(path)

    assert info.recognized is True
    assert info.header_row == 5
    assert info.start_column == 2
    assert info.next_row == 7


def test_reject_excel_that_is_not_gatling(tmp_path: Path) -> None:
    path = tmp_path / "other.xlsx"
    workbook = Workbook()
    workbook.active.title = "Sheet1"
    workbook.active["A1"] = "이름"
    workbook.save(path)

    info = recognize_gatling_workbook(path)

    assert info.recognized is False
    assert info.writable is False
    assert "마스터" in info.message
    assert "시트" in info.message


def test_paste_appends_after_existing_master_rows(tmp_path: Path) -> None:
    brand = write_brand_csv(
        tmp_path,
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        f'"단식원 가격","{QUESTION_SOURCE}",고요한아침,writer,질문형,,,,,가입인사\n',
    )
    gatling = write_gatling_xlsx(tmp_path)

    result, start_row = paste_manuscripts_into_gatling(brand, gatling)

    assert start_row == 9
    assert result.rows[0].title == "실제 원고 제목"
    workbook = load_workbook(gatling)
    sheet = workbook["마스터"]
    assert sheet.cell(8, 3).value == "이미 있는 글"
    assert sheet.cell(9, 2).value == TYPE_NEW_POST
    assert sheet.cell(9, 3).value == "실제 원고 제목"
    assert sheet.cell(9, 4).value == "실제 원고 본문"
    assert sheet.cell(10, 2).value == TYPE_COMMENT
    assert sheet.cell(10, 4).value == "첫 댓글"
    assert sheet.cell(14, 1).value == 2.1
    assert sheet.cell(14, 4).value == "깊은 답글"
    assert sheet.cell(15, 1).value == 2.2
    assert sheet.cell(15, 4).value == "더 깊은 답글"


def write_user_like_gatling(tmp_path: Path) -> Path:
    """마스터 헤더가 B6이고, 앞 새글은 채워져 있으며 댓글·대댓글이 뒤에 있는 실제 파일 형태."""
    path = tmp_path / "복붙용.xlsm"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "마스터"
    for offset, header in enumerate(("링크", "타입", "제목", "내용", "크롬번호", "아이디", "비번", "해시태그")):
        sheet.cell(6, 2 + offset, header)
    sheet.cell(7, 2, 800)
    sheet.cell(7, 3, "딜레이")
    sheet.cell(8, 3, TYPE_EDIT_POST)
    sheet.cell(8, 4, "이미 있는 수정 제목")
    sheet.cell(8, 5, "이미 있는 수정 본문")
    for row_number in range(9, 13):
        sheet.cell(row_number, 3, TYPE_NEW_POST)
        sheet.cell(row_number, 4, f"이미 있는 제목 {row_number}")
        sheet.cell(row_number, 5, f"이미 있는 본문 {row_number}")
    for row_number in range(13, 17):
        sheet.cell(row_number, 3, TYPE_NEW_POST)
        sheet.cell(row_number, 9, f"준비된 해시 {row_number}")
    for row_number in range(17, 19):
        sheet.cell(row_number, 3, TYPE_COMMENT)
        sheet.cell(row_number, 5, f"이미 있는 댓글 {row_number}")
    for row_number in range(19, 22):
        sheet.cell(row_number, 3, TYPE_COMMENT)
        sheet.cell(row_number, 9, f"댓글 해시 {row_number}")
    sheet.cell(22, 2, 1)
    sheet.cell(22, 3, TYPE_REPLY)
    sheet.cell(22, 5, "이미 있는 대댓글")
    for row_number in range(23, 26):
        sheet.cell(row_number, 3, TYPE_REPLY)
    sheet.cell(27, 3, TYPE_COMMENT)
    sheet.cell(27, 5, "뒤에 있는 댓글")
    sheet.cell(28, 3, TYPE_REPLY)
    sheet.cell(28, 5, "뒤에 있는 대댓글")
    workbook.save(path)
    return path


def test_recognize_uses_first_empty_title_not_last_comment(tmp_path: Path) -> None:
    path = write_user_like_gatling(tmp_path)
    info = recognize_gatling_workbook(path)

    assert info.recognized is True
    assert info.header_row == 6
    assert info.start_column == 2
    assert info.next_row == 13
    assert "13행" in info.message


def test_paste_fills_empty_title_and_keeps_comment_blocks(tmp_path: Path) -> None:
    path = write_user_like_gatling(tmp_path)
    brand = write_brand_csv(
        tmp_path,
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        f'"단식원 가격","{QUESTION_SOURCE}",고요한아침,writer,질문형,,,,,가입인사\n',
    )

    result, start_row = paste_manuscripts_into_gatling(brand, path)

    assert start_row == 13
    workbook = load_workbook(path)
    sheet = workbook["마스터"]
    assert sheet.cell(12, 4).value == "이미 있는 제목 12"
    assert sheet.cell(13, 3).value == TYPE_NEW_POST
    assert sheet.cell(13, 4).value == "실제 원고 제목"
    assert sheet.cell(13, 5).value == "실제 원고 본문"
    assert sheet.cell(14, 3).value == TYPE_NEW_POST
    assert sheet.cell(14, 4) is None or sheet.cell(14, 4).value is None
    assert sheet.cell(17, 5).value == "이미 있는 댓글 17"
    assert sheet.cell(19, 3).value == TYPE_COMMENT
    assert sheet.cell(19, 5).value == "첫 댓글"
    assert sheet.cell(22, 5).value == "이미 있는 대댓글"
    assert sheet.cell(23, 3).value == TYPE_REPLY
    assert sheet.cell(23, 5).value == "첫 답글"
    assert sheet.cell(27, 5).value == "뒤에 있는 댓글"
    assert sheet.cell(28, 5).value == "뒤에 있는 대댓글"
    assert result.rows[0].title == "실제 원고 제목"


def test_filled_row_thirty_one_is_not_treated_as_empty(tmp_path: Path) -> None:
    path = tmp_path / "숨긴행처럼보임.xlsm"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "마스터"
    for offset, header in enumerate(("링크", "타입", "제목", "내용")):
        sheet.cell(6, 2 + offset, header)
    for row_number in range(7, 63):
        sheet.cell(row_number, 3, TYPE_NEW_POST)
        sheet.cell(row_number, 4, f"이미 있는 제목 {row_number}")
        sheet.cell(row_number, 5, f"이미 있는 본문 {row_number}")
    for row_number in range(63, 70):
        sheet.cell(row_number, 3, TYPE_NEW_POST)
        sheet.cell(row_number, 9, "다이어트 카페")
    for row_number in range(94, 100):
        sheet.cell(row_number, 3, TYPE_COMMENT)
        sheet.cell(row_number, 5, f"이미 있는 댓글 {row_number}")
    for row_number in range(149, 152):
        sheet.cell(row_number, 3, TYPE_COMMENT)
    workbook.save(path)

    info = recognize_gatling_workbook(path)
    assert info.next_row == 63
    assert sheet.cell(31, 4).value == "이미 있는 제목 31"

    brand = write_brand_csv(
        tmp_path,
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        f'"단식원 가격","{QUESTION_SOURCE}",고요한아침,writer,질문형,,,,,가입인사\n',
    )
    _, start_row = paste_manuscripts_into_gatling(brand, path)
    workbook = load_workbook(path)
    sheet = workbook["마스터"]
    assert start_row == 63
    assert sheet.cell(31, 4).value == "이미 있는 제목 31"
    assert sheet.cell(62, 4).value == "이미 있는 제목 62"
    assert sheet.cell(63, 4).value == "실제 원고 제목"
    assert sheet.cell(94, 5).value == "이미 있는 댓글 94"
    assert sheet.cell(149, 5).value == "첫 댓글"


def test_real_user_master_starts_at_row_sixty_three() -> None:
    source = Path("/tmp/gatling-user/user.xlsm")
    if not source.exists():
        pytest.skip("사용자 기관총 파일이 없습니다")
    path = Path("/tmp/gatling-user/user-copy.xlsm")
    path.write_bytes(source.read_bytes())

    info = recognize_gatling_workbook(path)
    assert info.next_row == 63
    assert "63행" in info.message


def test_xlsb_is_recognized_but_not_writable(tmp_path: Path) -> None:
    source = Path("/tmp/gatling/gatling.bin")
    if not source.exists():
        pytest.skip("기관총 원본 샘플이 없습니다")
    path = tmp_path / "gatling.xlsb"
    path.write_bytes(source.read_bytes())

    info = recognize_gatling_workbook(path)

    assert info.recognized is True
    assert info.writable is False
    assert info.kind == "xlsb"
    assert "마스터" in info.sheet_names
    assert info.headers[:4] == ["링크", "타입", "제목", "내용"]
    with pytest.raises(GatlingPasteError, match="xlsb"):
        require_writable_gatling(path)


def test_marked_labels_are_recognized_but_marks_are_not_pasted() -> None:
    article = parse_gatling_article(
        "스위치온 쉐이크",
        """#제목 :
"스위치온 쉐이크 3주 먹어본 진짜 후기"

**본문 :**
#첫 줄입니다
둘째 줄입니다

#댓글 :
"첫 댓글"
**대댓글 :**
답글입니다
""",
    )
    assert article.title == "스위치온 쉐이크 3주 먹어본 진짜 후기"
    assert article.body == "첫 줄입니다\n둘째 줄입니다"
    assert article.comments[0].text == "첫 댓글"
    assert article.comments[0].children[0].text == "답글입니다"
    assert "#" not in article.title
    assert '"' not in article.body
    assert "*" not in article.comments[0].text


def test_quoted_and_fullwidth_labels_still_parse() -> None:
    article = parse_gatling_article(
        "키워드",
        '"제목 ： 멋진 제목\n본문 ： 멋진 본문\n댓글 : 댓글 내용"',
    )
    assert article.title == "멋진 제목"
    assert article.body == "멋진 본문"
    assert article.comments[0].text == "댓글 내용"


def test_plain_manuscript_stays_unchanged() -> None:
    article = parse_gatling_article("키워드", QUESTION_SOURCE)
    assert article.title == "실제 원고 제목"
    assert article.body == "실제 원고 본문"
    assert [node.text for node in article.comments] == ["첫 댓글", "둘째 댓글"]


def test_body_sentence_with_제목_is_not_a_label() -> None:
    source = normalize_gatling_source("제목 : 제목\n본문 :\n그래서 제목 : 이렇게 지었어요")
    article = parse_gatling_article("키워드", source)
    assert article.body == "그래서 제목 : 이렇게 지었어요"


def test_load_jobs_accepts_marked_title_lines(tmp_path: Path) -> None:
    path = write_brand_csv(
        tmp_path,
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        '"키워드","#제목 :\n실제 제목\n#본문 :\n실제 본문",고요한아침,writer,질문형,,,,,가입인사\n',
    )
    jobs, skipped = load_gatling_brand_jobs(path)
    assert skipped == []
    assert jobs[0].article.title == "실제 제목"
    assert jobs[0].article.body == "실제 본문"


def test_image_tokens_become_gatling_placeholder() -> None:
    assert replace_image_tokens("사진 {키워드} 끝") == "사진 {이미지} 끝"
    assert replace_image_tokens("{A열 키워드}\n본문") == "{이미지}\n본문"
    assert replace_image_tokens("{A열키워드}") == "{이미지}"
    assert replace_image_tokens("전 {B/A} 후") == "전 {이미지} 후"
    assert replace_image_tokens("이미 {이미지} 있음") == "이미 {이미지} 있음"


def test_image_tokens_are_rewritten_in_title_body_and_comments() -> None:
    source = (
        "제목 :\n제목 {키워드}\n\n본문 :\n본문 {A열 키워드}\n\n"
        "댓글1:\n댓글 {키워드}\n대댓글1:\n답글 {A열 키워드}\n"
    )
    job = GatlingBrandJob(
        row_number=2,
        keyword="엉덩이 종기",
        article=parse_article("엉덩이 종기", source),
        cafe="고요한아침",
        board="가입인사",
        article_type="질문형",
    )

    rows = build_master_rows(job, include_daily_new_post=False)

    assert rows[0].title == "제목 {이미지}"
    assert rows[0].body == "본문 {이미지}"
    assert rows[1].body == "댓글 {이미지}"
    assert rows[2].body == "답글 {이미지}"


def test_revision_order_is_daily_edit_all_comments_then_replies() -> None:
    rows = build_master_rows(make_job(cafe="씨씨앙", board="자유수다방"))
    types = [row.type for row in rows]

    assert types[:2] == [TYPE_NEW_POST, TYPE_EDIT_POST]
    assert TYPE_EDIT_POST not in types[2:]
    assert types[2:] == [TYPE_COMMENT, TYPE_COMMENT] + [TYPE_REPLY] * 4
    assert rows[0].title == "일상 제목"
    assert rows[1].title == "실제 원고 제목"


def test_replies_are_grouped_by_depth_like_the_template() -> None:
    comments = []
    for index in range(1, 4):
        child = CommentNode(label=f"대댓글{index}", text=f"답글{index}", depth=1, index=index)
        comments.append(
            CommentNode(
                label=f"댓글{index}",
                text=f"댓글{index} 내용",
                depth=0,
                index=index,
                children=[child],
            )
        )
    comments[1].children[0].children.append(
        CommentNode(label="대대댓글2", text="깊은 답글", depth=2, index=2)
    )
    job = GatlingBrandJob(
        row_number=2,
        keyword="키워드",
        article=ParsedArticle(
            title="원고 제목",
            body="원고 내용",
            keyword="키워드",
            tag="",
            comments=comments,
        ),
        cafe="씨씨앙",
        board="자유수다방",
        article_type="질문형",
        daily_post=DailyPost(row_number=2, cafe="씨씨앙", title="일상 글 제목", body="일상 글 내용"),
    )

    rows = build_master_rows(job)
    replies = [row for row in rows if row.type == TYPE_REPLY]

    assert [row.type for row in rows[:5]] == [
        TYPE_NEW_POST,
        TYPE_EDIT_POST,
        TYPE_COMMENT,
        TYPE_COMMENT,
        TYPE_COMMENT,
    ]
    assert [row.link for row in replies] == [1, 2, 3, 2.1]
    assert [row.body for row in replies] == ["답글1", "답글2", "답글3", "깊은 답글"]


def test_completed_f_column_rows_are_kept_for_gatling(tmp_path: Path) -> None:
    brand = write_brand_csv(
        tmp_path,
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        f'"엉덩이 종기","{QUESTION_SOURCE}",양평맘,writer,질문형,'
        "https://v2r.daboja.im/nc/articleDetail/1,,,,이모저모이야기\n",
    )

    jobs, skipped = load_gatling_brand_jobs(brand)
    result = build_gatling_master(brand, manuscript_only=True)

    assert skipped == []
    assert len(jobs) == 1
    assert jobs[0].keyword == "엉덩이 종기"
    assert jobs[0].completion_url.startswith("https://v2r.daboja.im")
    assert result.jobs[0].keyword == "엉덩이 종기"


def test_skip_completed_explains_why_nothing_is_left(tmp_path: Path) -> None:
    brand = write_brand_csv(
        tmp_path,
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        f'"엉덩이 종기","{QUESTION_SOURCE}",양평맘,writer,질문형,'
        "https://v2r.daboja.im/nc/articleDetail/1,,,,이모저모이야기\n",
    )

    with pytest.raises(GatlingPasteError, match="F열 완료 링크"):
        load_gatling_brand_jobs(brand, skip_completed=True)


def test_collected_images_use_pipe_separated_paths(tmp_path: Path) -> None:
    source_dir = tmp_path / "src"
    source_dir.mkdir()
    first = source_dir / "1.jpg"
    second = source_dir / "31.jpg"
    third = source_dir / "51.jpg"
    first.write_bytes(b"one")
    second.write_bytes(b"two")
    third.write_bytes(b"three")
    dest = tmp_path / "복불용_images"
    resolved = [
        ResolvedImage(0, "키워드", "a", "1.jpg", first),
        ResolvedImage(1, "키워드", "b", "31.jpg", second),
        ResolvedImage(2, "B/A", "c", "51.jpg", third),
    ]

    joined = collect_resolved_images(resolved, dest)

    assert dest.joinpath("1.jpg").read_bytes() == b"one"
    assert dest.joinpath("31.jpg").read_bytes() == b"two"
    assert dest.joinpath("51.jpg").read_bytes() == b"three"
    assert joined == "|".join(
        [str(dest / "1.jpg"), str(dest / "31.jpg"), str(dest / "51.jpg")]
    )


def test_image_folder_sits_next_to_the_excel(tmp_path: Path) -> None:
    path = tmp_path / "기관총 카페봇 복불용.xlsm"
    assert gatling_image_folder(path) == tmp_path / "기관총 카페봇 복불용_images"


def test_manuscript_row_gets_hashtag_and_image_location() -> None:
    job = make_job(cafe="씨씨앙", board="자유수다방")
    rows = build_master_rows(job, image_location=r"G:\image\1.jpg|G:\image\31.jpg")

    assert rows[0].hashtag == ""
    assert rows[0].image_location == ""
    assert rows[1].hashtag == "단식원 가격"
    assert rows[1].image_location == r"G:\image\1.jpg|G:\image\31.jpg"
    assert rows[1].cells()[13] == r"G:\image\1.jpg|G:\image\31.jpg"
    assert all(not row.hashtag and not row.image_location for row in rows[2:])


def test_brand_is_read_from_sheet_filename(tmp_path: Path) -> None:
    path = tmp_path / "카페 원고 작성 시트 (뉴더미스).csv"
    path.write_text(
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        f'"엉덩이 종기","{QUESTION_SOURCE}",양평맘,writer,질문형,,,,,이모저모이야기\n',
        encoding="utf-8-sig",
    )

    jobs, _ = load_gatling_brand_jobs(path)
    assert jobs[0].brand == "뉴더미스"


def test_same_title_and_body_are_not_pasted_twice(tmp_path: Path) -> None:
    brand = write_brand_csv(
        tmp_path,
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        f'"단식원 가격","{QUESTION_SOURCE}",고요한아침,writer,질문형,,,,,가입인사\n'
        f'"다른 키워드","{QUESTION_SOURCE}",고요한아침,writer,질문형,,,,,가입인사\n',
    )
    gatling = write_gatling_xlsx(tmp_path)

    first, _ = paste_manuscripts_into_gatling(brand, gatling)
    assert len(first.jobs) == 1
    assert any("제목·본문이 이미 있어" in item for item in first.skipped)

    with pytest.raises(GatlingPasteError, match="제목·본문이 같은 원고"):
        paste_manuscripts_into_gatling(brand, gatling)


def test_same_title_with_different_body_is_kept(tmp_path: Path) -> None:
    other = (
        "제목 :\n실제 원고 제목\n\n본문 :\n다른 본문\n\n댓글1:\n첫 댓글\n"
    )
    brand = write_brand_csv(
        tmp_path,
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        f'"단식원 가격","{QUESTION_SOURCE}",고요한아침,writer,질문형,,,,,가입인사\n'
        f'"단식원 가격","{other}",고요한아침,writer,질문형,,,,,가입인사\n',
    )

    result = build_gatling_master(brand, manuscript_only=True)
    assert len(result.jobs) == 2
    assert result.skipped == []


def _simple_source(title: str, body: str, comment: str = "댓글본문") -> str:
    return f"제목 :\n{title}\n\n본문 :\n{body}\n\n댓글1:\n{comment}\n대댓글1:\n답글본문\n"


def test_empty_excel_type_is_filled_and_all_sheet_jobs_are_pasted(tmp_path: Path) -> None:
    path = tmp_path / "empty-type.xlsm"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "마스터"
    for offset, header in enumerate(("링크", "타입", "제목", "내용")):
        sheet.cell(6, 2 + offset, header)
    sheet.cell(7, 3, TYPE_NEW_POST)
    sheet.cell(7, 4, "이미 있는 제목")
    sheet.cell(7, 5, "이미 있는 본문")
    sheet.cell(8, 3, TYPE_NEW_POST)
    sheet.cell(9, 2, 1)
    for row_number in range(10, 16):
        sheet.cell(row_number, 3, None)
    workbook.save(path)

    first = _simple_source("첫째 제목", "첫째 본문", "첫째 댓글")
    second = _simple_source("둘째 제목", "둘째 본문", "둘째 댓글")
    third = _simple_source("셋째 제목", "셋째 본문", "셋째 댓글")
    brand = write_brand_csv(
        tmp_path,
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        f'"키워드1","{first}",고요한아침,writer,,,,,,가입인사\n'
        f'"키워드2","{second}",고요한아침,writer,질문형,,,,,가입인사\n'
        f'"키워드3","{third}",고요한아침,writer,후기형,,,,,가입인사\n',
    )

    result, start_row = paste_manuscripts_into_gatling(brand, path)
    workbook = load_workbook(path)
    sheet = workbook["마스터"]

    assert len(result.jobs) == 3
    assert start_row == 8
    assert sheet.cell(7, 4).value == "이미 있는 제목"
    assert sheet.cell(8, 3).value == TYPE_NEW_POST
    assert sheet.cell(8, 4).value == "첫째 제목"
    assert sheet.cell(8, 5).value == "첫째 본문"
    assert sheet.cell(9, 2).value is None
    assert sheet.cell(9, 3).value == TYPE_COMMENT
    assert sheet.cell(9, 5).value == "첫째 댓글"
    assert sheet.cell(10, 3).value == TYPE_REPLY
    assert sheet.cell(10, 5).value == "답글본문"
    assert sheet.cell(11, 3).value == TYPE_NEW_POST
    assert sheet.cell(11, 4).value == "둘째 제목"
    assert sheet.cell(12, 3).value == TYPE_COMMENT
    assert sheet.cell(12, 5).value == "둘째 댓글"
    assert sheet.cell(13, 3).value == TYPE_REPLY
    assert sheet.cell(14, 3).value == TYPE_NEW_POST
    assert sheet.cell(14, 4).value == "셋째 제목"
    assert sheet.cell(15, 3).value == TYPE_COMMENT
    assert sheet.cell(15, 5).value == "셋째 댓글"


def test_blank_type_after_comment_block_is_used_for_next_article(tmp_path: Path) -> None:
    path = tmp_path / "blank-after-comments.xlsm"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "마스터"
    for offset, header in enumerate(("링크", "타입", "제목", "내용")):
        sheet.cell(6, 2 + offset, header)
    sheet.cell(7, 3, TYPE_NEW_POST)
    sheet.cell(7, 4, "이미 있는 제목")
    sheet.cell(7, 5, "이미 있는 본문")
    sheet.cell(8, 3, TYPE_COMMENT)
    sheet.cell(8, 5, "이미 있는 댓글")
    sheet.cell(9, 3, None)
    sheet.cell(10, 3, None)
    sheet.cell(11, 3, None)
    workbook.save(path)

    info = recognize_gatling_workbook(path)
    assert info.next_row == 9

    brand = write_brand_csv(
        tmp_path,
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        f'"키워드1","{_simple_source("새 제목", "새 본문")}",고요한아침,writer,질문형,,,,,가입인사\n',
    )
    _, start_row = paste_manuscripts_into_gatling(brand, path)
    workbook = load_workbook(path)
    sheet = workbook["마스터"]
    assert start_row == 9
    assert sheet.cell(8, 5).value == "이미 있는 댓글"
    assert sheet.cell(9, 3).value == TYPE_NEW_POST
    assert sheet.cell(9, 4).value == "새 제목"
    assert sheet.cell(10, 3).value == TYPE_COMMENT
    assert sheet.cell(11, 3).value == TYPE_REPLY


def test_extra_manuscripts_still_append_with_types(tmp_path: Path) -> None:
    path = write_gatling_xlsx(tmp_path)
    first = _simple_source("첫째 제목", "첫째 본문")
    second = _simple_source("둘째 제목", "둘째 본문")
    brand = write_brand_csv(
        tmp_path,
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        f'"키워드1","{first}",고요한아침,writer,질문형,,,,,가입인사\n'
        f'"키워드2","{second}",고요한아침,writer,질문형,,,,,가입인사\n',
    )

    result, start_row = paste_manuscripts_into_gatling(brand, path)
    workbook = load_workbook(path)
    sheet = workbook["마스터"]
    titles = [
        sheet.cell(row_number, 3).value
        for row_number in range(7, sheet.max_row + 1)
        if sheet.cell(row_number, 2).value == TYPE_NEW_POST
    ]

    assert len(result.jobs) == 2
    assert start_row == 9
    assert "이미 있는 글" in titles
    assert "첫째 제목" in titles
    assert "둘째 제목" in titles
    assert sheet.cell(9, 2).value == TYPE_NEW_POST
    assert sheet.cell(10, 2).value == TYPE_COMMENT
    assert sheet.cell(11, 2).value == TYPE_REPLY


def test_load_keeps_rows_without_article_type(tmp_path: Path) -> None:
    brand = write_brand_csv(
        tmp_path,
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        f'"단식원 가격","{_simple_source("제목만", "본문만")}",고요한아침,writer,,,,,,가입인사\n',
    )

    jobs, skipped = load_gatling_brand_jobs(brand)
    assert skipped == []
    assert len(jobs) == 1
    assert jobs[0].article.title == "제목만"


def write_proxy_xlsx(
    tmp_path: Path,
    comment_count: int = 6,
    self_comment_count: int = 6,
    affiliate_author_count: int = 1,
    self_author_count: int = 1,
) -> Path:
    path = tmp_path / "proxy.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["아이피:포트", "크롬번", "아이디", "비번", "카테고리"])
    sheet.append(["10.0.0.1:3030", 1, "writer", "pw-writer", "제휴"])
    for offset in range(1, affiliate_author_count):
        name = f"writer{offset + 1}"
        sheet.append(
            [f"10.0.1.{offset}:3030", 100 + offset, name, f"pw-{name}", "제휴"]
        )
    sheet.append(["10.0.0.2:3030", 10, "selfwriter", "pw-self", "자사"])
    for offset in range(1, self_author_count):
        name = f"selfwriter{offset + 1}"
        sheet.append(
            [f"10.0.2.{offset}:3030", 200 + offset, name, f"pw-{name}", "자사"]
        )
    for offset in range(comment_count):
        name = f"c{offset + 1}"
        sheet.append(
            [f"10.0.0.{offset + 11}:3030", 11 + offset, name, f"pw-{name}", "제휴 댓"]
        )
    for offset in range(self_comment_count):
        name = f"s{offset + 1}"
        sheet.append(
            [f"10.0.0.{offset + 31}:3030", 31 + offset, name, f"pw-{name}", "자사 댓"]
        )
    workbook.save(path)
    return path


def test_proxy_csv_is_recognized(tmp_path: Path) -> None:
    path = tmp_path / "proxy.csv"
    path.write_text(
        "아이피:포트,크롬번,아이디,비번,카테고리\n"
        "10.0.0.1:3030,1,writer,pw-writer,제휴\n"
        "10.0.0.11:3030,11,c1,pw-c1,제휴 댓\n",
        encoding="utf-8-sig",
    )
    info = recognize_proxy_workbook(path)
    assert info.recognized is True
    assert info.find_author("writer", "씨씨앙").chrome_number == 1


def test_proxy_excel_is_recognized(tmp_path: Path) -> None:
    path = write_proxy_xlsx(tmp_path)
    info = recognize_proxy_workbook(path)

    assert info.recognized is True
    assert "제휴 작성 1개" in info.message
    assert "제휴 댓글 6개" in info.message
    assert "자사 작성 1개" in info.message
    assert "자사 댓글 6개" in info.message
    assert info.find_author("writer", "씨씨앙").chrome_number == 1
    assert info.find_author("selfwriter", "고요한아침").chrome_number == 10


def test_question_type_shares_comment_two_and_puts_sixth_on_2_2(tmp_path: Path) -> None:
    job = make_job(cafe="씨씨앙", board="자유수다방")
    job.account = "writer"
    job.article_type = "질문형"
    book = recognize_proxy_workbook(write_proxy_xlsx(tmp_path))
    rng = random.Random(1)
    chosen = rng.sample(book.comment_accounts(), 6)
    rows = build_master_rows(job, proxy_book=book, rng=random.Random(1))

    comments = [row for row in rows if row.type == TYPE_COMMENT]
    replies = [row for row in rows if row.type == TYPE_REPLY]
    assert rows[0].account == "writer"
    assert rows[0].chrome_number == 1
    assert rows[1].account == "writer"
    assert comments[0].account == chosen[0].account
    assert comments[1].account == chosen[1].account
    assert replies[0].account == "writer"
    assert replies[1].account == "writer"
    assert replies[2].account == chosen[1].account
    assert replies[3].account == chosen[2].account
    assert replies[2].link == 2.1
    assert replies[3].link == 2.2


def test_review_type_puts_sixth_on_2_1_and_author_on_2_2(tmp_path: Path) -> None:
    job = make_job(cafe="양평맘", board="이모저모이야기")
    job.account = "writer"
    job.article_type = "후기형"
    book = recognize_proxy_workbook(write_proxy_xlsx(tmp_path))
    rng = random.Random(2)
    chosen = rng.sample(book.comment_accounts(), 6)
    rows = build_master_rows(job, proxy_book=book, rng=random.Random(2))

    comments = [row for row in rows if row.type == TYPE_COMMENT]
    replies = [row for row in rows if row.type == TYPE_REPLY]
    assert comments[1].account == chosen[1].account
    assert replies[2].account == chosen[2].account
    assert replies[3].account == "writer"
    assert replies[0].account == "writer"
    assert replies[1].account == "writer"


def test_self_owned_question_type_uses_self_comment_ids(tmp_path: Path) -> None:
    job = make_job(cafe="고요한아침", board="가입인사", with_daily=False)
    job.account = "selfwriter"
    job.article_type = "질문형"
    book = recognize_proxy_workbook(write_proxy_xlsx(tmp_path))
    rng = random.Random(1)
    chosen = rng.sample(book.self_comment_accounts(), 6)
    rows = build_master_rows(job, proxy_book=book, rng=random.Random(1))

    comments = [row for row in rows if row.type == TYPE_COMMENT]
    replies = [row for row in rows if row.type == TYPE_REPLY]
    assert rows[0].account == "selfwriter"
    assert rows[0].chrome_number == 10
    assert comments[0].account == chosen[0].account
    assert comments[1].account == chosen[1].account
    assert replies[0].account == "selfwriter"
    assert replies[1].account == "selfwriter"
    assert replies[2].account == chosen[1].account
    assert replies[3].account == chosen[2].account
    assert {row.account for row in comments} <= {item.account for item in chosen}
    assert "c1" not in {row.account for row in comments}


def test_reply_accounts_follow_sheet_author_not_random_self(tmp_path: Path) -> None:
    job = make_job(cafe="고요한아침", board="가입인사", with_daily=False)
    job.account = "sheetwriter"
    job.article_type = "질문형"
    book = recognize_proxy_workbook(write_proxy_xlsx(tmp_path))
    rows = build_master_rows(job, proxy_book=book, rng=random.Random(1))
    replies = [row for row in rows if row.type == TYPE_REPLY]
    assert rows[0].account == "sheetwriter"
    assert replies[0].account == "sheetwriter"
    assert replies[1].account == "sheetwriter"
    assert "selfwriter" not in {row.account for row in rows if row.type != TYPE_COMMENT}


def test_self_owned_review_type_matches_affiliate_order(tmp_path: Path) -> None:
    job = make_job(cafe="고요한아침", board="가입인사", with_daily=False)
    job.account = "selfwriter"
    job.article_type = "후기형"
    book = recognize_proxy_workbook(write_proxy_xlsx(tmp_path))
    rng = random.Random(2)
    chosen = rng.sample(book.self_comment_accounts(), 6)
    rows = build_master_rows(job, proxy_book=book, rng=random.Random(2))

    comments = [row for row in rows if row.type == TYPE_COMMENT]
    replies = [row for row in rows if row.type == TYPE_REPLY]
    assert comments[1].account == chosen[1].account
    assert replies[2].account == chosen[2].account
    assert replies[3].account == "selfwriter"
    assert replies[0].account == "selfwriter"
    assert replies[1].account == "selfwriter"
    assert not {row.account for row in comments} & {f"c{i}" for i in range(1, 7)}


def test_missing_author_is_explained(tmp_path: Path) -> None:
    brand = write_brand_csv(
        tmp_path,
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        f'"단식원 가격","{_simple_source("제목", "본문")}",고요한아침,unknown,질문형,,,,,가입인사\n',
    )
    book = recognize_proxy_workbook(write_proxy_xlsx(tmp_path))
    result = build_gatling_master(brand, manuscript_only=True, proxy_book=book)
    assert any("프록시 엑셀에서 찾지 못했습니다" in item for item in result.skipped)


def test_too_few_self_comment_accounts_raises(tmp_path: Path) -> None:
    brand = write_brand_csv(
        tmp_path,
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        f'"단식원 가격","{QUESTION_SOURCE}",고요한아침,selfwriter,질문형,,,,,가입인사\n',
    )
    book = recognize_proxy_workbook(
        write_proxy_xlsx(tmp_path, self_comment_count=2)
    )
    with pytest.raises(GatlingPasteError, match="자사 카페 댓글 아이디가 6개"):
        build_gatling_master(brand, manuscript_only=True, proxy_book=book)


def test_self_comment_pool_drops_affiliate_ids(tmp_path: Path) -> None:
    path = tmp_path / "mixed.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["아이피:포트", "크롬번", "아이디", "비번", "카테고리"])
    sheet.append(["10.0.0.2:3030", 10, "selfwriter", "pw-self", "자사"])
    for offset, name in enumerate(["s1", "s2", "s3", "s4", "s5", "quilliant"]):
        sheet.append(
            [f"10.0.0.{offset + 31}:3030", 31 + offset, name, f"pw-{name}", "자사 댓"]
        )
    workbook.save(path)
    book = recognize_proxy_workbook(path)
    names = {item.account for item in book.self_comment_accounts()}
    assert "quilliant" not in names
    assert names == {"s1", "s2", "s3", "s4", "s5"}


def test_too_few_comment_accounts_raises(tmp_path: Path) -> None:
    brand = write_brand_csv(
        tmp_path,
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        f'"단식원 가격","{QUESTION_SOURCE}",씨씨앙,writer,질문형,,,,,자유수다방\n',
    )
    book = recognize_proxy_workbook(write_proxy_xlsx(tmp_path, comment_count=2))
    with pytest.raises(GatlingPasteError, match="댓글 아이디가 6개"):
        build_gatling_master(brand, manuscript_only=True, proxy_book=book)


def test_paste_writes_chrome_id_and_password(tmp_path: Path) -> None:
    brand = write_brand_csv(
        tmp_path,
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        f'"단식원 가격","{_simple_source("첫째 제목", "첫째 본문")}",고요한아침,selfwriter,후기형,,,,,가입인사\n',
    )
    gatling = write_gatling_xlsx(tmp_path)
    book = recognize_proxy_workbook(write_proxy_xlsx(tmp_path))

    _, start_row = paste_manuscripts_into_gatling(brand, gatling, proxy_book=book)
    workbook = load_workbook(gatling)
    sheet = workbook["마스터"]
    assert start_row == 9
    assert sheet.cell(9, 5).value == 10
    assert sheet.cell(9, 6).value == "selfwriter"
    assert sheet.cell(9, 7).value == "pw-self"


def test_auto_id_count_accepts_two_to_ten() -> None:
    assert normalize_auto_id_count(2) == 2
    assert normalize_auto_id_count(6) == 6
    assert normalize_auto_id_count(10) == 10
    with pytest.raises(GatlingPasteError, match="2~10"):
        normalize_auto_id_count(1)
    with pytest.raises(GatlingPasteError, match="2~10"):
        normalize_auto_id_count(11)


def test_auto_id_labels_keep_six_id_order() -> None:
    assert comment_auto_labels("질문형", 6) == [
        "댓글1",
        "댓글2",
        "대대대댓글2",
        "댓글3",
        "댓글4",
        "댓글5",
    ]
    assert comment_auto_labels("후기형", 6) == [
        "댓글1",
        "댓글2",
        "대대댓글2",
        "댓글3",
        "댓글4",
        "댓글5",
    ]
    assert comment_auto_labels("질문형", 2) == ["댓글1", "댓글2"]
    assert comment_auto_labels("후기형", 3) == ["댓글1", "댓글2", "대대댓글2"]
    assert comment_auto_labels("질문형", 10) == [
        "댓글1",
        "댓글2",
        "대대대댓글2",
        "댓글3",
        "댓글4",
        "댓글5",
        "댓글6",
        "댓글7",
        "댓글8",
        "댓글9",
    ]


def test_two_auto_ids_fill_only_first_comments(tmp_path: Path) -> None:
    job = make_job(cafe="씨씨앙", board="자유수다방")
    job.account = "writer"
    job.article_type = "질문형"
    book = recognize_proxy_workbook(write_proxy_xlsx(tmp_path, comment_count=2))
    rng = random.Random(1)
    chosen = rng.sample(book.comment_accounts(), 2)
    rows = build_master_rows(
        job,
        proxy_book=book,
        rng=random.Random(1),
        comment_id_count=2,
    )

    comments = [row for row in rows if row.type == TYPE_COMMENT]
    replies = [row for row in rows if row.type == TYPE_REPLY]
    assert comments[0].account == chosen[0].account
    assert comments[1].account == chosen[1].account
    assert replies[2].account == chosen[1].account
    assert replies[3].account in {"", None}
    assert {row.account for row in comments} == {item.account for item in chosen}


def test_three_auto_ids_include_special_reply(tmp_path: Path) -> None:
    job = make_job(cafe="씨씨앙", board="자유수다방")
    job.account = "writer"
    job.article_type = "질문형"
    book = recognize_proxy_workbook(write_proxy_xlsx(tmp_path, comment_count=3))
    rng = random.Random(4)
    chosen = rng.sample(book.comment_accounts(), 3)
    rows = build_master_rows(
        job,
        proxy_book=book,
        rng=random.Random(4),
        comment_id_count=3,
    )
    replies = [row for row in rows if row.type == TYPE_REPLY]
    assert replies[2].account == chosen[1].account
    assert replies[3].account == chosen[2].account


def test_ten_auto_ids_need_ten_proxy_accounts(tmp_path: Path) -> None:
    brand = write_brand_csv(
        tmp_path,
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        f'"단식원 가격","{QUESTION_SOURCE}",씨씨앙,writer,질문형,,,,,자유수다방\n',
    )
    book = recognize_proxy_workbook(write_proxy_xlsx(tmp_path, comment_count=10))
    result = build_gatling_master(
        brand,
        manuscript_only=True,
        proxy_book=book,
        comment_id_count=10,
        rng=random.Random(5),
    )
    comments = [row for row in result.rows if row.type == TYPE_COMMENT]
    assert len({row.account for row in comments if row.account}) == 2
    assert all(row.account for row in comments)

    thin = recognize_proxy_workbook(write_proxy_xlsx(tmp_path, comment_count=4))
    with pytest.raises(GatlingPasteError, match="댓글 아이디가 10개"):
        build_gatling_master(
            brand,
            manuscript_only=True,
            proxy_book=thin,
            comment_id_count=10,
        )


def test_selected_count_overrides_default_six(tmp_path: Path) -> None:
    brand = write_brand_csv(
        tmp_path,
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        f'"단식원 가격","{QUESTION_SOURCE}",고요한아침,selfwriter,질문형,,,,,가입인사\n',
    )
    book = recognize_proxy_workbook(
        write_proxy_xlsx(tmp_path, self_comment_count=2)
    )
    result = build_gatling_master(
        brand,
        manuscript_only=True,
        proxy_book=book,
        comment_id_count=2,
    )
    comments = [row for row in result.rows if row.type == TYPE_COMMENT]
    assert all(row.account for row in comments)
    assert {row.account for row in comments} <= {"s1", "s2"}


def test_empty_account_uses_affiliate_authors_not_comment_ids(tmp_path: Path) -> None:
    first = QUESTION_SOURCE.replace("실제 원고 제목", "첫째 제목").replace("실제 원고 본문", "첫째 본문")
    second = QUESTION_SOURCE.replace("실제 원고 제목", "둘째 제목").replace("실제 원고 본문", "둘째 본문")
    third = QUESTION_SOURCE.replace("실제 원고 제목", "셋째 제목").replace("실제 원고 본문", "셋째 본문")
    fourth = QUESTION_SOURCE.replace("실제 원고 제목", "넷째 제목").replace("실제 원고 본문", "넷째 본문")
    brand = write_brand_csv(
        tmp_path,
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        f'"단식원 가격","{first}",씨씨앙,,질문형,,,,,자유수다방\n'
        f'"단식원 후기","{second}",씨씨앙,,질문형,,,,,자유수다방\n'
        f'"단식원 비용","{third}",씨씨앙,,질문형,,,,,자유수다방\n'
        f'"단식원 비교","{fourth}",씨씨앙,,질문형,,,,,자유수다방\n',
    )
    book = recognize_proxy_workbook(
        write_proxy_xlsx(tmp_path, affiliate_author_count=6)
    )
    expected = random.Random(1).sample(book.affiliate_authors(), 2)
    result = build_gatling_master(
        brand,
        manuscript_only=True,
        proxy_book=book,
        author_id_count=2,
        rng=random.Random(1),
    )
    edits = [row for row in result.rows if row.type == TYPE_EDIT_POST]
    comments = [row for row in result.rows if row.type == TYPE_COMMENT]
    assert [row.account for row in edits] == [
        expected[0].account,
        expected[1].account,
        expected[0].account,
        expected[1].account,
    ]
    author_names = {item.account for item in book.affiliate_authors()}
    comment_names = {item.account for item in book.comment_accounts()}
    assert set(row.account for row in edits) <= author_names
    assert set(row.account for row in comments).isdisjoint(author_names)
    assert {row.account for row in comments} <= comment_names


def test_empty_account_uses_self_authors_not_comment_ids(tmp_path: Path) -> None:
    first = QUESTION_SOURCE.replace("실제 원고 제목", "첫째 제목").replace("실제 원고 본문", "첫째 본문")
    second = QUESTION_SOURCE.replace("실제 원고 제목", "둘째 제목").replace("실제 원고 본문", "둘째 본문")
    third = QUESTION_SOURCE.replace("실제 원고 제목", "셋째 제목").replace("실제 원고 본문", "셋째 본문")
    brand = write_brand_csv(
        tmp_path,
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        f'"단식원 가격","{first}",고요한아침,,질문형,,,,,가입인사\n'
        f'"단식원 후기","{second}",고요한아침,,질문형,,,,,가입인사\n'
        f'"단식원 비용","{third}",고요한아침,,질문형,,,,,가입인사\n',
    )
    book = recognize_proxy_workbook(
        write_proxy_xlsx(tmp_path, self_author_count=6)
    )
    expected = random.Random(3).sample(book.self_authors(), 2)
    result = build_gatling_master(
        brand,
        manuscript_only=True,
        proxy_book=book,
        author_id_count=2,
        rng=random.Random(3),
    )
    posts = [row for row in result.rows if row.type == TYPE_NEW_POST]
    comments = [row for row in result.rows if row.type == TYPE_COMMENT]
    assert [row.account for row in posts] == [
        expected[0].account,
        expected[1].account,
        expected[0].account,
    ]
    assert {row.account for row in posts} <= {item.account for item in book.self_authors()}
    assert {row.account for row in comments} <= {
        item.account for item in book.self_comment_accounts()
    }
    assert {row.account for row in posts}.isdisjoint(
        {item.account for item in book.self_comment_accounts()}
    )


def test_filled_sheet_author_is_kept_when_others_auto_assign(tmp_path: Path) -> None:
    first = QUESTION_SOURCE.replace("실제 원고 제목", "첫째 제목").replace("실제 원고 본문", "첫째 본문")
    second = QUESTION_SOURCE.replace("실제 원고 제목", "둘째 제목").replace("실제 원고 본문", "둘째 본문")
    third = QUESTION_SOURCE.replace("실제 원고 제목", "셋째 제목").replace("실제 원고 본문", "셋째 본문")
    brand = write_brand_csv(
        tmp_path,
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        f'"단식원 가격","{first}",씨씨앙,writer,질문형,,,,,자유수다방\n'
        f'"단식원 후기","{second}",씨씨앙,,질문형,,,,,자유수다방\n'
        f'"단식원 비용","{third}",씨씨앙,,질문형,,,,,자유수다방\n',
    )
    book = recognize_proxy_workbook(
        write_proxy_xlsx(tmp_path, affiliate_author_count=6)
    )
    result = build_gatling_master(
        brand,
        manuscript_only=True,
        proxy_book=book,
        author_id_count=2,
        rng=random.Random(7),
    )
    edits = [row for row in result.rows if row.type == TYPE_EDIT_POST]
    assert edits[0].account == "writer"
    auto_names = {edits[1].account, edits[2].account}
    assert auto_names <= {item.account for item in book.affiliate_authors()}
    assert auto_names.isdisjoint({item.account for item in book.comment_accounts()})


def test_author_count_still_needs_six_comment_ids(tmp_path: Path) -> None:
    brand = write_brand_csv(
        tmp_path,
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        f'"단식원 가격","{QUESTION_SOURCE}",고요한아침,selfwriter,질문형,,,,,가입인사\n',
    )
    book = recognize_proxy_workbook(
        write_proxy_xlsx(tmp_path, self_comment_count=2, self_author_count=6)
    )
    with pytest.raises(GatlingPasteError, match="자사 카페 댓글 아이디가 6개"):
        build_gatling_master(
            brand,
            manuscript_only=True,
            proxy_book=book,
            author_id_count=2,
        )


def test_too_few_authors_for_auto_assign_raises(tmp_path: Path) -> None:
    brand = write_brand_csv(
        tmp_path,
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        f'"단식원 가격","{QUESTION_SOURCE}",고요한아침,,질문형,,,,,가입인사\n',
    )
    book = recognize_proxy_workbook(write_proxy_xlsx(tmp_path))
    with pytest.raises(GatlingPasteError, match="자사 작성 아이디가 6개"):
        build_gatling_master(brand, manuscript_only=True, proxy_book=book)
    with pytest.raises(GatlingPasteError, match="자사 작성 아이디가 2개"):
        build_gatling_master(
            brand,
            manuscript_only=True,
            proxy_book=book,
            author_id_count=2,
        )


def test_comment_categories_accept_댓_and_댓글(tmp_path: Path) -> None:
    path = tmp_path / "proxy-labels.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["아이피:포트", "크롬번", "아이디", "비번", "카테고리"])
    sheet.append(["10.0.0.1:3030", 1, "writer", "pw-writer", "제휴"])
    sheet.append(["10.0.0.2:3030", 10, "selfwriter", "pw-self", "자사"])
    affiliate_labels = ["제휴 댓", "제휴댓", "제휴 댓글", "제휴댓글", "제휴 댓", "제휴댓글"]
    self_labels = ["자사 댓", "자사댓", "자사 댓글", "자사댓글", "자사 댓", "자사댓글"]
    for offset, label in enumerate(affiliate_labels):
        name = f"ac{offset + 1}"
        sheet.append(
            [f"10.0.3.{offset}:3030", 40 + offset, name, f"pw-{name}", label]
        )
    for offset, label in enumerate(self_labels):
        name = f"sc{offset + 1}"
        sheet.append(
            [f"10.0.4.{offset}:3030", 50 + offset, name, f"pw-{name}", label]
        )
    workbook.save(path)
    book = recognize_proxy_workbook(path)
    assert {item.account for item in book.comment_accounts()} == {
        "ac1",
        "ac2",
        "ac3",
        "ac4",
        "ac5",
        "ac6",
    }
    assert {item.account for item in book.self_comment_accounts()} == {
        "sc1",
        "sc2",
        "sc3",
        "sc4",
        "sc5",
        "sc6",
    }


FULL_TXT_SOURCE = """제목 :
실제 원고 제목

본문 :
실제 원고 본문

댓글1:
첫 댓글
대댓글1:
첫 답글
댓글2:
둘째 댓글
대댓글2:
둘째 답글
대대댓글2:
깊은 답글
대대대댓글2:
더 깊은 답글
댓글3:
셋째 댓글
댓글4:
넷째 댓글
댓글5:
다섯째 댓글
"""


def test_parse_export_mode_accepts_excel_txt_or_both() -> None:
    assert parse_export_mode("excel") == EXPORT_EXCEL
    assert parse_export_mode("txt") == EXPORT_TXT
    assert parse_export_mode("both") == EXPORT_BOTH
    with pytest.raises(GatlingPasteError, match="엑셀, TXT"):
        parse_export_mode("pdf")


def test_split_manuscript_txt_parts_makes_four_files() -> None:
    article = parse_article("단식원 가격", FULL_TXT_SOURCE)
    parts = split_manuscript_txt_parts(article)
    assert set(parts) == {TXT_TITLE_BODY, TXT_COMMENTS_123, TXT_COMMENTS_45, TXT_REPLIES}
    assert parts[TXT_TITLE_BODY] == "실제 원고 제목\n\n실제 원고 본문\n"
    assert "제목 :" not in parts[TXT_TITLE_BODY]
    assert "본문 :" not in parts[TXT_TITLE_BODY]
    assert parts[TXT_COMMENTS_123] == "첫 댓글\n\n둘째 댓글\n\n셋째 댓글\n"
    assert "댓글1" not in parts[TXT_COMMENTS_123]
    assert parts[TXT_COMMENTS_45] == "넷째 댓글\n\n다섯째 댓글\n"
    assert parts[TXT_REPLIES] == "첫 답글\n\n둘째 답글\n\n깊은 답글\n\n더 깊은 답글\n"
    assert "첫 댓글" not in parts[TXT_REPLIES]
    assert reply_target_value(article.comments[1].children[0].children[0]) == 2.1
    assert reply_target_value(article.comments[1].children[0].children[0].children[0]) == 2.2


def test_parse_completion_link_reads_cafe_and_article() -> None:
    assert parse_completion_cafe_article(
        "https://cafe.naver.com/cantsb/3541968"
    ) == ("cantsb", "3541968")
    assert completion_txt_stem(
        "https://m.cafe.naver.com/cantsb/3541968?from=list"
    ) == "cantsb_3541968"
    with pytest.raises(GatlingPasteError, match="완료 링크"):
        parse_completion_cafe_article("")
    with pytest.raises(GatlingPasteError, match="카페명과 글 번호"):
        parse_completion_cafe_article("https://example.com/nope")


def test_write_gatling_txt_uses_cafe_article_folders(tmp_path: Path) -> None:
    job = GatlingBrandJob(
        row_number=2,
        keyword="단식원 가격",
        article=parse_article("단식원 가격", FULL_TXT_SOURCE),
        cafe="씨씨앙",
        board="자유수다방",
        article_type="질문형",
        completion_url="https://cafe.naver.com/cantsb/3541968",
    )
    written = write_gatling_txt_files([job], tmp_path / "out")
    rels = sorted(path.relative_to(tmp_path / "out").as_posix() for path in written)
    assert rels == [
        "대댓글/cantsb_3541968.txt",
        "댓글1,2,3/cantsb_3541968.txt",
        "댓글4,5/cantsb_3541968.txt",
        "제목본문/cantsb_3541968.txt",
    ]
    comment_123 = (tmp_path / "out" / "댓글1,2,3" / "cantsb_3541968.txt").read_text(
        encoding="utf-8-sig"
    )
    assert comment_123.startswith("첫 댓글")
    assert "댓글1" not in comment_123
    replies = (tmp_path / "out" / "대댓글" / "cantsb_3541968.txt").read_text(
        encoding="utf-8-sig"
    )
    assert "깊은 답글" in replies
    assert "더 깊은 답글" in replies
    assert "대댓글" not in replies
    assert "대대댓글" not in replies


def test_write_gatling_txt_skips_empty_comment_groups(tmp_path: Path) -> None:
    source = "제목 :\n제목만\n\n본문 :\n본문만\n\n댓글1:\n첫 댓글\n"
    job = GatlingBrandJob(
        row_number=2,
        keyword="코숨핏",
        article=parse_article("코숨핏", source),
        cafe="씨씨앙",
        board="자유수다방",
        completion_url="https://cafe.naver.com/cantsb/3541968",
    )
    rels = [path.relative_to(tmp_path).as_posix() for path in write_gatling_txt_files([job], tmp_path)]
    assert "제목본문/cantsb_3541968.txt" in rels
    assert "댓글1,2,3/cantsb_3541968.txt" in rels
    assert "댓글4,5/cantsb_3541968.txt" not in rels
    assert "대댓글/cantsb_3541968.txt" not in rels
    assert (tmp_path / "댓글4,5").is_dir()
    assert (tmp_path / "대댓글").is_dir()


def test_safe_txt_keyword_strips_path_chars() -> None:
    assert safe_txt_keyword("코숨/핏:후기") == "코숨_핏_후기"
    assert safe_txt_keyword("   ") == "키워드"


def test_duplicate_completion_links_get_numbered_txt_names(tmp_path: Path) -> None:
    source = "제목 :\n제목\n\n본문 :\n본문\n\n댓글1:\n댓글\n"
    jobs = [
        GatlingBrandJob(
            row_number=2,
            keyword="코숨핏",
            article=parse_article("코숨핏", source),
            cafe="씨씨앙",
            board="자유수다방",
            completion_url="https://cafe.naver.com/cantsb/3541968",
        ),
        GatlingBrandJob(
            row_number=3,
            keyword="코숨핏",
            article=parse_article("코숨핏", source),
            cafe="씨씨앙",
            board="자유수다방",
            completion_url="https://cafe.naver.com/cantsb/3541968",
        ),
    ]
    names = [path.name for path in write_gatling_txt_files(jobs, tmp_path)]
    assert "cantsb_3541968.txt" in names
    assert "cantsb_3541968_2.txt" in names


def test_write_gatling_txt_requires_completion_link(tmp_path: Path) -> None:
    job = GatlingBrandJob(
        row_number=4,
        keyword="코숨핏",
        article=parse_article("코숨핏", "제목 :\n제목\n\n본문 :\n본문\n\n댓글1:\n댓글\n"),
        cafe="씨씨앙",
        board="자유수다방",
    )
    with pytest.raises(GatlingPasteError, match="행 4"):
        write_gatling_txt_files([job], tmp_path)


def test_export_txt_only_does_not_need_gatling(tmp_path: Path) -> None:
    brand = write_brand_csv(
        tmp_path,
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        f'"단식원 가격","{FULL_TXT_SOURCE}",씨씨앙,writer,질문형,https://cafe.naver.com/cantsb/3541968,,, ,자유수다방\n',
    )
    exported = export_gatling_output(
        brand,
        mode=EXPORT_TXT,
        txt_dir=tmp_path / "txt",
    )
    assert exported.start_row is None
    assert exported.build.rows
    rels = {path.relative_to(tmp_path / "txt").as_posix() for path in exported.txt_paths}
    assert "제목본문/cantsb_3541968.txt" in rels
    assert "댓글1,2,3/cantsb_3541968.txt" in rels
    assert "댓글4,5/cantsb_3541968.txt" in rels
    assert "대댓글/cantsb_3541968.txt" in rels


def test_export_both_writes_excel_and_txt(tmp_path: Path) -> None:
    brand = write_brand_csv(
        tmp_path,
        "키워드,본문,카페명,작성계정,원고유형,완료 링크,말머리,계정유형,이미지 없음,게시판명\n"
        f'"단식원 가격","{FULL_TXT_SOURCE}",고요한아침,writer,질문형,https://cafe.naver.com/cantsb/3541968,,,,"약과 영양, 병원의 기억"\n',
    )
    gatling = write_gatling_xlsx(tmp_path)
    exported = export_gatling_output(
        brand,
        mode=EXPORT_BOTH,
        gatling_path=gatling,
        txt_dir=gatling_txt_folder(gatling),
    )
    assert exported.start_row == 9
    assert (gatling_txt_folder(gatling) / "댓글1,2,3" / "cantsb_3541968.txt").exists()
    workbook = load_workbook(gatling)
    sheet = workbook["마스터"]
    assert sheet.cell(9, 3).value == "실제 원고 제목"
    workbook.close()


def test_gatling_gui_has_export_mode_controls() -> None:
    source = (
        Path(__file__).resolve().parents[1] / "v2r_auto" / "gui_gatling_paste.py"
    ).read_text(encoding="utf-8")
    assert "받는 형식" in source
    assert "TXT 저장 폴더" in source
    assert "cantsb_3541968.txt" in source
    assert "export_gatling_output" in source
