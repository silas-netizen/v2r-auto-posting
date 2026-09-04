from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from v2r_auto.cafe_posts import (
    EXCEL_HEADERS,
    PAGE_LAST_KEYS,
    CafeComment,
    CafePostError,
    CafePostRow,
    apply_intro_cafe_name,
    article_checkbox_ids_from_html,
    article_page_looks_deleted,
    article_ids_from_html,
    article_ids_from_payload,
    article_url,
    board_menu_ids_from_payload,
    cafe_name_from_info_payload,
    cafe_name_from_intro_html,
    decode_naver_payload,
    comments_from_payload,
    duplicate_content_key,
    excel_filename,
    format_comments,
    html_to_text,
    last_page_from_payload,
    list_looks_truncated,
    list_total_from_payload,
    menu_ids_from_html,
    newer_duplicate_rows,
    next_list_last_page,
    parse_cafe_board_target,
    parse_cafe_datetime,
    payload_looks_like_article_list,
    post_from_payload,
    write_cafe_posts_xlsx,
)


def test_parse_cafe_home_url() -> None:
    target = parse_cafe_board_target("https://cafe.naver.com/cantsb")
    assert target.slug == "cantsb"
    assert target.menu_id == 0
    assert target.home_url == "https://cafe.naver.com/cantsb"


def test_parse_board_url_keeps_menu() -> None:
    target = parse_cafe_board_target(
        "https://cafe.naver.com/cantsb",
        "https://cafe.naver.com/f-e/cafes/25016228/menus/12",
    )
    assert target.cafe_id == 25016228
    assert target.menu_id == 12
    assert target.slug == "cantsb"


def test_parse_iframe_board_query() -> None:
    target = parse_cafe_board_target(
        "",
        "https://cafe.naver.com/cantsb?iframe_url=/ArticleList.nhn?search.clubid=25016228&search.menuid=88",
    )
    assert target.slug == "cantsb"
    assert target.cafe_id == 25016228
    assert target.menu_id == 88


def test_empty_address_is_error() -> None:
    with pytest.raises(CafePostError):
        parse_cafe_board_target("", "")


def test_article_url_uses_slug() -> None:
    target = parse_cafe_board_target("https://cafe.naver.com/cantsb")
    assert article_url(target, 345) == "https://cafe.naver.com/cantsb/345"


def test_format_comments_is_nickname_then_text() -> None:
    text = format_comments(
        [
            CafeComment("홍길동", "잘 보고 갑니다"),
            CafeComment("김철수", "감사합니다"),
        ]
    )
    assert text == "홍길동\n잘 보고 갑니다\n\n김철수\n감사합니다"


def test_html_body_keeps_line_breaks() -> None:
    assert "첫째" in html_to_text("<p>첫째</p><br>둘째<script>x()</script>")
    assert "둘째" in html_to_text("<p>첫째</p><br>둘째<script>x()</script>")
    assert "x()" not in html_to_text("<p>첫째</p><br>둘째<script>x()</script>")


def test_article_ids_from_list_payload() -> None:
    payload = {
        "result": {
            "articleList": [
                {"articleId": 11, "subject": "글1"},
                {"articleId": 12, "subject": "글2"},
                {"commentId": 99, "content": "댓글"},
            ]
        }
    }
    assert article_ids_from_payload(payload) == [11, 12]
    assert article_ids_from_html('href="/articles/11" articleid=12') == [11, 12]


def test_comments_from_payload_keep_order() -> None:
    payload = {
        "comments": {
            "items": [
                {"commentId": 1, "nickname": "첫째", "content": "안녕"},
                {"commentId": 2, "nickName": "둘째", "content": "<b>반가워</b>"},
            ]
        }
    }
    comments = comments_from_payload(payload)
    assert [(item.nickname, item.text) for item in comments] == [
        ("첫째", "안녕"),
        ("둘째", "반가워"),
    ]


def test_post_from_payload_reads_fields() -> None:
    payload = {
        "result": {
            "article": {
                "articleId": 77,
                "subject": "제목입니다",
                "contentHtml": "<p>본문입니다</p>",
                "writeDateTimestamp": 1693700000000,
                "writer": {"nickname": "작성닉"},
                "menuName": "자유게시판",
                "cafeName": "씨씨앙",
            },
            "comments": {
                "items": [{"commentId": 3, "nickname": "댓닉", "content": "댓글이다"}]
            },
        }
    }
    row = post_from_payload(payload)
    assert row.cafe_name == "씨씨앙"
    assert row.board_name == "자유게시판"
    assert row.author == "작성닉"
    assert row.title == "제목입니다"
    assert row.body == "본문입니다"
    assert row.comment_text() == "댓닉\n댓글이다"
    assert row.article_id == 77
    assert row.written_at == parse_cafe_datetime(1693700000000)


def test_write_excel_has_requested_columns(tmp_path: Path) -> None:
    path = write_cafe_posts_xlsx(
        tmp_path / excel_filename("씨씨앙", datetime(2026, 9, 3, 10, 28)),
        [
            CafePostRow(
                cafe_name="씨씨앙",
                board_name="자유게시판",
                author="작성닉",
                title="제목",
                body="본문",
                comments=[
                    CafeComment("첫째", "댓글1"),
                    CafeComment("둘째", "댓글2"),
                ],
            )
        ],
    )
    assert path.name == "카페글_씨씨앙_20260903_1028.xlsx"
    book = load_workbook(path)
    sheet = book.active
    assert [cell.value for cell in sheet[1]] == list(EXCEL_HEADERS)
    assert [cell.value for cell in sheet[2]] == [
        "씨씨앙",
        "자유게시판",
        "작성닉",
        "제목",
        "본문",
        "첫째\n댓글1\n\n둘째\n댓글2",
    ]


def test_gui_mentions_excel_and_login() -> None:
    source = Path("v2r_auto/gui_cafe_posts.py").read_text(encoding="utf-8")
    assert "네이버 로그인" in source
    assert "게시판 주소(선택)" in source
    assert "엑셀" in source
    assert "카페소개" in source
    assert "중지를 눌러도" in source
    assert "_write_collected_excel" in source
    assert "중복 글 삭제" in source
    assert "제목·본문이 같은 최신 글 삭제" in source
    assert "마지막 페이지까지" in source
    assert "글이 없는 페이지는 더 보지 않습니다" in source


def test_cp949_intro_html_decodes_to_hangul() -> None:
    html = (
        '<table class="tbl_cafe_info">'
        "<tr><th scope=\"row\">카페 이름</th>"
        '<td><strong class="cafe_name">웨딩 노트</strong></td></tr>'
        "</table>"
    )
    raw = html.encode("cp949")
    broken = raw.decode("utf-8", errors="replace")
    assert cafe_name_from_intro_html(broken) == ""
    assert cafe_name_from_intro_html(decode_naver_payload(raw, "text/html")) == "웨딩 노트"
    assert apply_intro_cafe_name(
        [CafePostRow("임시", "자유", "닉", "제목", "본문")],
        broken,
    )[0].cafe_name == "임시"


def test_intro_html_uses_cafe_name_row() -> None:
    html = """
    <table class="tbl_cafe_info">
      <tr>
        <th scope="row">카페 이름</th>
        <td><strong class="cafe_name">웨딩 노트</strong> <a>수정</a></td>
      </tr>
      <tr>
        <th scope="row">카페 주소</th>
        <td>https://cafe.naver.com/sharfova</td>
      </tr>
      <tr>
        <th scope="row">카페 매니저</th>
        <td>웨딩 노트M</td>
      </tr>
    </table>
    """
    assert cafe_name_from_intro_html(html) == "웨딩 노트"


def test_intro_text_skips_address_and_manager() -> None:
    text = """
    카페소개
    카페 이름
    웨딩 노트
    수정
    카페 주소
    https://cafe.naver.com/sharfova
    카페 매니저
    웨딩 노트M
    """
    assert cafe_name_from_intro_html(text) == "웨딩 노트"


def test_intro_payload_reads_cafe_info_view() -> None:
    payload = {
        "message": {
            "result": {
                "cafeInfoView": {
                    "cafeId": 123,
                    "cafeName": "웨딩 노트",
                    "cafeUrl": "sharfova",
                }
            }
        }
    }
    assert cafe_name_from_info_payload(payload) == "웨딩 노트"


def test_apply_intro_name_overwrites_rows() -> None:
    rows = [
        CafePostRow("sharfova", "자유", "닉", "제목1", "본문1"),
        CafePostRow("웨딩노트", "공지", "닉", "제목2", "본문2"),
    ]
    apply_intro_cafe_name(rows, "웨딩 노트")
    assert [row.cafe_name for row in rows] == ["웨딩 노트", "웨딩 노트"]


def test_browser_reads_intro_name_and_saves_on_stop() -> None:
    source = Path("v2r_auto/cafe_post_browser.py").read_text(encoding="utf-8")
    assert "CafeProfileView.nhn" in source
    assert "_load_intro_cafe_name" in source
    assert "중지를 눌러 여기까지 모은 글만 저장합니다" in source
    assert "row.cafe_name = self.target.slug" not in source
    assert "TextDecoder('euc-kr')" in source
    assert "next_list_last_page" in source
    assert "_list_board_menu_ids" in source
    assert "글이 있는 마지막 페이지까지 모았습니다" in source
    assert "더 이상 새 글이 없어 목록을 끝냅니다" in source


def test_last_page_uses_total_not_navigation_ten() -> None:
    payload = {
        "result": {
            "articleList": [{"articleId": 1, "subject": "글"}],
            "pageInfo": {
                "pageSize": 50,
                "lastNavigationPageNumber": 10,
                "lastPage": 10,
                "totalCount": 1368,
            },
        }
    }
    assert "lastNavigationPageNumber" not in PAGE_LAST_KEYS
    assert list_total_from_payload(payload) == 1368
    assert last_page_from_payload(payload) == 28


def test_last_page_ignores_navigation_only_payload() -> None:
    payload = {"pageInfo": {"lastNavigationPageNumber": 10}}
    assert last_page_from_payload(payload) is None
    assert list_total_from_payload(payload) is None


def test_comment_total_does_not_hide_article_total() -> None:
    payload = {
        "result": {
            "articleList": [{"articleId": 1}],
            "pageInfo": {
                "pageSize": 50,
                "lastNavigationPageNumber": 10,
                "totalCount": 1368,
            },
            "comments": {"totalCount": 3},
        }
    }
    assert last_page_from_payload(payload) == 28


def test_full_page_keeps_checking_past_guessed_ten() -> None:
    assert next_list_last_page(10, 50, 10) == 400
    assert next_list_last_page(1, 50, 28) == 28
    assert next_list_last_page(28, 18, 28) == 28
    assert next_list_last_page(11, 0, 28) == 10


def test_empty_after_full_page_looks_truncated() -> None:
    assert list_looks_truncated(50, 0) is True
    assert list_looks_truncated(18, 0) is False
    assert list_looks_truncated(50, 12) is False


def test_empty_article_list_is_trusted() -> None:
    payload = {"result": {"articleList": []}}
    assert payload_looks_like_article_list(payload) is True
    assert article_ids_from_payload(payload) == []


def test_board_menu_ids_skip_folder_and_all() -> None:
    payload = {
        "message": {
            "result": {
                "menus": [
                    {"menuId": 0, "menuName": "전체글", "menuType": "A"},
                    {"menuId": 1, "menuName": "공지", "menuType": "B"},
                    {"menuId": 2, "menuName": "자유게시판", "menuType": "B"},
                    {"menuId": 10, "menuName": "앨범모음", "menuType": "F"},
                    {"menuId": 11, "menuName": "외부링크", "menuType": "L"},
                ]
            }
        }
    }
    assert board_menu_ids_from_payload(payload) == [1, 2]
    assert menu_ids_from_html(
        'href="/ArticleList.nhn?search.menuid=1" menuid=2 menuid=0'
    ) == [1, 2]


def test_stop_still_writes_partial_excel(tmp_path: Path) -> None:
    rows = apply_intro_cafe_name(
        [
            CafePostRow("임시", "자유", "닉", "제목", "본문"),
        ],
        "웨딩 노트",
    )
    path = write_cafe_posts_xlsx(tmp_path / excel_filename("웨딩 노트"), rows)
    book = load_workbook(path)
    assert book.active["A2"].value == "웨딩 노트"


def test_newer_duplicates_keep_oldest_date() -> None:
    older = CafePostRow(
        "웨딩 노트",
        "자유",
        "닉",
        "같은 제목",
        "같은 본문",
        article_id=10,
        written_at=datetime(2026, 1, 1, 10, 0),
    )
    newer = CafePostRow(
        "웨딩 노트",
        "자유",
        "닉",
        "같은 제목",
        "같은 본문",
        article_id=20,
        written_at=datetime(2026, 8, 1, 10, 0),
    )
    newest = CafePostRow(
        "웨딩 노트",
        "자유",
        "닉",
        " 같은   제목 ",
        "같은\n본문",
        article_id=30,
        written_at=datetime(2026, 9, 1, 10, 0),
    )
    unique = CafePostRow(
        "웨딩 노트",
        "자유",
        "닉",
        "같은 제목",
        "다른 본문",
        article_id=40,
        written_at=datetime(2026, 9, 2, 10, 0),
    )
    deleted = newer_duplicate_rows([newer, unique, newest, older])
    assert [row.article_id for row in deleted] == [20, 30]
    assert duplicate_content_key("같은 제목", "같은 본문") == duplicate_content_key(
        " 같은   제목 ", "같은\n본문"
    )


def test_title_only_match_is_not_duplicate() -> None:
    rows = [
        CafePostRow("카페", "게시판", "닉", "제목", "본문A", article_id=1),
        CafePostRow("카페", "게시판", "닉", "제목", "본문B", article_id=2),
    ]
    assert newer_duplicate_rows(rows) == []


def test_checkbox_ids_from_staff_list_html() -> None:
    html = """
    <table class="article-table">
      <tr>
        <td><input type="checkbox" name="articleid" value="101"></td>
        <td><a class="article" href="/ArticleRead.nhn?articleid=101">첫째</a></td>
      </tr>
      <tr>
        <td><input type="checkbox" name="articleid" value="202"></td>
        <td><a href="/articles/202">둘째</a></td>
      </tr>
    </table>
    """
    assert article_checkbox_ids_from_html(html) == [101, 202]


def test_wedding_note_style_161_posts_delete_newer_only() -> None:
    rows: list[CafePostRow] = []
    expected: list[int] = []
    article_id = 1000
    day = 1
    for index in range(40):
        title = f"웨딩 중복 {index}"
        body = f"본문 내용 {index} " * 3
        older_id = article_id
        article_id += 1
        newer_id = article_id
        article_id += 1
        rows.append(
            CafePostRow(
                "웨딩 노트",
                "자유",
                "닉",
                title,
                body,
                article_id=older_id,
                written_at=datetime(2026, 1, min(day, 28), 9, 0),
            )
        )
        rows.append(
            CafePostRow(
                "웨딩 노트",
                "자유",
                "닉",
                title,
                body,
                article_id=newer_id,
                written_at=datetime(2026, 8, min(day, 28), 9, 0),
            )
        )
        expected.append(newer_id)
        day += 1
    for index in range(20):
        title = f"웨딩 삼중복 {index}"
        body = f"세 번 올린 본문 {index}"
        first = article_id
        article_id += 1
        second = article_id
        article_id += 1
        third = article_id
        article_id += 1
        for offset, aid in enumerate((first, second, third)):
            rows.append(
                CafePostRow(
                    "웨딩 노트",
                    "자유",
                    "닉",
                    title,
                    body,
                    article_id=aid,
                    written_at=datetime(2026, 2 + offset, min(index + 1, 28), 10, 0),
                )
            )
        expected.extend([second, third])
    for index in range(21):
        rows.append(
            CafePostRow(
                "웨딩 노트",
                "자유",
                "닉",
                f"유일 글 {index}",
                f"유일한 본문 {index}",
                article_id=article_id,
                written_at=datetime(2026, 3, min(index + 1, 28), 11, 0),
            )
        )
        article_id += 1
    assert len(rows) == 161
    deleted = newer_duplicate_rows(rows)
    assert len(deleted) == 80
    assert sorted(row.article_id for row in deleted) == sorted(expected)
    assert 1000 not in {row.article_id for row in deleted}
    assert 1001 in {row.article_id for row in deleted}


def test_browser_deletes_with_checkbox() -> None:
    source = Path("v2r_auto/cafe_post_browser.py").read_text(encoding="utf-8")
    assert "board.removeArticles" in source
    assert "delete_newer_duplicates" in source
    assert "관리자 계정" in source
    assert "window.confirm = accept" in source
    assert "_override_confirms" in source
    assert "_article_looks_gone" in source
    assert "전체 글 목록은 넘기지 않고" in source
    assert "text === '예'" in source


def test_article_page_looks_deleted() -> None:
    assert article_page_looks_deleted(
        "삭제되었거나 존재하지 않는 게시글입니다."
    )
    assert article_page_looks_deleted("삭제된 게시글입니다")
    assert not article_page_looks_deleted("안녕하세요 본문입니다")


def test_gui_mentions_delete_confirm() -> None:
    source = Path("v2r_auto/gui_cafe_posts.py").read_text(encoding="utf-8")
    assert "삭제 확인 창까지" in source
    assert "그 글 화면으로 갑니다" in source
