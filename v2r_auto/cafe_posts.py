from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from datetime import datetime
from html import unescape
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import parse_qs, unquote, urlparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


NAVER_LOGIN_URL = "https://nid.naver.com/nidlogin.login"
NAVER_LOGIN_COOKIES = ("NID_AUT", "NID_SES")
RESERVED_CAFE_SLUGS = {
    "articlelist.nhn",
    "articlesearchlist.nhn",
    "articlewrite.nhn",
    "ca-cafes",
    "ca-fe",
    "ca-fes",
    "cafeprofileview.nhn",
    "f-e",
    "managehome.nhn",
}
CAFE_ID_PATTERNS = (
    re.compile(r"/cafes/(\d+)", re.I),
    re.compile(r"(?:clubid|cafeid)=(\d+)", re.I),
    re.compile(r'"cafeId"\s*:\s*"?(\d+)'),
    re.compile(r'"clubId"\s*:\s*"?(\d+)'),
    re.compile(r"g_sClubId\s*=\s*['\"]?(\d+)"),
)
MENU_ID_PATTERNS = (
    re.compile(r"/menus/(\d+)", re.I),
    re.compile(r"(?:menuid|menuId|menu_id)=(\d+)", re.I),
)
EXCEL_HEADERS = ("카페명", "게시판", "작성자 닉네임", "제목", "본문", "댓글")
ARTICLE_ID_KEYS = ("articleId", "articleid", "article_id")
TITLE_KEYS = ("subject", "title", "articleSubject", "articleTitle")
AUTHOR_KEYS = (
    "writerNickname",
    "nickname",
    "nickName",
    "memberNickname",
    "authorNickname",
    "nick",
)
BOARD_KEYS = ("menuName", "boardName", "menuNm", "menu")
CAFE_NAME_KEYS = ("cafeName", "clubName", "cafe")
BODY_KEYS = ("contentHtml", "content", "articleContent", "body")
COMMENT_ID_KEYS = ("commentId", "commentid", "comment_id")
COMMENT_TEXT_KEYS = ("content", "commentContent", "comment", "body", "text")
PAGE_LAST_KEYS = (
    "lastNavigationPageNumber",
    "lastPage",
    "totalPages",
    "pageCount",
    "endPage",
)
DATE_KEYS = (
    "writeDateTimestamp",
    "writeDate",
    "addDate",
    "createdAt",
    "createdDate",
)
DATE_FORMATS = (
    "%Y.%m.%d. %H:%M:%S",
    "%Y.%m.%d. %H:%M",
    "%Y.%m.%d %H:%M:%S",
    "%Y.%m.%d %H:%M",
    "%Y.%m.%d.",
    "%Y.%m.%d",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%dT%H:%M:%S",
    "%Y-%m-%d",
)


class CafePostError(ValueError):
    pass


@dataclass(slots=True)
class CafeBoardTarget:
    home_url: str
    cafe_id: int | None = None
    slug: str = ""
    menu_id: int = 0
    board_url: str = ""

    def with_cafe_id(self, cafe_id: int) -> "CafeBoardTarget":
        return CafeBoardTarget(
            home_url=self.home_url,
            cafe_id=int(cafe_id),
            slug=self.slug,
            menu_id=self.menu_id,
            board_url=self.board_url,
        )


@dataclass(slots=True)
class CafeComment:
    nickname: str
    text: str


@dataclass(slots=True)
class CafePostRow:
    cafe_name: str
    board_name: str
    author: str
    title: str
    body: str
    comments: list[CafeComment] = field(default_factory=list)
    article_id: int = 0
    written_at: datetime | None = None

    def comment_text(self) -> str:
        return format_comments(self.comments)


def clean_text(value: str) -> str:
    return re.sub(r"[ \t]+", " ", unescape(str(value or ""))).strip()


def html_to_text(value: str) -> str:
    text = str(value or "")
    text = re.sub(r"(?is)<script[^>]*>.*?</script>", "", text)
    text = re.sub(r"(?is)<style[^>]*>.*?</style>", "", text)
    text = re.sub(r"(?is)<br\s*/?>", "\n", text)
    text = re.sub(r"(?is)</p>", "\n", text)
    text = re.sub(r"(?is)</div>", "\n", text)
    text = re.sub(r"(?is)<[^>]+>", "", text)
    text = unescape(text)
    text = text.replace("\xa0", " ")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def format_comments(comments: Iterable[CafeComment | tuple[str, str]]) -> str:
    blocks: list[str] = []
    for item in comments:
        if isinstance(item, CafeComment):
            nick, text = item.nickname, item.text
        else:
            nick, text = item
        nick = clean_text(nick)
        text = clean_text(text)
        if not nick and not text:
            continue
        if nick and text:
            blocks.append(f"{nick}\n{text}")
        else:
            blocks.append(nick or text)
    return "\n\n".join(blocks)


def normalize_cafe_url(text: str) -> str:
    raw = (text or "").strip()
    if not raw:
        return ""
    if "://" not in raw:
        raw = "https://" + raw
    return raw


def _first_query_int(url: str, *names: str) -> int | None:
    parsed = urlparse(url)
    chunks = [parsed.query]
    iframe = parse_qs(parsed.query).get("iframe_url") or []
    if iframe:
        chunks.append(unquote(iframe[0]).split("?", 1)[-1] if "?" in unquote(iframe[0]) else "")
        chunks.append(unquote(iframe[0]))
    for chunk in chunks:
        query = parse_qs(chunk)
        for name in names:
            raw = (query.get(name) or [""])[0]
            if str(raw).isdigit():
                return int(raw)
    return None


def cafe_id_from_text(text: str) -> int | None:
    for pattern in CAFE_ID_PATTERNS:
        match = pattern.search(text or "")
        if match:
            return int(match.group(1))
    return None


def menu_id_from_text(text: str) -> int | None:
    found = _first_query_int(text, "search.menuid", "menuid", "menuId", "menu_id")
    if found is not None:
        return found
    for pattern in MENU_ID_PATTERNS:
        match = pattern.search(text or "")
        if match:
            return int(match.group(1))
    return None


def slug_from_url(url: str) -> str:
    match = re.search(r"(?:m\.)?cafe\.naver\.com/([^/?#]+)", url or "", re.I)
    if not match:
        return ""
    candidate = match.group(1)
    if candidate.casefold() in RESERVED_CAFE_SLUGS or candidate.isdigit():
        return ""
    return candidate


def parse_cafe_board_target(cafe_url: str, board_url: str = "") -> CafeBoardTarget:
    cafe_raw = normalize_cafe_url(cafe_url)
    board_raw = normalize_cafe_url(board_url)
    sources = [part for part in (board_raw, cafe_raw) if part]
    if not sources:
        raise CafePostError("카페 주소 또는 게시판 주소를 넣어 주세요")
    joined = " ".join(sources)
    if "cafe.naver.com" not in joined.casefold():
        raise CafePostError(
            "네이버 카페 주소를 넣어 주세요. 예: https://cafe.naver.com/cantsb"
        )
    cafe_id = cafe_id_from_text(joined)
    slug = slug_from_url(cafe_raw) or slug_from_url(board_raw)
    menu_id = menu_id_from_text(board_raw) if board_raw else 0
    if menu_id is None:
        menu_id = 0
    if slug:
        home_url = f"https://cafe.naver.com/{slug}"
    elif cafe_id:
        home_url = f"https://cafe.naver.com/f-e/cafes/{cafe_id}"
    else:
        home_url = (cafe_raw or board_raw).split("?")[0].rstrip("/")
    return CafeBoardTarget(
        home_url=home_url,
        cafe_id=cafe_id,
        slug=slug,
        menu_id=int(menu_id),
        board_url=board_raw,
    )


def article_url(target: CafeBoardTarget, article_id: int | str) -> str:
    aid = str(article_id or "").strip()
    if not aid:
        raise CafePostError("글 번호가 없습니다")
    if target.slug:
        return f"https://cafe.naver.com/{target.slug}/{aid}"
    if target.cafe_id:
        return f"https://cafe.naver.com/f-e/cafes/{target.cafe_id}/articles/{aid}"
    raise CafePostError("카페 번호를 찾지 못했습니다. 카페 주소를 다시 확인해 주세요")


def walk_dicts(node: Any) -> Iterable[dict[str, Any]]:
    if isinstance(node, dict):
        yield node
        for value in node.values():
            yield from walk_dicts(value)
    elif isinstance(node, list):
        for item in node:
            yield from walk_dicts(item)


def _first_str(obj: dict[str, Any], keys: tuple[str, ...]) -> str:
    for key in keys:
        value = obj.get(key)
        if isinstance(value, str) and value.strip():
            return value.strip()
        if isinstance(value, dict):
            nested = _first_str(value, keys)
            if nested:
                return nested
    return ""


def _first_int(obj: dict[str, Any], keys: tuple[str, ...]) -> int | None:
    for key in keys:
        value = obj.get(key)
        if isinstance(value, bool):
            continue
        if isinstance(value, int) and value > 0:
            return value
        if isinstance(value, str) and value.isdigit() and int(value) > 0:
            return int(value)
    return None


def loads_maybe_json(text: str) -> Any:
    raw = (text or "").strip()
    if not raw:
        return None
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        return None


def article_ids_from_payload(payload: Any) -> list[int]:
    found: list[int] = []
    seen: set[int] = set()
    for obj in walk_dicts(payload):
        if _first_int(obj, COMMENT_ID_KEYS) and not _first_str(obj, TITLE_KEYS):
            continue
        article_id = _first_int(obj, ARTICLE_ID_KEYS)
        if article_id is None or article_id in seen:
            continue
        seen.add(article_id)
        found.append(article_id)
    return found


def article_ids_from_html(html: str) -> list[int]:
    found: list[int] = []
    seen: set[int] = set()
    for match in re.finditer(
        r"(?:/articles/|articleid=|articleId=)(\d+)", html or "", flags=re.I
    ):
        article_id = int(match.group(1))
        if article_id in seen:
            continue
        seen.add(article_id)
        found.append(article_id)
    return found


def last_page_from_payload(payload: Any) -> int | None:
    best = 0
    for obj in walk_dicts(payload):
        for key in PAGE_LAST_KEYS:
            value = obj.get(key)
            if isinstance(value, int) and value > best:
                best = value
            elif isinstance(value, str) and value.isdigit() and int(value) > best:
                best = int(value)
    return best or None


def comments_from_payload(payload: Any) -> list[CafeComment]:
    comments: list[CafeComment] = []
    seen: set[int] = set()
    for obj in walk_dicts(payload):
        comment_id = _first_int(obj, COMMENT_ID_KEYS)
        if comment_id is None or comment_id in seen:
            continue
        nick = _first_str(obj, AUTHOR_KEYS)
        text = html_to_text(_first_str(obj, COMMENT_TEXT_KEYS))
        if not nick and not text:
            continue
        seen.add(comment_id)
        comments.append(CafeComment(nickname=nick, text=text))
    return comments


def decode_naver_payload(data: bytes, content_type: str = "") -> str:
    raw = bytes(data or b"")
    ctype = (content_type or "").lower()
    utf8 = raw.decode("utf-8", errors="replace")
    if "json" in ctype or "utf-8" in ctype:
        return utf8
    korean = raw.decode("cp949", errors="replace")
    if any(token in ctype for token in ("euc-kr", "ks_c_5601", "ksc5601", "korean")):
        return korean
    if utf8.count("\ufffd") > korean.count("\ufffd"):
        return korean
    return utf8


def looks_broken_cafe_name(value: str) -> bool:
    name = str(value or "")
    if not name or "\ufffd" in name:
        return True
    if re.search(r"[가-힣]", name):
        return False
    return bool(re.search(r"[\u0100-\u024f]", name))


def clean_intro_cafe_name(value: str) -> str:
    name = clean_text(html_to_text(value))
    name = re.sub(r"(?:\s*(?:수정|EDIT))+$", "", name, flags=re.I).strip()
    if "\n" in name:
        name = clean_text(name.split("\n", 1)[0])
    if not name or name.startswith("http"):
        return ""
    if re.fullmatch(r"카페\s*(이름|주소|매니저|소개|주제|창립일|설명)", name):
        return ""
    if looks_broken_cafe_name(name):
        return ""
    return name


def cafe_name_from_intro_html(html: str) -> str:
    text = html or ""
    match = re.search(
        r'(?is)<table[^>]*class="[^"]*(?:tbl_cafe_info|cafe_info)[^"]*"[^>]*>(.*?)</table>',
        text,
    )
    if match:
        named = re.search(
            r'(?is)<strong[^>]*class="[^"]*cafe_name[^"]*"[^>]*>(.*?)</strong>',
            match.group(1),
        )
        if named:
            name = clean_intro_cafe_name(named.group(1))
            if name:
                return name
        row = re.search(
            r"(?is)<th[^>]*>\s*카페\s*이름\s*</th>\s*<td[^>]*>(.*?)</td>",
            match.group(1),
        )
        if row:
            name = clean_intro_cafe_name(row.group(1))
            if name:
                return name
    row = re.search(
        r"(?is)<th[^>]*>\s*카페\s*이름\s*</th>\s*<td[^>]*>(.*?)</td>",
        text,
    )
    if row:
        name = clean_intro_cafe_name(row.group(1))
        if name:
            return name
    named = re.search(
        r'(?is)<(?:strong|span|a)[^>]*class="[^"]*cafe_name[^"]*"[^>]*>(.*?)</(?:strong|span|a)>',
        text,
    )
    if named:
        name = clean_intro_cafe_name(named.group(1))
        if name:
            return name
    return cafe_name_from_intro_text(text)


def cafe_name_from_intro_text(text: str) -> str:
    lines = [clean_text(line) for line in (text or "").splitlines()]
    lines = [line for line in lines if line]
    for index, line in enumerate(lines):
        if not re.fullmatch(r"카페\s*이름", line):
            continue
        for nxt in lines[index + 1 :]:
            if nxt.casefold() in {"수정", "edit"}:
                continue
            name = clean_intro_cafe_name(nxt)
            if name:
                return name
            break
    match = re.search(
        r"카페\s*이름\s+(.+?)(?:\s+수정\b|\s+카페\s+주소\b|$)",
        re.sub(r"\s+", " ", text or ""),
        flags=re.I,
    )
    if match:
        return clean_intro_cafe_name(match.group(1))
    return ""


def cafe_name_from_info_payload(payload: Any) -> str:
    for obj in walk_dicts(payload):
        view = obj.get("cafeInfoView")
        if isinstance(view, dict):
            name = clean_intro_cafe_name(_first_str(view, CAFE_NAME_KEYS))
            if name:
                return name
        if not any(key in obj for key in ("cafeUrl", "clubUrl", "cafeId", "clubId")):
            continue
        name = clean_intro_cafe_name(_first_str(obj, CAFE_NAME_KEYS))
        if name:
            return name
    return ""


def apply_intro_cafe_name(
    rows: Iterable[CafePostRow], cafe_name: str
) -> list[CafePostRow]:
    name = clean_intro_cafe_name(cafe_name)
    collected = list(rows)
    if not name:
        return collected
    for row in collected:
        row.cafe_name = name
    return collected


def parse_cafe_datetime(value: Any) -> datetime | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        ts = float(value)
        if ts <= 0:
            return None
        if ts > 1e12:
            ts /= 1000.0
        if ts < 1e9:
            return None
        try:
            return datetime.fromtimestamp(ts)
        except (OSError, OverflowError, ValueError):
            return None
    text = clean_text(str(value))
    if not text or "전" in text:
        return None
    text = re.sub(r"\s*[오전오후]\s*", " ", text).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def written_at_from_obj(obj: dict[str, Any]) -> datetime | None:
    for key in DATE_KEYS:
        parsed = parse_cafe_datetime(obj.get(key))
        if parsed:
            return parsed
    return None


def article_dates_from_payload(payload: Any) -> dict[int, datetime]:
    found: dict[int, datetime] = {}
    for obj in walk_dicts(payload):
        if _first_int(obj, COMMENT_ID_KEYS) and not _first_str(obj, TITLE_KEYS):
            continue
        article_id = _first_int(obj, ARTICLE_ID_KEYS)
        written_at = written_at_from_obj(obj)
        if article_id and written_at and article_id not in found:
            found[article_id] = written_at
    return found


def post_from_payload(
    payload: Any,
    *,
    fallback: CafePostRow | None = None,
) -> CafePostRow:
    title = ""
    author = ""
    board = ""
    cafe = ""
    body = ""
    article_id = 0
    written_at: datetime | None = None
    for obj in walk_dicts(payload):
        title = title or _first_str(obj, TITLE_KEYS)
        author = author or _first_str(obj, AUTHOR_KEYS)
        board = board or _first_str(obj, BOARD_KEYS)
        cafe = cafe or _first_str(obj, CAFE_NAME_KEYS)
        raw_body = _first_str(obj, BODY_KEYS)
        if raw_body and len(raw_body) > len(body):
            body = raw_body
        if not article_id:
            if _first_int(obj, COMMENT_ID_KEYS) and not _first_str(obj, TITLE_KEYS):
                continue
            article_id = _first_int(obj, ARTICLE_ID_KEYS) or 0
        if written_at is None:
            written_at = written_at_from_obj(obj)
    base = fallback or CafePostRow("", "", "", "", "")
    return CafePostRow(
        cafe_name=clean_text(cafe) or base.cafe_name,
        board_name=clean_text(board) if isinstance(board, str) else base.board_name,
        author=clean_text(author) or base.author,
        title=clean_text(title) or base.title,
        body=html_to_text(body) or base.body,
        comments=comments_from_payload(payload) or list(base.comments),
        article_id=article_id or base.article_id,
        written_at=written_at or base.written_at,
    )


def normalize_duplicate_text(value: str) -> str:
    return re.sub(r"\s+", " ", clean_text(value)).casefold()


def duplicate_content_key(title: str, body: str) -> str:
    title_key = normalize_duplicate_text(title)
    body_key = normalize_duplicate_text(body)
    if not title_key or not body_key:
        return ""
    return f"{title_key}\n{body_key}"


def row_recency_key(row: CafePostRow) -> tuple[float, int]:
    article_id = int(row.article_id or 0)
    if row.written_at is not None:
        return (row.written_at.timestamp(), article_id)
    return (float(article_id), article_id)


def newer_duplicate_rows(rows: Iterable[CafePostRow]) -> list[CafePostRow]:
    groups: dict[str, list[CafePostRow]] = {}
    for row in rows:
        key = duplicate_content_key(row.title, row.body)
        if not key:
            continue
        groups.setdefault(key, []).append(row)
    to_delete: list[CafePostRow] = []
    for group in groups.values():
        if len(group) < 2:
            continue
        ordered = sorted(group, key=row_recency_key)
        to_delete.extend(ordered[1:])
    return to_delete


def article_checkbox_ids_from_html(html: str) -> list[int]:
    found: list[int] = []
    seen: set[int] = set()
    for match in re.finditer(
        r"<input\b[^>]*type=['\"]checkbox['\"][^>]*>",
        html or "",
        flags=re.I,
    ):
        tag = match.group(0)
        value_match = re.search(r"\bvalue=['\"](\d+)['\"]", tag, flags=re.I)
        if not value_match:
            value_match = re.search(
                r"data-article-?id=['\"](\d+)['\"]", tag, flags=re.I
            )
        if not value_match:
            continue
        article_id = int(value_match.group(1))
        if article_id <= 0 or article_id in seen:
            continue
        seen.add(article_id)
        found.append(article_id)
    for match in re.finditer(r"(?is)<tr\b[^>]*>(.*?)</tr>", html or ""):
        row = match.group(1)
        if "checkbox" not in row.lower():
            continue
        link = re.search(
            r"(?:articleid=|/articles/|articleId=)(\d+)", row, flags=re.I
        )
        if not link:
            continue
        article_id = int(link.group(1))
        if article_id in seen:
            continue
        seen.add(article_id)
        found.append(article_id)
    return found


def excel_filename(cafe_name: str, when: datetime | None = None) -> str:
    stamp = (when or datetime.now()).strftime("%Y%m%d_%H%M")
    slug = re.sub(r"[\\/:*?\"<>|]+", "", clean_text(cafe_name) or "카페")
    slug = re.sub(r"\s+", "", slug) or "카페"
    return f"카페글_{slug}_{stamp}.xlsx"


def write_cafe_posts_xlsx(path: Path, rows: Iterable[CafePostRow]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "카페글"
    sheet.append(list(EXCEL_HEADERS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        sheet.append(
            [
                row.cafe_name,
                row.board_name,
                row.author,
                row.title,
                row.body,
                row.comment_text(),
            ]
        )
    widths = (16, 18, 16, 40, 60, 50)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    book.save(path)
    return path


def cookies_show_naver_login(cookie_names: Iterable[str]) -> bool:
    names = set(cookie_names)
    return any(name in names for name in NAVER_LOGIN_COOKIES)
