from __future__ import annotations

import csv
import re
from collections import deque
from pathlib import Path

from openpyxl import load_workbook

from .content import ContentFormatError, ParsedArticle, parse_article
from .models import ImmediateJob, JobStatus
from .sheet import SheetSchemaError


BRAND_REQUIRED_COLUMNS = {
    "keyword": "키워드",
    "body": "본문",
    "cafe": "카페명",
    "account": "작성계정",
    "article_type": "원고유형",
    "completion_url": "완료 링크",
}
BRAND_OPTIONAL_COLUMNS = {
    "prefix": "말머리",
    "account_type": "계정유형",
    "image_disabled": "이미지 없음",
}
BOARD_HEADERS = {"게시판명", "게시판", "메뉴", "메뉴명"}
DAILY_HEADERS = ("카페명", "게시판명", "각색제목", "각색본문")
DAILY_OPTIONAL_ACCOUNT_HEADER = "작성계정"
DAILY_OPTIONAL_BOARD_URL_HEADER = "게시판링크"
INFORMATIONAL_SHEET_ID = "1vSON0Rej9anDQXcAOXyBrCr50B4MMqZ79FahF4cDPJw"
INFORMATIONAL_SHEET_GID = "1193993260"
ACCOUNT_TEST_SHEET_ID = "1UgcAvHFCpC5N9joC9T5WCATK834F3XAtRrepFv6XbEs"
ACCOUNT_TEST_SHEET_GID = "218285244"
ACCOUNT_TEST_HEADERS = {
    "number": "번호",
    "account": "ID",
    "selected": "테스트 선택",
    "result": "테스트 결과",
    "link": "테스트 링크",
    "tested_at": "테스트 일시",
}
KOREAN_SENTENCE_ENDINGS = tuple(
    sorted(
        {
            "더라고요",
            "더라구요",
            "거든요",
            "했습니다",
            "였습니다",
            "없습니다",
            "있습니다",
            "같습니다",
            "됩니다",
            "랍니다",
            "합니다",
            "했어요",
            "됐어요",
            "였어요",
            "있어요",
            "없어요",
            "같아요",
            "좋아요",
            "보였어요",
            "싶어요",
            "이에요",
            "해요",
            "돼요",
            "예요",
            "아요",
            "어요",
            "네요",
            "군요",
            "까요",
            "했답니다",
            "했다",
            "됐다",
            "한다",
            "된다",
            "있다",
            "없다",
            "같다",
            "좋다",
            "싶다",
            "이다",
            "입니다",
            "죠",
        },
        key=len,
        reverse=True,
    )
)


def _cell(value) -> str:
    return str(value or "").replace("\r\n", "\n").replace("\r", "\n").strip()


def _find_header(headers: list[str], candidates: set[str]) -> str:
    return next((header for header in headers if header.strip() in candidates), "")


def _column_letter(zero_based_index: int) -> str:
    value = zero_based_index + 1
    letters = ""
    while value:
        value, remainder = divmod(value - 1, 26)
        letters = chr(ord("A") + remainder) + letters
    return letters


def is_informational_sheet(sheet_url: str) -> bool:
    return (
        f"/d/{INFORMATIONAL_SHEET_ID}/" in sheet_url
        and f"gid={INFORMATIONAL_SHEET_GID}" in sheet_url
    )


def is_account_test_sheet(sheet_url: str) -> bool:
    return (
        f"/d/{ACCOUNT_TEST_SHEET_ID}/" in sheet_url
        and f"gid={ACCOUNT_TEST_SHEET_GID}" in sheet_url
    )


def format_daily_body(body: str) -> str:
    """Clean daily text and add readable Korean paragraphs."""
    cleaned = body.replace("…", "")
    cleaned = re.sub(r"\.{2,}", "", cleaned)
    cleaned = re.sub(r"(?<!\d)\.(?!\d)", "", cleaned)
    cleaned = re.sub(r"(?<!\d),(?!\d)", "", cleaned)
    cleaned = re.sub(
        r"[\U0001F000-\U0001FAFF\u2600-\u27BF\uFE0F\u200D]",
        "",
        cleaned,
    )
    cleaned = re.sub(r"[ \t]+", " ", cleaned)
    cleaned = "\n".join(line.strip() for line in cleaned.splitlines()).strip()
    if "\n" in cleaned:
        return cleaned
    sentences: list[str] = []
    current: list[str] = []
    for token in re.findall(r"\S+", cleaned):
        current.append(token)
        ending_token = re.sub(r"[!?~ㅋㅎㅠㅜ]+$", "", token)
        if ending_token.endswith(KOREAN_SENTENCE_ENDINGS):
            sentences.append(" ".join(current))
            current = []
    if current:
        sentences.append(" ".join(current))
    if len(sentences) <= 1:
        return cleaned
    if len(sentences) == 2:
        return "\n\n".join(sentences) if len(cleaned) > 100 else cleaned
    paragraphs = [
        "\n".join(sentences[index : index + 2])
        for index in range(0, len(sentences), 2)
    ]
    return "\n\n".join(paragraphs)


def load_brand_immediate_jobs(
    path: str | Path,
    *,
    brand: str,
    format_body: bool = False,
    use_comment_ai: bool = False,
) -> list[ImmediateJob]:
    csv_path = Path(path)
    with csv_path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        raw_headers = list(reader.fieldnames or [])
        headers = [header for header in raw_headers if header]
        missing = [
            header
            for header in BRAND_REQUIRED_COLUMNS.values()
            if header not in headers
        ]
        board_header = _find_header(headers, BOARD_HEADERS)
        if not board_header:
            missing.append("게시판명")
        if missing:
            raise SheetSchemaError(
                "즉시 발행용 시트 열을 찾지 못했습니다: " + ", ".join(missing)
            )

        jobs: list[ImmediateJob] = []
        for row_number, row in enumerate(reader, start=2):
            keyword = _cell(row.get(BRAND_REQUIRED_COLUMNS["keyword"]))
            source = _cell(row.get(BRAND_REQUIRED_COLUMNS["body"]))
            cafe = _cell(row.get(BRAND_REQUIRED_COLUMNS["cafe"]))
            board = _cell(row.get(board_header))
            article_type = _cell(row.get(BRAND_REQUIRED_COLUMNS["article_type"]))
            if not any((keyword, source, cafe, board, article_type)):
                continue
            if not all((keyword, source, cafe, board)):
                continue
            try:
                article = parse_article(keyword, source)
            except ContentFormatError as exc:
                article = ParsedArticle(
                    title="",
                    body="",
                    keyword=keyword,
                    tag=re.sub(r"\s+", "", keyword),
                    comments=[],
                )
                format_error = f"원고 형식 오류: {exc}"
            else:
                format_error = ""
            if format_body:
                article.body = format_daily_body(article.body)
            job = ImmediateJob(
                row_number=row_number,
                article=article,
                cafe=cafe,
                board=board,
                account=_cell(row.get(BRAND_REQUIRED_COLUMNS["account"])),
                article_type=article_type,
                prefix=_cell(row.get(BRAND_OPTIONAL_COLUMNS["prefix"])),
                account_type=_cell(row.get(BRAND_OPTIONAL_COLUMNS["account_type"])),
                image_disabled=(
                    _cell(row.get(BRAND_OPTIONAL_COLUMNS["image_disabled"])).casefold()
                    == "y"
                ),
                brand=brand,
                completion_url=_cell(
                    row.get(BRAND_REQUIRED_COLUMNS["completion_url"])
                ),
                source_kind="brand",
                source_name=csv_path.name,
                use_comment_ai=use_comment_ai,
            )
            if format_error:
                job.status = JobStatus.FAILED
                job.message = format_error
            elif job.completion_url:
                job.status = JobStatus.SKIPPED
                job.message = "완료 링크가 있어 건너뜀"
            elif not job.account and job.account_type not in {"실명", "비실명"}:
                job.status = JobStatus.SKIPPED
                job.message = "작성계정과 계정유형이 모두 비어 있음"
            jobs.append(job)
    if not jobs:
        raise SheetSchemaError("즉시 발행할 브랜드 원고가 없습니다")
    return jobs


def interleave_daily_jobs_by_cafe(
    jobs: list[ImmediateJob],
) -> list[ImmediateJob]:
    """Alternate cafes while preserving each cafe's original row order."""
    queues: dict[str, deque[ImmediateJob]] = {}
    for job in jobs:
        queues.setdefault(job.cafe, deque()).append(job)
    result: list[ImmediateJob] = []
    while queues:
        for cafe in list(queues):
            queue = queues[cafe]
            result.append(queue.popleft())
            if not queue:
                queues.pop(cafe)
    return result


def load_daily_excel_jobs(path: str | Path) -> list[ImmediateJob]:
    workbook_path = Path(path)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    jobs: list[ImmediateJob] = []
    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            try:
                first = next(rows)
            except StopIteration:
                continue
            headers = tuple(_cell(value) for value in first)
            missing = [header for header in DAILY_HEADERS if header not in headers]
            if missing:
                raise SheetSchemaError(
                    f"Excel '{worksheet.title}' 필수 열이 없습니다: "
                    + ", ".join(missing)
                )
            duplicates = [
                header for header in headers if header and headers.count(header) > 1
            ]
            if duplicates:
                raise SheetSchemaError(
                    f"Excel '{worksheet.title}' 열 이름이 중복됩니다: "
                    + ", ".join(sorted(set(duplicates)))
                )
            column = {header: index for index, header in enumerate(headers) if header}

            def value(values, header: str) -> str:
                index = column.get(header)
                return (
                    _cell(values[index])
                    if index is not None and index < len(values)
                    else ""
                )

            for row_number, values in enumerate(rows, start=2):
                cafe, board, title, body = (
                    value(values, header) for header in DAILY_HEADERS
                )
                if not any((cafe, board, title, body)):
                    continue
                if not all((cafe, board, title, body)):
                    continue
                board_url = value(values, DAILY_OPTIONAL_BOARD_URL_HEADER)
                link_match = re.search(
                    r"/cafes/(\d+)/menus/(\d+)",
                    board_url,
                )
                jobs.append(
                    ImmediateJob(
                        row_number=row_number,
                        article=ParsedArticle(
                            title=title,
                            body=format_daily_body(body),
                            keyword="",
                            tag="",
                            comments=[],
                        ),
                        cafe=cafe,
                        board=board,
                        account=value(values, DAILY_OPTIONAL_ACCOUNT_HEADER),
                        cafe_id=int(link_match.group(1)) if link_match else 0,
                        menu_id=int(link_match.group(2)) if link_match else 0,
                        image_disabled=True,
                        source_kind="daily",
                        source_name=f"{workbook_path.name}:{worksheet.title}",
                    )
                )
    finally:
        workbook.close()
    if not jobs:
        raise SheetSchemaError("Excel에 즉시 발행할 일상 글이 없습니다")
    return interleave_daily_jobs_by_cafe(jobs)


def load_account_test_jobs(path: str | Path) -> list[ImmediateJob]:
    csv_path = Path(path)
    with csv_path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        raw_headers = list(reader.fieldnames or [])
        headers = [header for header in raw_headers if header]
        missing = [
            header
            for header in ACCOUNT_TEST_HEADERS.values()
            if header not in headers
        ]
        if missing:
            raise SheetSchemaError(
                "한줄테스트 시트 열을 찾지 못했습니다: " + ", ".join(missing)
            )
        duplicates = [
            header
            for header in ACCOUNT_TEST_HEADERS.values()
            if raw_headers.count(header) > 1
        ]
        if duplicates:
            raise SheetSchemaError(
                "한줄테스트 시트 열 이름이 중복됩니다: "
                + ", ".join(duplicates)
            )
        result_column = _column_letter(
            raw_headers.index(ACCOUNT_TEST_HEADERS["result"])
        )
        link_column = _column_letter(
            raw_headers.index(ACCOUNT_TEST_HEADERS["link"])
        )
        tested_at_column = _column_letter(
            raw_headers.index(ACCOUNT_TEST_HEADERS["tested_at"])
        )
        jobs: list[ImmediateJob] = []
        for row_number, row in enumerate(reader, start=2):
            selected = _cell(row.get(ACCOUNT_TEST_HEADERS["selected"])).casefold()
            if selected not in {"true", "y", "1", "yes"}:
                continue
            number = _cell(row.get(ACCOUNT_TEST_HEADERS["number"]))
            account = _cell(row.get(ACCOUNT_TEST_HEADERS["account"]))
            if not number or not account:
                continue
            text = f"김천kb보험 그라래{number}"
            jobs.append(
                ImmediateJob(
                    row_number=row_number,
                    article=ParsedArticle(
                        title=text,
                        body=f"{text}\n{text}",
                        keyword="",
                        tag="",
                        comments=[],
                    ),
                    cafe="",
                    board="자유게시판",
                    account=account,
                    image_disabled=True,
                    source_kind="account_test",
                    source_name=csv_path.name,
                    use_comment_ai=False,
                    account_test_result_column=result_column,
                    account_test_link_column=link_column,
                    account_test_time_column=tested_at_column,
                )
            )
    if not jobs:
        raise SheetSchemaError("체크된 한줄테스트 계정이 없습니다")
    return jobs
