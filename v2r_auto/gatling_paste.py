from __future__ import annotations

import csv
import random
import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .cafe_catalog import korean_name, normalized_name
from .content import CommentNode, ContentFormatError, ParsedArticle, parse_article
from .daily_posts import DailyPostSheetError, load_daily_posts
from .gatling_accounts import (
    COMMENT_ID_COUNT,
    DEFAULT_AUTO_ID_COUNT,
    ProxyAccount,
    ProxyBook,
    account_for_comment_node,
    assign_auto_authors,
    normalize_auto_id_count,
    resolve_job_accounts,
)
from .images import (
    GoogleDriveImageResolver,
    ResolvedImage,
    brand_from_sheet_title,
)
from .models import DailyPost


DAILY_POST_SHEET_URL = (
    "https://docs.google.com/spreadsheets/d/"
    "1vSON0Rej9anDQXcAOXyBrCr50B4MMqZ79FahF4cDPJw/edit?gid=1842684291#gid=1842684291"
)
AFFILIATE_CAFES = {"씨씨앙", "양평맘"}
AFFILIATE_EXACT_BOARDS = {
    "씨씨앙": "자유 수다방",
    "양평맘": "이모저모 이야기💕",
}
AFFILIATE_BOARD_LINKS = {
    "씨씨앙": "https://cafe.naver.com/f-e/cafes/25016228/menus/328?viewType=L",
    "양평맘": "https://cafe.naver.com/f-e/cafes/22788814/menus/14?viewType=L",
}
SELF_OWNED_CAFE_IDS = {
    "러브인썸": 26616683,
    "마이웨딩드림": 26680163,
    "고요한아침": 14567700,
    "헬씨트리": 23708088,
    "송도포털": 16149995,
    "글로시마이": 15175096,
    "웨딩노트": 15441090,
    "쌍둥이맘모여라": 10174516,
}
SELF_OWNED_BOARD_MENUS = {
    "러브인썸": 18,
    "마이웨딩드림": 1,
    "고요한아침": 29,
    "헬씨트리": 27,
    "송도포털": 19,
    "글로시마이": 50,
    "웨딩노트": 32,
    "쌍둥이맘모여라": 664,
}
SELF_OWNED_EXACT_BOARDS = {
    "러브인썸": "뷰티&미용",
    "마이웨딩드림": "뷰티&다이어트",
    "고요한아침": "약과 영양, 병원의 기억",
    "헬씨트리": "자유로운 건강 수다방",
    "송도포털": "친해지는 수다",
    "글로시마이": "다이어트 · 운동 톡",
    "웨딩노트": "뷰티 · 다이어트",
    "쌍둥이맘모여라": "ㄴ가족업체 자유게시판",
}


def _menu_board_url(cafe_id: int, menu_id: int) -> str:
    return f"https://cafe.naver.com/f-e/cafes/{cafe_id}/menus/{menu_id}?viewType=L"


SELF_OWNED_BOARD_LINKS = {
    key: _menu_board_url(SELF_OWNED_CAFE_IDS[key], menu)
    for key, menu in SELF_OWNED_BOARD_MENUS.items()
}
KNOWN_EXACT_BOARDS = (
    "자유 수다방",
    "이모저모 이야기💕",
    "웨딩홀탑방기",
)
BOARD_NAME_ALIASES = {
    normalized_name("웨딩홀탐방기"): "웨딩홀탑방기",
}

MASTER_SHEET_NAME = "마스터"
MASTER_HEADER_ROW = 6
MASTER_HEADERS = (
    "링크",
    "타입",
    "제목",
    "내용",
    "크롬번호",
    "아이디",
    "비번",
    "해시태그",
    "말머리",
    "게시판이름",
    "멤버공개",
    "댓글비허용",
    "첨부비디오위치",
    "첨부이미지위치",
    "결과",
    "결과링크",
    "아이피",
    "비고",
    "메모",
)
TYPE_NEW_POST = "새글"
TYPE_EDIT_POST = "글수정"
TYPE_COMMENT = "댓글"
TYPE_REPLY = "대댓글"
TYPE_DELAY = "딜레이"
ARTICLE_TYPES = {TYPE_NEW_POST, TYPE_EDIT_POST}
COMMENT_ALLOWED = "허용"
REQUIRED_MASTER_HEADERS = MASTER_HEADERS[:4]
WRITABLE_KINDS = {"xlsx", "xlsm"}
IMAGE_PLACEHOLDER = "{이미지}"
IMAGE_TOKEN_PATTERN = re.compile(r"\{(?:A열\s*)?키워드\}|\{B\s*/\s*A\}|\{BA\}")
XLSB_WRITE_MESSAGE = (
    "고른 파일은 기관총 원본(.xlsb)입니다. "
    "이 형식은 매크로 파일이라 프로그램이 제목·본문·댓글을 직접 넣을 수 없습니다. "
    "엑셀에서 '다른 이름으로 저장'으로 .xlsx 또는 .xlsm 파일을 만든 뒤 그 파일을 선택해 주세요."
)

BRAND_REQUIRED_COLUMNS = {
    "keyword": "키워드",
    "body": "본문",
    "cafe": "카페명",
    "article_type": "원고유형",
}
BOARD_HEADERS = {"게시판명", "게시판", "메뉴", "메뉴명"}
OPTIONAL_COLUMNS = {
    "account": "작성계정",
    "completion_url": "완료 링크",
    "prefix": "말머리",
    "account_type": "계정유형",
    "image_disabled": "이미지 없음",
}


class GatlingPasteError(ValueError):
    pass


@dataclass(slots=True)
class GatlingBrandJob:
    row_number: int
    keyword: str
    article: ParsedArticle
    cafe: str
    board: str
    account: str = ""
    article_type: str = ""
    prefix: str = ""
    account_type: str = ""
    image_disabled: bool = False
    brand: str = ""
    completion_url: str = ""
    cafe_article_url: str = ""
    daily_post: DailyPost | None = None


@dataclass(slots=True)
class MasterRow:
    link: str | int | float | None = None
    type: str = ""
    title: str = ""
    body: str = ""
    chrome_number: int | str = ""
    account: str = ""
    password: str = ""
    hashtag: str = ""
    prefix: str = ""
    board_name: str = ""
    comment_policy: str = ""
    result_link: str = ""
    image_location: str = ""
    ip: str = ""

    def cells(self) -> list[object]:
        return [
            self.link,
            self.type,
            self.title,
            self.body,
            self.chrome_number if self.chrome_number != "" else None,
            self.account or None,
            self.password or None,
            self.hashtag,
            self.prefix,
            self.board_name,
            None,
            self.comment_policy,
            None,
            self.image_location or None,
            None,
            self.result_link or None,
            self.ip or None,
            None,
            None,
        ]


@dataclass(slots=True)
class GatlingBuildResult:
    rows: list[MasterRow]
    jobs: list[GatlingBrandJob] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)
    image_count: int = 0
    account_count: int = 0

    def type_counts(self) -> dict[str, int]:
        counts: dict[str, int] = {}
        for row in self.rows:
            counts[row.type] = counts.get(row.type, 0) + 1
        return counts


@dataclass(slots=True)
class GatlingFileInfo:
    path: Path
    kind: str
    sheet_names: list[str]
    headers: list[str]
    last_data_row: int
    recognized: bool
    writable: bool
    message: str
    master_sheet: str = MASTER_SHEET_NAME
    header_row: int = MASTER_HEADER_ROW
    start_column: int = 1
    row6_preview: str = ""
    next_article_row: int = 0

    @property
    def next_row(self) -> int:
        if self.next_article_row:
            return self.next_article_row
        return max(self.last_data_row + 1, self.header_row + 1)


def is_affiliate_cafe(cafe: str) -> bool:
    return cafe.strip() in AFFILIATE_CAFES


def affiliate_board_link(cafe: str) -> str:
    """양평맘·씨씨앙 새글 링크 열에 넣는 게시판 주소."""
    return AFFILIATE_BOARD_LINKS.get((cafe or "").strip(), "")


def self_owned_cafe_key(cafe: str) -> str:
    hangul = korean_name(cafe)
    compact = normalized_name(cafe)
    for key in SELF_OWNED_BOARD_LINKS:
        if hangul == key or compact == normalized_name(key) or hangul.startswith(key):
            return key
    return ""


def self_owned_board_link(cafe: str, board: str = "") -> str:
    """자사 카페 새글 링크 열에 넣는, 카페마다 하나인 게시판 주소."""
    key = self_owned_cafe_key(cafe)
    if not key:
        return ""
    return SELF_OWNED_BOARD_LINKS[key]


def exact_board_name(
    sheet_board: str,
    cafe: str,
    extra_exact_names: list[str] | tuple[str, ...] = (),
) -> str:
    """Compare our sheet board without spaces, write the real cafe board name."""
    wanted = (sheet_board or "").strip()
    cafe_name = cafe.strip()
    known: list[str] = []
    default = AFFILIATE_EXACT_BOARDS.get(cafe_name, "")
    if default:
        known.append(default)
    self_board = SELF_OWNED_EXACT_BOARDS.get(self_owned_cafe_key(cafe_name), "")
    if self_board:
        known.append(self_board)
    known.extend(KNOWN_EXACT_BOARDS)
    known.extend(name.strip() for name in extra_exact_names if str(name).strip())

    unique_known: list[str] = []
    seen: set[str] = set()
    for name in known:
        key = normalized_name(name)
        if key and key not in seen:
            unique_known.append(name)
            seen.add(key)

    if not wanted:
        if default:
            return default
        if self_board:
            return self_board
        raise GatlingPasteError(f"{cafe_name} 게시판명이 비어 있습니다")

    alias = BOARD_NAME_ALIASES.get(normalized_name(wanted))
    if alias:
        return alias

    wanted_key = normalized_name(wanted)
    matches = [name for name in unique_known if normalized_name(name) == wanted_key]
    if len(matches) == 1:
        return matches[0]
    if len(matches) > 1:
        raise GatlingPasteError(
            f"{cafe_name} 게시판명이 여러 개와 맞습니다: {wanted}"
        )
    if default:
        raise GatlingPasteError(
            f"{cafe_name} 게시판명을 실제 카페 게시판과 맞출 수 없습니다: {wanted}. "
            f"기관총에는 '{default}'처럼 정확한 이름이 필요합니다"
        )
    return wanted


def replace_image_tokens(text: str) -> str:
    """기관총은 {이미지}만 인식하므로 시트 표기를 맞춰 넣는다."""
    return IMAGE_TOKEN_PATTERN.sub(IMAGE_PLACEHOLDER, text or "")


def manuscript_text_key(title: object, body: object) -> tuple[str, str]:
    return (replace_image_tokens(_cell(title)), replace_image_tokens(_cell(body)))


def drop_duplicate_manuscripts(
    jobs: list[GatlingBrandJob],
    existing: set[tuple[str, str]] | None = None,
) -> tuple[list[GatlingBrandJob], list[str]]:
    seen = set(existing or ())
    kept: list[GatlingBrandJob] = []
    skipped: list[str] = []
    for job in jobs:
        key = manuscript_text_key(job.article.title, job.article.body)
        if not (key[0] and key[1]):
            kept.append(job)
            continue
        if key in seen:
            skipped.append(f"행 {job.row_number}: 제목·본문이 이미 있어 건너뜀")
            continue
        seen.add(key)
        kept.append(job)
    return kept, skipped


def load_existing_manuscript_keys(path: str | Path) -> set[tuple[str, str]]:
    info = recognize_gatling_workbook(path)
    if not info.recognized:
        return set()
    workbook = load_workbook(info.path, data_only=True)
    try:
        if info.master_sheet not in workbook.sheetnames:
            return set()
        sheet = workbook[info.master_sheet]
        keys: set[tuple[str, str]] = set()
        type_col = info.start_column + 1
        title_col = info.start_column + 2
        body_col = info.start_column + 3
        for row_number in range(info.header_row + 1, sheet.max_row + 1):
            typ = _cell(sheet.cell(row_number, type_col).value)
            if typ not in ARTICLE_TYPES:
                continue
            key = manuscript_text_key(
                sheet.cell(row_number, title_col).value,
                sheet.cell(row_number, body_col).value,
            )
            if key[0] and key[1]:
                keys.add(key)
        return keys
    finally:
        workbook.close()


def reply_target_value(node: CommentNode) -> int | float:
    """기관총 대댓글 A열에 원래 쓰이는 대상 번호. URL을 넣지 않는다."""
    if node.depth <= 0:
        raise GatlingPasteError(f"대댓글 대상 번호가 없습니다: {node.label}")
    if node.depth == 1:
        return node.index
    if node.depth in {2, 3}:
        return float(f"{node.index}.{node.depth - 1}")
    raise GatlingPasteError(f"지원하지 않는 댓글 깊이입니다: {node.label}")


def _cell(value: object) -> str:
    return str(value or "").replace("\r\n", "\n").replace("\r", "\n").strip()


_LABEL_WRAP_CHARS = set(
    "#\"'`*~_=[]()【】「」『』“”‘’·•※★☆∙\ufeff"
)
_PASTE_MARK_PATTERN = re.compile(r'[#"“”‘’*`]+')
_SECTION_LINE_PATTERN = re.compile(
    r"^(?P<head>.*?)(?P<label>제목|본문)\s*[:：]\s*(?P<tail>.*)$"
)
_COMMENT_LINE_PATTERN = re.compile(
    r"^(?P<head>.*?)(?P<dashes>대*)댓글(?:\s*(?P<number>\d+))?\s*[:：]\s*(?P<tail>.*)$"
)
_WRAP_PAIRS = (('"', '"'), ("'", "'"), ("“", "”"), ("‘", "’"))


def _head_is_only_marks(head: str) -> bool:
    return all(ch in _LABEL_WRAP_CHARS or ch.isspace() for ch in head)


def strip_gatling_marks(text: str) -> str:
    """Drop decorative # / quotes / * so they are never pasted into 기관총."""
    cleaned = _PASTE_MARK_PATTERN.sub("", text or "")
    cleaned = re.sub(r"[ \t]+\n", "\n", cleaned)
    cleaned = re.sub(r"[ \t]{2,}", " ", cleaned)
    return cleaned.strip()


def _unwrap_source_quotes(source: str) -> str:
    text = (source or "").strip()
    for left, right in _WRAP_PAIRS:
        if len(text) >= 2 and text.startswith(left) and text.endswith(right):
            return text[len(left) : -len(right)].strip()
    return text


def normalize_gatling_source(source: str) -> str:
    """Keep 제목 / 본문 / 댓글 labels even when #, quotes, or * wrap them."""
    text = _unwrap_source_quotes(source.replace("\r\n", "\n").replace("\r", "\n"))
    comment_index = 0
    reply_at_depth: dict[int, int] = {}
    lines: list[str] = []
    for line in text.split("\n"):
        section = _SECTION_LINE_PATTERN.match(line)
        if section and _head_is_only_marks(section.group("head")):
            tail = strip_gatling_marks(section.group("tail"))
            lines.append(f"{section.group('label')} : {tail}".rstrip())
            continue
        comment = _COMMENT_LINE_PATTERN.match(line)
        if comment and _head_is_only_marks(comment.group("head")):
            depth = len(comment.group("dashes") or "")
            number = comment.group("number")
            if number:
                index = int(number)
            elif depth == 0:
                comment_index += 1
                index = comment_index
                reply_at_depth = {}
            else:
                reply_at_depth[depth] = reply_at_depth.get(depth, 0) + 1
                index = reply_at_depth[depth]
            tail = strip_gatling_marks(comment.group("tail"))
            lines.append(f"{'대' * depth}댓글{index}: {tail}".rstrip())
            continue
        lines.append(line)
    return "\n".join(lines)


def parse_gatling_article(keyword: str, source: str) -> ParsedArticle:
    article = parse_article(keyword, normalize_gatling_source(source))
    article.title = strip_gatling_marks(article.title)
    article.body = strip_gatling_marks(article.body)

    def clean(nodes: list[CommentNode]) -> None:
        for node in nodes:
            node.text = strip_gatling_marks(node.text)
            clean(node.children)

    clean(article.comments)
    return article


def _find_header(headers: list[str], candidates: set[str]) -> str:
    return next((header for header in headers if header.strip() in candidates), "")


def workbook_kind(path: str | Path) -> str:
    suffix = Path(path).suffix.casefold()
    if suffix in {".xlsx", ".xlsm", ".xlsb"}:
        return suffix[1:]
    return "unknown"


def is_master_header_row(values: list[object]) -> bool:
    return find_header_start(_cell(value) for value in values) is not None


def find_header_start(values) -> int | None:
    cleaned = [_cell(value) for value in values]
    wanted = list(REQUIRED_MASTER_HEADERS)
    for start, _value in enumerate(cleaned):
        if cleaned[start : start + len(wanted)] == wanted:
            return start
    return None


def _choose_master_sheet(sheet_names: list[str]) -> str:
    for name in sheet_names:
        if _cell(name) == MASTER_SHEET_NAME:
            return name
    return ""


def _preview_values(values: list[object], limit: int = 6) -> str:
    items = [_cell(value) or "(빈칸)" for value in values[:limit]]
    return ", ".join(items) if items else "(비어 있음)"


def _column_value(values: list[object], start_column: int, offset: int) -> object:
    index = start_column - 1 + offset
    if index < 0 or index >= len(values):
        return None
    return values[index]


def _row_type(values: list[object], start_column: int) -> str:
    return _cell(_column_value(values, start_column, 1))


def _row_has_title_or_body(values: list[object], start_column: int) -> bool:
    title = _column_value(values, start_column, 2)
    body = _column_value(values, start_column, 3)
    return any(_cell(value) for value in (title, body))


def _row_has_manuscript(values: list[object], start_column: int) -> bool:
    """A row is used when 링크, 제목, or 내용 has text. 타입만 있는 칸은 빈 칸으로 본다."""
    link = _column_value(values, start_column, 0)
    return bool(_cell(link)) or _row_has_title_or_body(values, start_column)


def _paste_slot_kind(values: list[object], start_column: int) -> str:
    """Empty title/body rows can take a manuscript. 딜레이와 이미 쓴 칸은 뺀다."""
    typ = _row_type(values, start_column)
    if typ == TYPE_DELAY or _row_has_title_or_body(values, start_column):
        return ""
    if typ in ARTICLE_TYPES:
        return "article"
    if typ == TYPE_COMMENT:
        return "comment"
    if typ == TYPE_REPLY:
        return "reply"
    if typ == "":
        return "blank"
    return ""


def _is_empty_article_values(
    values: list[object],
    start_column: int,
    comment_block_started: bool = False,
) -> bool:
    """Empty 새글/글수정 slot, or a blank 타입 row that can become one."""
    del comment_block_started
    return _paste_slot_kind(values, start_column) in {"article", "blank"}


def _scan_sheet_rows(
    rows: list[tuple[int, list[object]]],
) -> tuple[int, int, list[str], int, str, int]:
    header_row = 0
    start_column = 1
    headers: list[str] = []
    last_row = 0
    row6_preview = ""
    next_article_row = 0
    for row_number, values in rows:
        if row_number == MASTER_HEADER_ROW:
            row6_preview = _preview_values(values)
        start = find_header_start(values)
        if start is not None and not headers:
            header_row = row_number
            start_column = start + 1
            headers = [_cell(value) for value in values[start : start + len(MASTER_HEADERS)]]
            last_row = row_number
            continue
        if not headers:
            continue
        if _row_has_manuscript(values, start_column):
            last_row = row_number
        if not next_article_row and _is_empty_article_values(values, start_column):
            next_article_row = row_number
    return header_row, start_column, headers, last_row, row6_preview, next_article_row


def _xlsx_master_preview(
    path: Path,
) -> tuple[list[str], str, int, int, list[str], int, str, int]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        names = list(workbook.sheetnames)
        sheet_name = _choose_master_sheet(names)
        if not sheet_name:
            return names, "", 0, 1, [], 0, "", 0
        sheet = workbook[sheet_name]
        rows: list[tuple[int, list[object]]] = []
        for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_col=30), start=1):
            rows.append((row_number, [cell.value for cell in row]))
        (
            header_row,
            start_column,
            headers,
            last_row,
            row6_preview,
            next_article_row,
        ) = _scan_sheet_rows(rows)
        return (
            names,
            sheet_name,
            header_row,
            start_column,
            headers,
            last_row,
            row6_preview,
            next_article_row,
        )
    finally:
        workbook.close()


def _xlsb_master_preview(
    path: Path,
) -> tuple[list[str], str, int, int, list[str], int, str, int]:
    try:
        from pyxlsb import open_workbook
    except ImportError as exc:
        raise GatlingPasteError(
            "기관총 원본(.xlsb)을 확인하려면 pyxlsb가 필요합니다"
        ) from exc

    with open_workbook(str(path)) as workbook:
        names = list(workbook.sheets)
        sheet_name = _choose_master_sheet(names)
        if not sheet_name:
            return names, "", 0, 1, [], 0, "", 0
        rows: list[tuple[int, list[object]]] = []
        with workbook.get_sheet(sheet_name) as sheet:
            for row_number, row in enumerate(sheet.rows(), start=1):
                values = [cell.v for cell in row]
                cells = values[1:] if values and values[0] is None else values
                rows.append((row_number, cells))
        (
            header_row,
            start_column,
            headers,
            last_row,
            row6_preview,
            next_article_row,
        ) = _scan_sheet_rows(rows)
        return (
            names,
            sheet_name,
            header_row,
            start_column,
            headers,
            last_row,
            row6_preview,
            next_article_row,
        )


def recognize_gatling_workbook(path: str | Path) -> GatlingFileInfo:
    """Confirm the chosen file is 기관총 by finding 링크/타입/제목/내용 headers."""
    file_path = Path(path)
    if not file_path.exists():
        raise GatlingPasteError(f"기관총 엑셀 파일이 없습니다: {file_path}")

    kind = workbook_kind(file_path)
    if kind == "unknown":
        raise GatlingPasteError(
            "기관총 엑셀은 .xlsx, .xlsm, .xlsb 파일만 선택할 수 있습니다"
        )

    if kind == "xlsb":
        (
            sheet_names,
            master_sheet,
            header_row,
            start_column,
            headers,
            last_data_row,
            row6_preview,
            next_article_row,
        ) = _xlsb_master_preview(file_path)
    else:
        try:
            (
                sheet_names,
                master_sheet,
                header_row,
                start_column,
                headers,
                last_data_row,
                row6_preview,
                next_article_row,
            ) = _xlsx_master_preview(file_path)
        except Exception as exc:
            raise GatlingPasteError(
                "엑셀 파일을 열지 못했습니다. 파일이 열려 있으면 닫고 다시 선택해 주세요"
            ) from exc

    recognized = bool(master_sheet and headers)
    writable = recognized and kind in WRITABLE_KINDS
    resolved_header = header_row or MASTER_HEADER_ROW
    resolved_last = last_data_row or resolved_header
    resolved_next_article = next_article_row or (resolved_last + 1)
    if not master_sheet:
        message = (
            "이 파일에는 '마스터' 시트가 없어 기관총 엑셀로 보지 않습니다. "
            f"지금 시트: {', '.join(sheet_names) or '(없음)'}"
        )
    elif not headers:
        message = (
            "마스터에서 '링크, 타입, 제목, 내용' 열을 찾지 못했습니다. "
            f"6행에 보이는 값: {row6_preview or '(비어 있음)'}. "
            "엑셀을 열어 마스터 시트에 그 네 칸이 있는지 확인해 주세요. "
            "파일이 열려 있으면 저장한 뒤 닫고 다시 선택해 주세요."
        )
    elif kind == "xlsb":
        message = (
            f"기관총 파일로 확인했습니다. 제목·본문이 비어 있는 칸은 "
            f"{resolved_next_article}행입니다. 타입이 비어 있어도 "
            "새글·글수정·댓글·대댓글을 알아서 적습니다. " + XLSB_WRITE_MESSAGE
        )
    else:
        message = (
            f"기관총 파일로 확인했습니다. 제목·본문이 비어 있는 "
            f"{resolved_next_article}행부터 시트 원고를 모두 넣고, "
            "타입이 비어 있으면 새글·글수정·댓글·대댓글을 알아서 적습니다"
        )

    return GatlingFileInfo(
        path=file_path,
        kind=kind,
        sheet_names=sheet_names,
        headers=headers,
        last_data_row=resolved_last,
        recognized=recognized,
        writable=writable,
        message=message,
        master_sheet=master_sheet or MASTER_SHEET_NAME,
        header_row=resolved_header,
        start_column=start_column,
        row6_preview=row6_preview,
        next_article_row=resolved_next_article,
    )


def require_writable_gatling(path: str | Path) -> GatlingFileInfo:
    info = recognize_gatling_workbook(path)
    if not info.recognized:
        raise GatlingPasteError(info.message)
    if not info.writable:
        raise GatlingPasteError(info.message)
    return info


def _sheet_row_values(sheet, row_number: int, start_column: int) -> list[object]:
    values: list[object] = [None] * (start_column + 3)
    values[start_column - 1] = sheet.cell(row_number, start_column).value
    values[start_column] = sheet.cell(row_number, start_column + 1).value
    values[start_column + 1] = sheet.cell(row_number, start_column + 2).value
    values[start_column + 2] = sheet.cell(row_number, start_column + 3).value
    return values


def _sheet_row_has_title_or_body(sheet, row_number: int, start_column: int) -> bool:
    return _row_has_title_or_body(_sheet_row_values(sheet, row_number, start_column), start_column)


def _collect_empty_type_slots(
    sheet, info: GatlingFileInfo
) -> tuple[list[int], list[int], list[int], list[int]]:
    last_sheet_row = max(sheet.max_row or info.header_row, info.header_row)
    article_slots: list[int] = []
    comment_slots: list[int] = []
    reply_slots: list[int] = []
    blank_slots: list[int] = []
    for row_number in range(info.header_row + 1, last_sheet_row + 1):
        values = _sheet_row_values(sheet, row_number, info.start_column)
        kind = _paste_slot_kind(values, info.start_column)
        if kind == "article":
            article_slots.append(row_number)
        elif kind == "comment":
            comment_slots.append(row_number)
        elif kind == "reply":
            reply_slots.append(row_number)
        elif kind == "blank":
            blank_slots.append(row_number)
    return article_slots, comment_slots, reply_slots, blank_slots


def _target_rows_for_paste(
    sheet, info: GatlingFileInfo, rows: list[MasterRow]
) -> list[int]:
    """Use matching empty rows, then empty-타입 rows, then append so every job is written."""
    article_slots, comment_slots, reply_slots, blank_slots = _collect_empty_type_slots(
        sheet, info
    )
    used: set[int] = set()
    append_at = max(sheet.max_row or info.header_row, info.header_row) + 1
    targets: list[int] = []

    def take(pool: list[int]) -> int:
        nonlocal append_at
        for candidate in (pool, blank_slots):
            for row_number in candidate:
                if row_number not in used:
                    used.add(row_number)
                    return row_number
        while append_at in used or _sheet_row_has_title_or_body(
            sheet, append_at, info.start_column
        ):
            append_at += 1
            if append_at > info.header_row + 20000:
                raise GatlingPasteError("마스터에서 이어 넣을 빈 행을 찾지 못했습니다")
        chosen = append_at
        used.add(chosen)
        append_at += 1
        return chosen

    for row in rows:
        if row.type in ARTICLE_TYPES:
            targets.append(take(article_slots))
        elif row.type == TYPE_COMMENT:
            targets.append(take(comment_slots))
        elif row.type == TYPE_REPLY:
            targets.append(take(reply_slots))
        else:
            targets.append(take([]))
    return targets


def _write_master_row(sheet, row_number: int, start_column: int, row: MasterRow) -> None:
    """Always write 링크/타입/제목/내용 so leftover 타입·링크도 원고에 맞게 바꾼다."""
    for offset, value in enumerate(row.cells()):
        cell = sheet.cell(row_number, start_column + offset)
        if offset < 4:
            cell.value = None if value in (None, "") else value
        elif value not in (None, ""):
            cell.value = value


def append_master_rows(path: str | Path, rows: list[MasterRow]) -> int:
    info = require_writable_gatling(path)
    keep_vba = info.kind == "xlsm"
    workbook = load_workbook(info.path, keep_vba=keep_vba)
    try:
        sheet = workbook[info.master_sheet]
        target_rows = _target_rows_for_paste(sheet, info, rows)
        for row_number, row in zip(target_rows, rows):
            _write_master_row(sheet, row_number, info.start_column, row)
        workbook.save(info.path)
        return target_rows[0]
    finally:
        workbook.close()


def load_gatling_brand_jobs(
    path: str | Path,
    *,
    skip_completed: bool = False,
    brand: str = "",
) -> tuple[list[GatlingBrandJob], list[str]]:
    csv_path = Path(path)
    if not csv_path.exists():
        raise GatlingPasteError(f"브랜드 시트 파일이 없습니다: {csv_path}")

    detected_brand = brand or brand_from_sheet_title(csv_path.stem)
    skipped: list[str] = []
    with csv_path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        headers = [header for header in (reader.fieldnames or []) if header]
        missing = [
            header
            for header in BRAND_REQUIRED_COLUMNS.values()
            if header not in headers
        ]
        board_header = _find_header(headers, BOARD_HEADERS)
        if not board_header:
            missing.append("게시판명")
        if missing:
            raise GatlingPasteError(
                "브랜드 시트 열을 찾지 못했습니다: " + ", ".join(missing)
            )

        jobs: list[GatlingBrandJob] = []
        for row_number, row in enumerate(reader, start=2):
            keyword = _cell(row.get(BRAND_REQUIRED_COLUMNS["keyword"]))
            source = _cell(row.get(BRAND_REQUIRED_COLUMNS["body"]))
            cafe = _cell(row.get(BRAND_REQUIRED_COLUMNS["cafe"]))
            article_type = _cell(row.get(BRAND_REQUIRED_COLUMNS["article_type"]))
            board = _cell(row.get(board_header))
            if not any((keyword, source, cafe, article_type, board)):
                continue
            if not all((keyword, source, cafe)):
                skipped.append(f"행 {row_number}: 키워드·본문·카페명이 비어 있음")
                continue
            completion_url = _cell(row.get(OPTIONAL_COLUMNS["completion_url"]))
            if skip_completed and completion_url:
                skipped.append(f"행 {row_number}: F열 완료 링크가 있어 건너뜀")
                continue
            try:
                article = parse_gatling_article(keyword, source)
            except ContentFormatError as exc:
                skipped.append(f"행 {row_number}: 원고 형식 오류 ({exc})")
                continue
            jobs.append(
                GatlingBrandJob(
                    row_number=row_number,
                    keyword=keyword,
                    article=article,
                    cafe=cafe,
                    board=board,
                    account=_cell(row.get(OPTIONAL_COLUMNS["account"])),
                    article_type=article_type,
                    prefix=_cell(row.get(OPTIONAL_COLUMNS["prefix"])),
                    account_type=_cell(row.get(OPTIONAL_COLUMNS["account_type"])),
                    image_disabled=(
                        _cell(row.get(OPTIONAL_COLUMNS["image_disabled"])).casefold()
                        == "y"
                    ),
                    brand=detected_brand,
                    completion_url=completion_url,
                )
            )
    if not jobs:
        if skipped:
            preview = "\n".join(skipped[:8])
            extra = f"\n외 {len(skipped) - 8}건" if len(skipped) > 8 else ""
            raise GatlingPasteError(
                "붙여넣을 브랜드 원고가 없습니다.\n" + preview + extra
            )
        raise GatlingPasteError(
            "붙여넣을 브랜드 원고가 없습니다. "
            "시트에 키워드·본문·카페명이 있는 행이 있는지 확인하세요."
        )
    return jobs, skipped


def assign_daily_posts(
    jobs: list[GatlingBrandJob],
    daily_posts: list[DailyPost],
    rng: random.Random | None = None,
) -> None:
    randomizer = rng or random.SystemRandom()
    for cafe in ("씨씨앙", "양평맘"):
        cafe_jobs = [
            job
            for job in jobs
            if is_affiliate_cafe(job.cafe)
            and job.cafe == cafe
            and job.daily_post is None
        ]
        if not cafe_jobs:
            continue
        candidates = [post for post in daily_posts if post.cafe == cafe]
        if len(candidates) < len(cafe_jobs):
            raise DailyPostSheetError(
                f"{cafe} 일상 글이 부족합니다: 필요 {len(cafe_jobs)}개, "
                f"사용 가능 {len(candidates)}개"
            )
        chosen = randomizer.sample(candidates, len(cafe_jobs))
        for job, post in zip(cafe_jobs, chosen):
            job.daily_post = post


def _with_account(row: MasterRow, account: ProxyAccount | None) -> MasterRow:
    if account is None:
        return row
    row.chrome_number = account.chrome_number
    row.account = account.account
    row.password = account.password
    row.ip = account.ip
    return row


def _article_row(
    *,
    type_name: str,
    title: str,
    body: str,
    job: GatlingBrandJob,
    board_name: str,
    link: str = "",
    with_hashtag: bool = False,
    image_location: str = "",
    author: ProxyAccount | None = None,
) -> MasterRow:
    return _with_account(
        MasterRow(
            link=link or None,
            type=type_name,
            title=replace_image_tokens(title),
            body=replace_image_tokens(body),
            hashtag=job.keyword if with_hashtag else "",
            prefix=job.prefix,
            board_name=board_name,
            comment_policy=COMMENT_ALLOWED,
            image_location=image_location if with_hashtag else "",
        ),
        author,
    )


def _comment_and_reply_rows(
    job: GatlingBrandJob,
    article_url: str,
    *,
    author: ProxyAccount | None = None,
    comment_map: dict[str, ProxyAccount] | None = None,
) -> list[MasterRow]:
    """수정 발행 순서: 댓글을 모두 넣은 뒤, 대댓글은 얕은 것부터 넣는다."""
    rows: list[MasterRow] = []
    mapping = comment_map or {}
    for comment in job.article.comments:
        rows.append(
            _with_account(
                MasterRow(
                    link=article_url or None,
                    type=TYPE_COMMENT,
                    body=replace_image_tokens(comment.text),
                ),
                account_for_comment_node(job, comment, mapping, author),
            )
        )
    current = [
        child
        for comment in job.article.comments
        for child in comment.children
    ]
    while current:
        nxt: list[CommentNode] = []
        for node in current:
            rows.append(
                _with_account(
                    MasterRow(
                        link=reply_target_value(node),
                        type=TYPE_REPLY,
                        body=replace_image_tokens(node.text),
                        result_link=article_url,
                    ),
                    account_for_comment_node(job, node, mapping, author),
                )
            )
            nxt.extend(node.children)
        current = nxt
    return rows


def build_master_rows(
    job: GatlingBrandJob,
    extra_exact_names: list[str] | tuple[str, ...] = (),
    *,
    include_daily_new_post: bool = True,
    image_location: str = "",
    proxy_book: ProxyBook | None = None,
    rng: random.Random | None = None,
    comment_id_count: int = COMMENT_ID_COUNT,
    author_id_count: int = DEFAULT_AUTO_ID_COUNT,
) -> list[MasterRow]:
    board_name = exact_board_name(
        job.board,
        job.cafe,
        extra_exact_names=extra_exact_names,
    )
    article_url = (job.cafe_article_url or "").strip()
    author, comment_map = resolve_job_accounts(
        job,
        proxy_book,
        rng=rng,
        comment_id_count=comment_id_count,
        author_id_count=author_id_count,
    )
    rows: list[MasterRow] = []

    if is_affiliate_cafe(job.cafe) and include_daily_new_post:
        if job.daily_post is None:
            raise GatlingPasteError(
                f"행 {job.row_number} {job.cafe}는 일상 글을 새글에 넣어야 합니다"
            )
        rows.append(
            _article_row(
                type_name=TYPE_NEW_POST,
                title=job.daily_post.title,
                body=job.daily_post.body,
                job=job,
                board_name=board_name,
                link=affiliate_board_link(job.cafe),
                author=author,
            )
        )
        rows.append(
            _article_row(
                type_name=TYPE_EDIT_POST,
                title=job.article.title,
                body=job.article.body,
                job=job,
                board_name=board_name,
                link=article_url,
                with_hashtag=True,
                image_location=image_location,
                author=author,
            )
        )
    elif is_affiliate_cafe(job.cafe):
        rows.append(
            _article_row(
                type_name=TYPE_EDIT_POST,
                title=job.article.title,
                body=job.article.body,
                job=job,
                board_name=board_name,
                link=article_url,
                with_hashtag=True,
                image_location=image_location,
                author=author,
            )
        )
    else:
        rows.append(
            _article_row(
                type_name=TYPE_NEW_POST,
                title=job.article.title,
                body=job.article.body,
                job=job,
                board_name=board_name,
                link=self_owned_board_link(job.cafe, job.board) or article_url,
                with_hashtag=True,
                image_location=image_location,
                author=author,
            )
        )

    rows.extend(
        _comment_and_reply_rows(
            job,
            article_url,
            author=author,
            comment_map=comment_map,
        )
    )
    return rows


def gatling_image_folder(gatling_path: str | Path) -> Path:
    path = Path(gatling_path)
    return path.with_name(f"{path.stem}_images")


EXPORT_EXCEL = "excel"
EXPORT_TXT = "txt"
EXPORT_BOTH = "both"
EXPORT_MODES = (EXPORT_EXCEL, EXPORT_TXT, EXPORT_BOTH)
TXT_TITLE_BODY = "제목본문"
TXT_COMMENTS_123 = "댓글1,2,3"
TXT_COMMENTS_45 = "댓글4,5"
TXT_REPLIES = "대댓글"
TXT_FOLDER_NAMES = (
    TXT_TITLE_BODY,
    TXT_COMMENTS_123,
    TXT_COMMENTS_45,
    TXT_REPLIES,
)
_UNSAFE_FILENAME_CHARS = re.compile(r'[<>:"/\\|?*]')
_RESERVED_CAFE_SLUGS = {"f-e", "ca-fe", "cafes", "articles", "menus"}
_COMPLETION_SLUG_ARTICLE_RE = re.compile(
    r"(?:https?://)?(?:m\.)?cafe\.naver\.com/([^/?#]+)/(\d+)(?:[/?#]|$)",
    re.IGNORECASE,
)
_COMPLETION_CAFES_ARTICLE_RE = re.compile(
    r"(?:https?://)?(?:m\.)?cafe\.naver\.com/(?:f-e/|ca-fe/)?"
    r"cafes/([^/]+)/articles/(\d+)",
    re.IGNORECASE,
)


def parse_export_mode(value: str) -> str:
    text = (value or "").strip()
    if text in EXPORT_MODES:
        return text
    raise GatlingPasteError("엑셀, TXT, 둘 다 중에서 고르세요")


def parse_row_range(value: str) -> tuple[int, int] | None:
    text = (value or "").strip().replace(" ", "")
    if not text:
        return None
    text = text.replace("~", "-").replace("–", "-").replace("—", "-")
    if "-" in text:
        left, right = text.split("-", 1)
        if not left.isdigit() or not right.isdigit():
            raise GatlingPasteError("행 범위는 137-146처럼 넣어 주세요")
        start, end = int(left), int(right)
    elif text.isdigit():
        start = end = int(text)
    else:
        raise GatlingPasteError("행 범위는 137-146처럼 넣어 주세요")
    if start < 2 or end < start:
        raise GatlingPasteError("행 범위는 2행부터, 작은 번호-큰 번호로 넣어 주세요")
    return start, end


def apply_row_range(
    jobs: list[GatlingBrandJob], row_range: tuple[int, int] | None
) -> list[GatlingBrandJob]:
    if row_range is None:
        return jobs
    start, end = row_range
    return [job for job in jobs if start <= job.row_number <= end]


def pick_txt_jobs(
    jobs: list[GatlingBrandJob],
) -> tuple[list[GatlingBrandJob], list[str]]:
    ready: list[GatlingBrandJob] = []
    skipped: list[str] = []
    for job in jobs:
        try:
            completion_txt_stem(job.completion_url)
        except GatlingPasteError as exc:
            skipped.append(f"행 {job.row_number}: {exc}")
            continue
        ready.append(job)
    return ready, skipped


def gatling_txt_folder(gatling_path: str | Path) -> Path:
    path = Path(gatling_path)
    return path.with_name(f"{path.stem}_txt")


def safe_txt_keyword(keyword: str) -> str:
    text = _UNSAFE_FILENAME_CHARS.sub("_", (keyword or "").strip())
    text = text.strip(" .")
    return text or "키워드"


def parse_completion_cafe_article(url: str) -> tuple[str, str]:
    text = (url or "").strip()
    if not text:
        raise GatlingPasteError(
            "완료 링크가 없습니다. "
            "https://cafe.naver.com/cantsb/3541968 형식으로 넣어 주세요"
        )
    match = _COMPLETION_SLUG_ARTICLE_RE.search(text)
    if match:
        slug, article = match.group(1), match.group(2)
        if slug.casefold() not in _RESERVED_CAFE_SLUGS:
            return slug, article
    match = _COMPLETION_CAFES_ARTICLE_RE.search(text)
    if match and not match.group(1).isdigit():
        return match.group(1), match.group(2)
    raise GatlingPasteError(
        "완료 링크에서 카페명과 글 번호를 찾지 못했습니다. "
        "https://cafe.naver.com/cantsb/3541968 형식으로 넣어 주세요"
    )


def completion_txt_stem(url: str) -> str:
    slug, article = parse_completion_cafe_article(url)
    slug = safe_txt_keyword(slug)
    return f"{slug}_{article}"


def require_txt_stems(jobs: list[GatlingBrandJob]) -> list[str]:
    stems: list[str] = []
    errors: list[str] = []
    for job in jobs:
        try:
            stems.append(completion_txt_stem(job.completion_url))
        except GatlingPasteError as exc:
            errors.append(f"행 {job.row_number}: {exc}")
    if errors:
        raise GatlingPasteError("\n".join(errors))
    return stems


def _unique_txt_stem(base: str, used: set[str]) -> str:
    stem = base
    index = 2
    while stem.casefold() in used:
        stem = f"{base}_{index}"
        index += 1
    used.add(stem.casefold())
    return stem


def _iter_comment_tree(nodes: list[CommentNode]) -> list[CommentNode]:
    items: list[CommentNode] = []
    for node in nodes:
        items.append(node)
        items.extend(_iter_comment_tree(node.children))
    return items


def _format_txt_body(text: str) -> str:
    return replace_image_tokens(text).strip()


def flatten_comment_txt(text: str) -> str:
    cleaned = replace_image_tokens(text or "")
    lines = [
        line.strip()
        for line in cleaned.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    ]
    return "$".join(line for line in lines if line)


def format_title_body_txt(article: ParsedArticle) -> str:
    title = _format_txt_body(article.title)
    body = _format_txt_body(article.body)
    if not title and not body:
        return ""
    if not title:
        return f"{body}\n"
    if not body:
        return f"{title}\n"
    return f"{title}\n\n{body}\n"


def format_comment_nodes_txt(nodes: list[CommentNode]) -> str:
    blocks: list[str] = []
    for node in nodes:
        body = flatten_comment_txt(node.text)
        if body:
            blocks.append(body)
    if not blocks:
        return ""
    return "\n\n".join(blocks) + "\n"


def split_manuscript_txt_parts(article: ParsedArticle) -> dict[str, str]:
    comments_123 = [node for node in article.comments if node.index in {1, 2, 3}]
    comments_45 = [node for node in article.comments if node.index not in {1, 2, 3}]
    replies = [
        node
        for root in article.comments
        for node in _iter_comment_tree(root.children)
    ]
    parts = {
        TXT_TITLE_BODY: format_title_body_txt(article),
        TXT_COMMENTS_123: format_comment_nodes_txt(comments_123),
        TXT_COMMENTS_45: format_comment_nodes_txt(comments_45),
        TXT_REPLIES: format_comment_nodes_txt(replies),
    }
    return {name: text for name, text in parts.items() if text.strip()}


def write_gatling_txt_files(
    jobs: list[GatlingBrandJob], dest_dir: str | Path
) -> list[Path]:
    folder = Path(dest_dir)
    folder.mkdir(parents=True, exist_ok=True)
    for name in TXT_FOLDER_NAMES:
        (folder / name).mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    used: set[str] = set()
    for job, base in zip(jobs, require_txt_stems(jobs), strict=True):
        stem = _unique_txt_stem(base, used)
        for name, content in split_manuscript_txt_parts(job.article).items():
            path = folder / name / f"{stem}.txt"
            path.write_text(content, encoding="utf-8-sig")
            written.append(path)
    return written


@dataclass(slots=True)
class GatlingExportResult:
    build: GatlingBuildResult
    start_row: int | None = None
    txt_paths: list[Path] = field(default_factory=list)


def collect_resolved_images(resolved: list[ResolvedImage], dest_dir: str | Path) -> str:
    folder = Path(dest_dir)
    folder.mkdir(parents=True, exist_ok=True)
    written: list[str] = []
    used: set[str] = set()
    for item in resolved:
        dest = folder / item.local_path.name
        if dest.name.casefold() in used or (
            dest.exists() and dest.resolve() != item.local_path.resolve()
        ):
            dest = folder / f"{item.file_id}_{item.local_path.name}"
        if dest.resolve() != item.local_path.resolve():
            dest.write_bytes(item.local_path.read_bytes())
        used.add(dest.name.casefold())
        written.append(str(dest))
    return "|".join(written)


def _resolve_job_images(
    job: GatlingBrandJob,
    *,
    image_resolver: GoogleDriveImageResolver | None,
    image_dir: str | Path | None,
    skipped: list[str],
) -> tuple[str, int]:
    if job.image_disabled or image_resolver is None or not (job.brand or "").strip():
        return "", 0
    source = f"{job.article.title}\n{job.article.body}"
    try:
        resolved = image_resolver.resolve_body(
            brand=job.brand,
            keyword=job.keyword,
            body=source,
            image_disabled=job.image_disabled,
            row_number=job.row_number,
        )
    except Exception as exc:
        skipped.append(f"행 {job.row_number}: 이미지 생략 ({exc})")
        return "", 0
    if not resolved:
        return "", 0
    if image_dir:
        return collect_resolved_images(resolved, image_dir), len(resolved)
    return "", len(resolved)


def build_gatling_master(
    brand_path: str | Path,
    daily_path: str | Path | None = None,
    *,
    skip_completed: bool = False,
    rng: random.Random | None = None,
    extra_exact_names: list[str] | tuple[str, ...] = (),
    manuscript_only: bool = False,
    brand: str = "",
    image_resolver: GoogleDriveImageResolver | None = None,
    image_dir: str | Path | None = None,
    existing_keys: set[tuple[str, str]] | None = None,
    proxy_book: ProxyBook | None = None,
    comment_id_count: int = COMMENT_ID_COUNT,
    author_id_count: int = DEFAULT_AUTO_ID_COUNT,
    row_range: tuple[int, int] | None = None,
    need_master_rows: bool = True,
) -> GatlingBuildResult:
    jobs, skipped = load_gatling_brand_jobs(
        brand_path,
        skip_completed=skip_completed,
        brand=brand,
    )
    jobs = apply_row_range(jobs, row_range)
    jobs, duplicate_skipped = drop_duplicate_manuscripts(jobs, existing_keys)
    skipped.extend(duplicate_skipped)
    if not need_master_rows:
        return GatlingBuildResult(
            rows=[],
            jobs=jobs,
            skipped=skipped,
            image_count=0,
            account_count=0,
        )
    include_daily_new_post = not manuscript_only
    affiliate_jobs = [job for job in jobs if is_affiliate_cafe(job.cafe)]
    if include_daily_new_post and affiliate_jobs:
        if not daily_path:
            raise GatlingPasteError(
                "제휴 카페 원고가 있어 일상 글 시트가 필요합니다"
            )
        assign_daily_posts(jobs, load_daily_posts(daily_path), rng=rng)
    needed_comments = normalize_auto_id_count(comment_id_count)
    needed_authors = normalize_auto_id_count(author_id_count)
    if proxy_book and affiliate_jobs and any(job.article.comments for job in affiliate_jobs):
        comment_count = len(proxy_book.comment_accounts())
        if comment_count < needed_comments:
            raise GatlingPasteError(
                f"양평맘·씨씨앙 댓글 아이디가 {needed_comments}개 필요합니다. "
                f"지금 {comment_count}개입니다"
            )
    self_jobs = [job for job in jobs if not is_affiliate_cafe(job.cafe)]
    if proxy_book and self_jobs and any(job.article.comments for job in self_jobs):
        self_comment_count = len(proxy_book.self_comment_accounts())
        if self_comment_count < needed_comments:
            raise GatlingPasteError(
                f"자사 카페 댓글 아이디가 {needed_comments}개 필요합니다. "
                f"지금 {self_comment_count}개입니다"
            )
    empty_affiliate = [
        job for job in affiliate_jobs if not (job.account or "").strip()
    ]
    empty_self = [job for job in self_jobs if not (job.account or "").strip()]
    if proxy_book and empty_affiliate:
        affiliate_author_count = len(proxy_book.affiliate_authors())
        if affiliate_author_count < needed_authors:
            raise GatlingPasteError(
                f"제휴 작성 아이디가 {needed_authors}개 필요합니다. "
                f"지금 {affiliate_author_count}개입니다"
            )
    if proxy_book and empty_self:
        self_author_count = len(proxy_book.self_authors())
        if self_author_count < needed_authors:
            raise GatlingPasteError(
                f"자사 작성 아이디가 {needed_authors}개 필요합니다. "
                f"지금 {self_author_count}개입니다"
            )
    assign_auto_authors(
        jobs, proxy_book, rng, author_id_count=needed_authors
    )

    rows: list[MasterRow] = []
    image_count = 0
    for job in jobs:
        if proxy_book and job.account and not proxy_book.find_author(job.account, job.cafe):
            skipped.append(
                f"행 {job.row_number}: 작성계정 {job.account}을 프록시 엑셀에서 찾지 못했습니다"
            )
        image_location, count = _resolve_job_images(
            job,
            image_resolver=image_resolver,
            image_dir=image_dir,
            skipped=skipped,
        )
        image_count += count
        rows.extend(
            build_master_rows(
                job,
                extra_exact_names=extra_exact_names,
                include_daily_new_post=include_daily_new_post,
                image_location=image_location,
                proxy_book=proxy_book,
                rng=rng,
                comment_id_count=needed_comments,
                author_id_count=needed_authors,
            )
        )
    return GatlingBuildResult(
        rows=rows,
        jobs=jobs,
        skipped=skipped,
        image_count=image_count,
        account_count=sum(1 for row in rows if row.account),
    )


def create_master_template(path: str | Path) -> Path:
    """Create an empty 기관총 마스터 workbook. Prefer .xlsm for paste-in-place."""
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = MASTER_SHEET_NAME
    sheet["A1"] = (
        "붙여넣기용 마스터입니다. 6행 열 이름은 기관총과 같습니다. "
        "제목·본문·댓글은 7행부터 이어 넣습니다."
    )
    for column, header in enumerate(MASTER_HEADERS, start=1):
        sheet.cell(MASTER_HEADER_ROW, column, header)
    workbook.save(output)
    return output


def write_master_xlsx(path: str | Path, rows: list[MasterRow]) -> Path:
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = MASTER_SHEET_NAME
    for column, header in enumerate(MASTER_HEADERS, start=1):
        sheet.cell(MASTER_HEADER_ROW, column, header)
    for offset, row in enumerate(rows):
        for column, value in enumerate(row.cells(), start=1):
            if value is None or value == "":
                continue
            sheet.cell(MASTER_HEADER_ROW + 1 + offset, column, value)
    workbook.save(output)
    return output


def build_and_write_master(
    brand_path: str | Path,
    output_path: str | Path,
    daily_path: str | Path | None = None,
    *,
    skip_completed: bool = False,
    rng: random.Random | None = None,
    extra_exact_names: list[str] | tuple[str, ...] = (),
    manuscript_only: bool = False,
    brand: str = "",
    image_resolver: GoogleDriveImageResolver | None = None,
    image_dir: str | Path | None = None,
    existing_keys: set[tuple[str, str]] | None = None,
    proxy_book: ProxyBook | None = None,
    comment_id_count: int = COMMENT_ID_COUNT,
    author_id_count: int = DEFAULT_AUTO_ID_COUNT,
) -> GatlingBuildResult:
    result = build_gatling_master(
        brand_path,
        daily_path,
        skip_completed=skip_completed,
        rng=rng,
        extra_exact_names=extra_exact_names,
        manuscript_only=manuscript_only,
        brand=brand,
        image_resolver=image_resolver,
        image_dir=image_dir,
        existing_keys=existing_keys,
        proxy_book=proxy_book,
        comment_id_count=comment_id_count,
        author_id_count=author_id_count,
    )
    write_master_xlsx(output_path, result.rows)
    return result


def paste_manuscripts_into_gatling(
    brand_path: str | Path,
    gatling_path: str | Path,
    daily_path: str | Path | None = None,
    *,
    skip_completed: bool = False,
    rng: random.Random | None = None,
    extra_exact_names: list[str] | tuple[str, ...] = (),
    manuscript_only: bool = True,
    brand: str = "",
    image_resolver: GoogleDriveImageResolver | None = None,
    image_dir: str | Path | None = None,
    proxy_book: ProxyBook | None = None,
    comment_id_count: int = COMMENT_ID_COUNT,
    author_id_count: int = DEFAULT_AUTO_ID_COUNT,
) -> tuple[GatlingBuildResult, int]:
    """Append Google Sheet title/body/comments into a recognized 기관총 마스터."""
    result = build_gatling_master(
        brand_path,
        daily_path,
        skip_completed=skip_completed,
        rng=rng,
        extra_exact_names=extra_exact_names,
        manuscript_only=manuscript_only,
        brand=brand,
        image_resolver=image_resolver,
        image_dir=image_dir,
        existing_keys=load_existing_manuscript_keys(gatling_path),
        proxy_book=proxy_book,
        comment_id_count=comment_id_count,
        author_id_count=author_id_count,
    )
    if not result.rows:
        preview = "\n".join(result.skipped[:8])
        extra = f"\n외 {len(result.skipped) - 8}건" if len(result.skipped) > 8 else ""
        raise GatlingPasteError(
            "제목·본문이 같은 원고는 이미 기관총에 있어 넣지 않았습니다."
            + (f"\n{preview}{extra}" if preview else "")
        )
    start_row = append_master_rows(gatling_path, result.rows)
    return result, start_row


def export_gatling_output(
    brand_path: str | Path,
    *,
    mode: str,
    gatling_path: str | Path | None = None,
    txt_dir: str | Path | None = None,
    daily_path: str | Path | None = None,
    skip_completed: bool = False,
    rng: random.Random | None = None,
    extra_exact_names: list[str] | tuple[str, ...] = (),
    brand: str = "",
    image_resolver: GoogleDriveImageResolver | None = None,
    image_dir: str | Path | None = None,
    proxy_book: ProxyBook | None = None,
    comment_id_count: int = COMMENT_ID_COUNT,
    author_id_count: int = DEFAULT_AUTO_ID_COUNT,
    row_range: tuple[int, int] | None = None,
) -> GatlingExportResult:
    chosen = parse_export_mode(mode)
    want_excel = chosen in {EXPORT_EXCEL, EXPORT_BOTH}
    want_txt = chosen in {EXPORT_TXT, EXPORT_BOTH}
    if want_excel and not gatling_path:
        raise GatlingPasteError("기관총 엑셀 파일을 선택해 주세요")
    if want_txt and not txt_dir:
        raise GatlingPasteError("TXT 저장 폴더를 선택해 주세요")
    existing_keys = (
        load_existing_manuscript_keys(gatling_path) if want_excel and gatling_path else None
    )
    result = build_gatling_master(
        brand_path,
        daily_path,
        skip_completed=skip_completed,
        rng=rng,
        extra_exact_names=extra_exact_names,
        manuscript_only=not want_excel,
        brand=brand,
        image_resolver=image_resolver,
        image_dir=image_dir,
        existing_keys=existing_keys,
        proxy_book=proxy_book,
        comment_id_count=comment_id_count,
        author_id_count=author_id_count,
        row_range=row_range,
        need_master_rows=want_excel,
    )
    start_row = None
    txt_paths: list[Path] = []
    txt_jobs = result.jobs
    if want_txt:
        txt_jobs, txt_skipped = pick_txt_jobs(result.jobs)
        result.skipped.extend(txt_skipped)
        if not txt_jobs:
            preview = "\n".join(txt_skipped[:8])
            extra = f"\n외 {len(txt_skipped) - 8}건" if len(txt_skipped) > 8 else ""
            raise GatlingPasteError(
                "TXT로 만들 완료 링크가 없습니다. "
                "https://cafe.naver.com/cantsb/3541968 형식으로 넣어 주세요"
                + (f"\n{preview}{extra}" if preview else "")
            )
    if want_excel:
        if not result.rows:
            preview = "\n".join(result.skipped[:8])
            extra = f"\n외 {len(result.skipped) - 8}건" if len(result.skipped) > 8 else ""
            raise GatlingPasteError(
                "제목·본문이 같은 원고는 이미 기관총에 있어 넣지 않았습니다."
                + (f"\n{preview}{extra}" if preview else "")
            )
        start_row = append_master_rows(gatling_path, result.rows)
    if want_txt:
        txt_paths = write_gatling_txt_files(txt_jobs, txt_dir)
    return GatlingExportResult(build=result, start_row=start_row, txt_paths=txt_paths)
