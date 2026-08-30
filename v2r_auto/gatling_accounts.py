from __future__ import annotations

import csv
import random
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from .content import CommentNode


KNOWN_AFFILIATE_COMMENT_IDS = {
    "quilliant",
    "hunnede",
    "prtchht",
    "chocobbn",
    "chenallo",
    "colpith",
}
AFFILIATE_COMMENT_COUNT = 6
COMMENT_ID_COUNT = AFFILIATE_COMMENT_COUNT
KIND_AFFILIATE_AUTHOR = "affiliate_author"
KIND_AFFILIATE_COMMENT = "affiliate_comment"
KIND_SELF_AUTHOR = "self_author"
KIND_SELF_COMMENT = "self_comment"
KIND_OTHER = "other"
COMMENT_KINDS = {KIND_AFFILIATE_COMMENT, KIND_SELF_COMMENT}

ID_HEADERS = {"아이디", "id", "계정", "작성계정"}
PASSWORD_HEADERS = {"비번", "비밀번호", "패스워드", "password", "pwd"}
CHROME_HEADERS = {"크롬번", "크롬번호", "크롬", "chrome"}
CATEGORY_HEADERS = {"카테고리"}
IP_HEADERS = {"아이피:포트", "아이피", "ip", "프록시", "proxy"}
AFFILIATE_AUTHOR_CATEGORIES = {"제휴"}
AFFILIATE_COMMENT_CATEGORIES = {"제휴 댓", "제휴댓", "제휴 댓글", "제휴댓글"}
SELF_AUTHOR_CATEGORIES = {"자사"}
SELF_COMMENT_CATEGORIES = {"자사 댓", "자사댓", "자사 댓글", "자사댓글"}
AFFILIATE_CAFES = {"씨씨앙", "양평맘"}


def _proxy_error(message: str) -> Exception:
    from .gatling_paste import GatlingPasteError

    return GatlingPasteError(message)


def _is_affiliate_cafe(cafe: str) -> bool:
    return (cafe or "").strip() in AFFILIATE_CAFES


@dataclass(slots=True)
class ProxyAccount:
    account: str
    password: str = ""
    chrome_number: int | str = ""
    ip: str = ""
    category: str = ""
    kind: str = KIND_OTHER


@dataclass(slots=True)
class ProxyBook:
    path: Path
    accounts: list[ProxyAccount] = field(default_factory=list)
    message: str = ""
    recognized: bool = False

    def comment_accounts(self) -> list[ProxyAccount]:
        blocked = {
            item.account.casefold()
            for item in self.accounts
            if item.kind == KIND_SELF_COMMENT
        }
        return [
            item
            for item in self.accounts
            if item.kind == KIND_AFFILIATE_COMMENT
            and item.account.casefold() not in blocked
        ]

    def self_comment_accounts(self) -> list[ProxyAccount]:
        blocked = {
            item.account.casefold()
            for item in self.accounts
            if item.kind == KIND_AFFILIATE_COMMENT
        }
        blocked.update(KNOWN_AFFILIATE_COMMENT_IDS)
        return [
            item
            for item in self.accounts
            if item.kind == KIND_SELF_COMMENT
            and item.account.casefold() not in blocked
        ]

    def affiliate_authors(self) -> list[ProxyAccount]:
        return [item for item in self.accounts if item.kind == KIND_AFFILIATE_AUTHOR]

    def self_authors(self) -> list[ProxyAccount]:
        return [item for item in self.accounts if item.kind == KIND_SELF_AUTHOR]

    def find_author(self, account: str, cafe: str = "") -> ProxyAccount | None:
        wanted = (account or "").strip().casefold()
        if not wanted:
            return None
        matches = [
            item for item in self.accounts if item.account.casefold() == wanted
        ]
        if not matches:
            return None
        if _is_affiliate_cafe(cafe):
            preferred = [
                item for item in matches if item.kind == KIND_AFFILIATE_AUTHOR
            ]
        else:
            preferred = [item for item in matches if item.kind == KIND_SELF_AUTHOR]
        authors_only = [
            item for item in matches if item.kind not in COMMENT_KINDS
        ]
        return (preferred or authors_only or matches)[0]


def _cell(value: object) -> str:
    return str(value or "").replace("\r\n", "\n").replace("\r", "\n").strip()


def _chrome_value(value: object) -> int | str:
    if value is None or value == "":
        return ""
    if isinstance(value, bool):
        return _cell(value)
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    text = _cell(value)
    if text.isdigit():
        return int(text)
    return text


def _header_index(headers: list[str], candidates: set[str]) -> int:
    wanted = {item.casefold() for item in candidates}
    for index, header in enumerate(headers):
        if header.casefold() in wanted:
            return index
    return -1


def _category_matches(label: str, compact: str, candidates: set[str]) -> bool:
    if label in candidates:
        return True
    return compact in {item.replace(" ", "") for item in candidates}


def _kind_from_category(category: str, account: str) -> str:
    label = _cell(category)
    compact = label.replace(" ", "")
    if _category_matches(label, compact, AFFILIATE_COMMENT_CATEGORIES):
        return KIND_AFFILIATE_COMMENT
    if _category_matches(label, compact, SELF_COMMENT_CATEGORIES):
        return KIND_SELF_COMMENT
    if label in AFFILIATE_AUTHOR_CATEGORIES:
        return KIND_AFFILIATE_AUTHOR
    if label in SELF_AUTHOR_CATEGORIES:
        return KIND_SELF_AUTHOR
    if account.casefold() in KNOWN_AFFILIATE_COMMENT_IDS:
        return KIND_AFFILIATE_COMMENT
    if label:
        return KIND_OTHER
    return KIND_AFFILIATE_AUTHOR


def _parse_proxy_rows(rows: list[list[object]]) -> list[ProxyAccount]:
    header_row = -1
    headers: list[str] = []
    for index, values in enumerate(rows):
        cleaned = [_cell(value) for value in values]
        if _header_index(cleaned, ID_HEADERS) >= 0 and (
            _header_index(cleaned, CHROME_HEADERS) >= 0
            or _header_index(cleaned, PASSWORD_HEADERS) >= 0
        ):
            header_row = index
            headers = cleaned
            break
    if header_row < 0:
        return []

    id_col = _header_index(headers, ID_HEADERS)
    password_col = _header_index(headers, PASSWORD_HEADERS)
    chrome_col = _header_index(headers, CHROME_HEADERS)
    category_col = _header_index(headers, CATEGORY_HEADERS)
    ip_col = _header_index(headers, IP_HEADERS)
    accounts: list[ProxyAccount] = []
    seen: set[str] = set()
    for values in rows[header_row + 1 :]:
        account = _cell(values[id_col] if id_col < len(values) else "")
        if not account or account.casefold() in seen:
            continue
        category = _cell(values[category_col] if 0 <= category_col < len(values) else "")
        item = ProxyAccount(
            account=account,
            password=_cell(values[password_col] if 0 <= password_col < len(values) else ""),
            chrome_number=_chrome_value(
                values[chrome_col] if 0 <= chrome_col < len(values) else ""
            ),
            ip=_cell(values[ip_col] if 0 <= ip_col < len(values) else ""),
            category=category,
            kind=_kind_from_category(category, account),
        )
        seen.add(account.casefold())
        accounts.append(item)
    return accounts


def _read_xlsx_rows(path: Path) -> list[list[object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows = [list(row) for row in sheet.iter_rows(max_col=20, values_only=True)]
            parsed = _parse_proxy_rows(rows)
            if parsed:
                return rows
        return []
    finally:
        workbook.close()


def _read_csv_rows(path: Path) -> list[list[object]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        return [list(row) for row in csv.reader(stream)]


def recognize_proxy_workbook(path: str | Path) -> ProxyBook:
    file_path = Path(path)
    if not file_path.exists():
        raise _proxy_error(f"프록시 엑셀 파일이 없습니다: {file_path}")
    suffix = file_path.suffix.casefold()
    if suffix == ".xls":
        raise _proxy_error(
            "프록시 파일은 .xlsx 또는 .csv로 저장해 주세요. .xls는 열 수 없습니다"
        )
    if suffix not in {".xlsx", ".xlsm", ".csv"}:
        raise _proxy_error("프록시 파일은 .xlsx, .xlsm, .csv만 선택할 수 있습니다")

    try:
        rows = _read_csv_rows(file_path) if suffix == ".csv" else _read_xlsx_rows(file_path)
    except Exception as exc:
        from .gatling_paste import GatlingPasteError

        if isinstance(exc, GatlingPasteError):
            raise
        raise _proxy_error(
            "프록시 엑셀을 열지 못했습니다. 파일이 열려 있으면 닫고 다시 선택해 주세요"
        ) from exc

    accounts = _parse_proxy_rows(rows)
    if not accounts:
        return ProxyBook(
            path=file_path,
            message=(
                "프록시 엑셀에서 아이디·크롬번호·비번 열을 찾지 못했습니다. "
                "첫 시트에 '아이디', '크롬번', '비번', '카테고리'가 있는지 확인해 주세요"
            ),
        )
    comments = sum(1 for item in accounts if item.kind == KIND_AFFILIATE_COMMENT)
    self_comments = sum(1 for item in accounts if item.kind == KIND_SELF_COMMENT)
    affiliate_authors = sum(1 for item in accounts if item.kind == KIND_AFFILIATE_AUTHOR)
    self_authors = sum(1 for item in accounts if item.kind == KIND_SELF_AUTHOR)
    message = (
        f"프록시 엑셀로 확인했습니다. 제휴 작성 {affiliate_authors}개 / "
        f"제휴 댓글 {comments}개 / 자사 작성 {self_authors}개 / "
        f"자사 댓글 {self_comments}개"
    )
    if comments and comments < COMMENT_ID_COUNT:
        message += (
            f". 양평맘·씨씨앙 댓글을 넣으려면 제휴 댓글 아이디가 "
            f"{COMMENT_ID_COUNT}개 필요합니다"
        )
    if self_comments and self_comments < COMMENT_ID_COUNT:
        message += (
            f". 자사 카페 댓글을 넣으려면 자사 댓글 아이디가 "
            f"{COMMENT_ID_COUNT}개 필요합니다"
        )
    return ProxyBook(path=file_path, accounts=accounts, message=message, recognized=True)


def require_proxy_workbook(path: str | Path) -> ProxyBook:
    book = recognize_proxy_workbook(path)
    if not book.recognized:
        raise _proxy_error(book.message)
    return book


def article_kind(article_type: str) -> str:
    text = (article_type or "").strip()
    if "후기" in text:
        return "후기형"
    if "질문" in text:
        return "질문형"
    return ""


def _comment_id_map(
    job,
    comment_accounts: list[ProxyAccount],
    rng: random.Random,
    *,
    pool_label: str,
) -> dict[str, ProxyAccount]:
    if not job.article.comments:
        return {}
    if len(comment_accounts) < COMMENT_ID_COUNT:
        raise _proxy_error(
            f"{pool_label} 댓글 아이디가 {COMMENT_ID_COUNT}개 필요합니다. "
            f"지금 {len(comment_accounts)}개입니다"
        )
    chosen = rng.sample(comment_accounts, COMMENT_ID_COUNT)
    special_label = "대대댓글2" if article_kind(job.article_type) == "후기형" else "대대대댓글2"
    labels = ("댓글1", "댓글2", special_label, "댓글3", "댓글4", "댓글5")
    return dict(zip(labels, chosen))


def account_for_comment_node(
    job,
    node: CommentNode,
    comment_map: dict[str, ProxyAccount],
    author: ProxyAccount | None,
) -> ProxyAccount | None:
    if node.depth == 0:
        return comment_map.get(node.label)
    if node.label == "대대댓글2":
        if article_kind(job.article_type) == "후기형":
            return comment_map.get("대대댓글2")
        return comment_map.get("댓글2")
    if node.label == "대대대댓글2":
        if article_kind(job.article_type) == "후기형":
            return author
        return comment_map.get("대대대댓글2")
    return author


def resolve_job_accounts(
    job,
    book: ProxyBook | None,
    rng: random.Random | None = None,
) -> tuple[ProxyAccount | None, dict[str, ProxyAccount]]:
    if book is None:
        return None, {}
    author = book.find_author(job.account, job.cafe)
    comment_map: dict[str, ProxyAccount] = {}
    if not job.article.comments:
        return author, comment_map
    rng = rng or random.SystemRandom()
    if _is_affiliate_cafe(job.cafe):
        comment_map = _comment_id_map(
            job,
            book.comment_accounts(),
            rng,
            pool_label="양평맘·씨씨앙",
        )
    else:
        comment_map = _comment_id_map(
            job,
            book.self_comment_accounts(),
            rng,
            pool_label="자사 카페",
        )
    return author, comment_map
